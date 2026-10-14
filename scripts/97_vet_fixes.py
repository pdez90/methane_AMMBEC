"""97_vet_fixes.py -- consistency and copy fixes from the full vet of the 14 Oct 2026 manuscript and SI
(applied to the output of script 96).

Main text
  * Introduction: the "(for example XX)" placeholder for the declining emission intensity of the
    DJB is replaced with a citation of Daley et al. (2026, the fall-2021 DJB aircraft study of
    emission rates, intensities and trends); references are renumbered in order of first citation.
  * 3.3: every flight in Table 1 has at least three contributing transects, so the sentence about
    single-transect flights "identified accordingly in Table 1" is reworded.
  * 3.2: "neighboring urban legs" -> "transects"; 4.1: "5,8, and 10 July L1" spacing.
  * Heading punctuation (2.2., 3.1., 3.2.), the author line, and stray double spaces / a missing
    space after a full stop (including inside tracked insertions; those keep their tracking).
SI
  * S7.1: the main text compares both inventories through HRRR-STILT footprints, not only EPA.
  * S1.1.1: the median ground-minus-airborne mixing-height difference is over the 13 flights.
  * S5 and Figure S12 caption punctuation; Table S5 flight labels in the date style of the other tables.
  * S9 heading names the transport analysis it now holds.
New text is red; pure spacing/punctuation fixes keep the formatting of the text they touch.

usage: python3 scripts/97_vet_fixes.py <main96.docx> <si96.docx> <main_out> <si_out>
"""
import sys, re, copy, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
from docx.text.run import Run
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
a = ap.parse_args()
W_ = lambda t: qn("w:" + t)

def t_fix(d, old, new, count, label):
    """Plain-text fix inside single w:t elements anywhere (direct runs, hyperlinks, tracked insertions)."""
    n = 0
    for p in L.all_paragraphs(d):
        for t in p._p.iter(W_("t")):
            if t.text and old in t.text:
                n += t.text.count(old); t.text = t.text.replace(old, new)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    if count is not None and n != count: sys.exit(f"{label}: {old!r} fixed {n}x, expected {count}x")
    return n

def seam_fix(d, left, right, new_left, label):
    """Fix a seam between two consecutive w:t elements (e.g. a tracked insertion next to plain text):
    the element whose text ends with `left`, followed by one starting with `right`, gets `new_left`."""
    n = 0
    for p in L.all_paragraphs(d):
        ts = list(p._p.iter(W_("t")))
        for x, y in zip(ts, ts[1:]):
            if (x.text or "").endswith(left) and (y.text or "").startswith(right):
                x.text = x.text[: len(x.text) - len(left)] + new_left; n += 1
                x.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    if n != 1: sys.exit(f"{label}: seam {left!r}|{right!r} found {n}x, expected 1x")

def dspace_fix(d, anchor, label):
    """Collapse the double space inside `anchor` (visible w:t text, may span runs or tracked insertions)."""
    hits = 0
    for p in L.all_paragraphs(d):
        ts = list(p._p.iter(W_("t"))); txt = "".join(t.text or "" for t in ts)
        i = txt.find(anchor)
        if i < 0: continue
        hits += 1; k = i + anchor.index("  ") + 1          # position of the second space
        pos = 0
        for t in ts:
            n = len(t.text or "")
            if pos <= k < pos + n:
                j = k - pos; assert t.text[j] == " "; t.text = t.text[:j] + t.text[j + 1:]
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); break
            pos += n
    if hits != 1: sys.exit(f"{label}: {anchor!r} found {hits}x, expected 1x")

def cite_after(d, anchor_old, anchor_new, sup, label):
    """Replace anchor_old by anchor_new (red) and put a red superscript citation run right after it."""
    O.do_replace(d, anchor_old, anchor_new, 1, label)
    p = O.find_para(d, anchor_new, label)
    r = next(r for r in p.runs if r.text == anchor_new)
    nr = copy.deepcopy(r._r); r._r.addnext(nr); run = Run(nr, p); run.text = sup
    run.font.color.rgb = L.RED; run.font.superscript = True

# ================================ main ================================
m = docx.Document(a.main_in); LM = "main text"
cite_after(m, "with a declining emission intensity over time (for example XX).", "with a declining emission intensity over time.", "20", LM)
O.do_replace(m, "Flights with only one qualifying transect retained a point estimate but could not support a transect-block bootstrap interval and are identified accordingly in Table 1.",
             "A flight with only one qualifying transect would keep a point estimate without a transect-block bootstrap interval; every flight in Table 1 has at least three contributing transects.", 1, LM)
O.do_replace(m, "the cleanest air sampled on neighboring urban legs,", "the cleanest air sampled on neighboring urban transects,", 1, LM)
t_fix(m, "5,8, and 10 July L1", "5, 8, and 10 July L1", 1, LM)
t_fix(m, "2.2.CDPHE facility ", "2.2. CDPHE facility ", 1, LM)
t_fix(m, "3.1 Constant-altitude flight transects", "3.1. Constant-altitude flight transects", 1, LM)
t_fix(m, "3.2 Background and enhancement", "3.2. Background and enhancement", 1, LM)
t_fix(m, "Xinrong Ren*,Rainer", "Xinrong Ren*, Rainer", None, LM)
seam_fix(m, ".", "The estimated", ". ", LM)                     # her tracked insertions: "campaign." + "The estimated"
seam_fix(m, "final ", " data based", "final", LM)
for old in ("crossed the  Denver", "(sections 3.3-3.4).  The", "Aeris  background", "measuring  C2H6",
            "(Figure S6).  The", "fossil signal.  Source"):
    dspace_fix(m, old, LM)
# author line: collapse runs of spaces between names (text only, no names changed)
ap_ = next(p for p in m.paragraphs if O.vis_text(p).startswith("Priyanka deSouza"))
for t in ap_._p.iter(W_("t")):
    if t.text: t.text = re.sub(r",\s{2,}", ", ", re.sub(r" {2,}", " ", t.text))
L.renumber_refs(m, LM, start="Abstract")
m.save(a.main_out)

# ================================ SI ================================
s = docx.Document(a.si_in); LS = "SI"
O.do_replace(s, "The main text therefore compares the inventories with the aircraft domain-wide and, for the EPA inventory, through HRRR-STILT footprints (section S9).",
             "The main text therefore compares the inventories with the aircraft domain-wide and, for both inventories, through HRRR-STILT footprints (section S9).", 1, LS)
O.do_replace(s, "by up to 2.9 km (median 0.75 km)", "by up to 2.9 km (median difference over the 13 flights 0.75 km)", 1, LS)
t_fix(s, "S5 DJB basin ethane-to-methane ratio in 2024", "S5. DJB basin ethane-to-methane ratio in 2024", 1, LS)
seam_fix(s, "Figure S12", " Ethane:methane", "Figure S12:", LS)
O.do_replace(s, "S9. Fossil fraction versus north-easterly flow", "S9. Fossil fraction versus north-easterly flow and modeled transport", 1, LS)
# Table S5 flight labels: 20240712 (L1) -> 2024-07-12 (L1)
tS5 = next(tb for tb in s.tables if tb.rows[0].cells[0].text.strip() == "Flight" and "Basin transects" in tb.rows[0].cells[1].text)
nfix = 0
for row in tS5.rows[1:]:
    for t in row.cells[0]._tc.iter(W_("t")):
        if t.text and re.match(r"2024\d{4}", t.text):
            t.text = re.sub(r"(2024)(\d\d)(\d\d)", r"\1-\2-\3", t.text, count=1); nfix += 1
assert nfix == 22, nfix
s.save(a.si_out)

# ---- checks ----------------------------------------------------------------------------------
tm = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.main_out)))
ts = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.si_out)))
for b in ("XX", "neighboring urban legs", "campaign.The", "2.2.CDPHE", "identified accordingly in Table 1"):
    assert b not in tm, f"main: {b!r} remains"
assert re.search(r"\n3\.1\. Constant", tm) and re.search(r"\n3\.2\. Background", tm), "main: 3.1./3.2. headings"
for b in ("for the EPA inventory, through HRRR", "S5 DJB", "Figure S12 Ethane", "\n20240"):
    assert b not in ts, f"SI: {b!r} remains"
print("main: Daley et al. cited; references renumbered. checks passed")
