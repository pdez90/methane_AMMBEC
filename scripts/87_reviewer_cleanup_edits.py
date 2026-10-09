"""87_reviewer_cleanup_edits.py -- reviewer's final substantive points and cleanup list (4 Oct 2026),
applied to the author-edited manuscript and SI.

Substance
  1. Table S9: the Carbon Mapper persistence-weighted rates are read from v2_pointsource_hotspots.csv
     (102 and 420 kg/h; emission_auto is already persistence-weighted, script 59), replacing the
     stale 68 and 332, which double-weighted persistence.
  2. SI S3 / section 3.2: the 16-27% bound on distribution gas applies to the strongest leg-level
     floor (13 July L1, 0.054) only; the weaker floors on the other three flights do not exclude an
     ethane-poor source (per-flight bounds from two_endmember_floor_flights.csv, script 52).
  3. Section 3.4: the altitude/time filters are omitted for the SI's reason (a single lidar 40 km
     from the city), not because "the ratio does not require a mixing-height term"; the SI test
     concerns the fossil fraction only, not the emission estimates.
  4. Section 3.1: legs are classified by their mean position and kept whole (tag_region); the
     altitude criterion is the leg's altitude range, maximum minus minimum (detect_level_legs).
  5. SI S2.2: a road plume may include overlapping sources; the all-drive median of the in-box
     ethane-bearing plumes (0.084) is what falls in the regional reference range (flight-window
     median 0.070 lies just below).
  6. Abstract: "regional ... gradient" softened to a comparison with the adjacent basin; the 0.059
     break-even sentence restored.
  Also: Table S2 29 June mixing height restored to the PUMAS value in denver_curtain_screen.csv
  (1214 m; it had reverted to 2000).
Cleanup
  "two conditional checks"; data-availability statement qualified (reproducibility from the
  aircraft files plus the ground, lidar and inventory inputs; ARL car and CDPHE access); "ARL car
  use"; "C2H6:CH4 fits"; "substantial,"; Figure S7 "rises at thresholds above 50 ppb";
  Table S3 caption; the three plain "GRA2PES" spellings follow the documents' GRA²PES. The equations were checked in the formatted document (Word math objects
  carry the Delta and inequality symbols); no change.
All inserted text is red.

usage: python3 scripts/87_reviewer_cleanup_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
from docx_edit_lib import R, T, runs_text, all_paragraphs

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.OUTDIR = str(OUT)
rd = lambda f: list(csv.DictReader(open(OUT / f)))

# ---- numbers from the outputs ------------------------------------------------------------------
hs = {r["nearest_facility"]: r for r in rd("v2_pointsource_hotspots.csv")}
cm_dads = str(round(float(hs["Denver Arapahoe Disposal Site (DADS)"]["carbonmapper_kg_hr"])))
cm_tower = str(round(float(hs["Tower Road landfill"]["carbonmapper_kg_hr"])))
assert (cm_dads, cm_tower) == ("102", "420"), (cm_dads, cm_tower)

cs = {r["flight"]: r for r in rd("denver_curtain_screen.csv")}
blh_0629 = str(round(float(cs["20240629_R0_L1"]["blh_m"]))); assert cs["20240629_R0_L1"]["blh_src"] == "lidar (PUMAS)"
assert blh_0629 == "1214", blh_0629

fl = {r["flight"]: r for r in rd("two_endmember_floor_flights.csv")}
pct = lambda v: round(100 * float(v))
strong = max(fl.values(), key=lambda r: float(r["beta_eff_floor"]))
assert strong["flight"] == "20240713_R0_L1"
s_lo, s_hi = pct(strong["p_dist_max_0.061_ec_lo"]), pct(strong["p_dist_max_0.061_ec_hi"])
s_lo2, s_hi2 = pct(strong["p_dist_max_0.0813_ec_lo"]), pct(strong["p_dist_max_0.0813_ec_hi"])
others = [r for r in fl.values() if r["flight"] != strong["flight"]]
o_lo = min(pct(r["p_dist_max_0.061_ec_lo"]) for r in others); o_hi = max(pct(r["p_dist_max_0.061_ec_hi"]) for r in others)
o_floor = (min(float(r["beta_eff_floor"]) for r in others), max(float(r["beta_eff_floor"]) for r in others))
assert (s_lo, s_hi, s_lo2, s_hi2, o_lo, o_hi) == (16, 27, 43, 60, 34, 100), (s_lo, s_hi, s_lo2, s_hi2, o_lo, o_hi)
assert o_hi == 100 and o_lo >= 50 - 16   # the weaker floors allow a majority of distribution gas

# ---- main text -----------------------------------------------------------------------------------
# 6. abstract
R("placing Denver toward the biogenic end of a regional fossil-to-biogenic gradient from the producing basin into the metropolitan area.",
  "indicating a more biogenic source mix over the sampled urban area than reported for the adjacent producing basin; the median "
  "exceeds 50% fossil only for an effective endmember below 0.059 mol mol−1.")
R("Together, these results place Denver toward the biogenic end of a regional fossil-to-biogenic gradient from the producing basin into the "
  "metropolitan area, rather than in conflict with basin-scale source apportionment.",
  "Together, these results indicate a more biogenic source mix over the sampled urban area than reported for the adjacent producing basin, "
  "rather than a conflict with basin-scale source apportionment.")
# cleanup: grammar
R("the CDPHE facility surveys and the ARL car uses it", "the CDPHE facility surveys and the ARL car use it")
R("a substantial-, and during this campaign apparently majority, component", "a substantial, and during this campaign apparently majority, component")
# 2. section 3.2: the floor bound is a per-leg statement; the strongest leg carries the 16-27% bound
R("provided two additional, although conditional, constraints on the effective endmember.",
  "provided two conditional checks on the effective endmember.")
R("across four flights, indicating that, under an assumed lean basin ratio of 0.061 mol mol−1, a very ethane-poor fossil source is unlikely to have dominated those legs (section S3).",
  f"across four flights. The strongest of these leg-level constraints (13 July L1, {float(strong['beta_eff_floor']):.3f} mol mol−1) indicates that, "
  "under an assumed lean basin ratio of 0.061 mol mol−1, a very ethane-poor fossil source is unlikely to have dominated that leg; the lower "
  "floors on the other three flights do not exclude it (section S3).")
# 5. the in-box ground plumes: it is their median that falls in the reference range
R("while those inside the analysis box (0.084 mol mol−1) fall within the regional reference range",
  "while the all-drive median of those inside the analysis box (0.084 mol mol−1) falls within the regional reference range")
# 4. leg definition and classification, as coded (detect_level_legs, tag_region)
R("stable heading, and limited end-to-end altitude variation.", "stable heading, and a limited altitude range.")
R("and holds its altitude to within 150 m end to end.", "and whose altitude range (maximum minus minimum GPS altitude over the leg) is at most 150 m.")
R("We then classified the legs geographically. Legs within the Denver-metro box (39.50 to 39.95°N, 105.20 to 104.55°W) were classified as ‘urban’. "
  "Legs north of 40.05°N, which lie over the adjacent gas field, were classified as ‘basin’.",
  "We then classified each leg geographically by its mean position and retained the whole leg: legs whose mean position lies within the "
  "Denver-metro box (39.50 to 39.95°N, 105.20 to 104.55°W) were classified as ‘urban’, and legs whose mean position lies north of 40.05°N, "
  "over the adjacent gas field, as ‘basin’.")
R("at least three urban legs and at least 50 in-box urban-leg samples.", "at least three urban legs and at least 50 urban-leg samples.")
# 3. section 3.4: why the filters are omitted (the SI's reasoning); the test concerns the fossil fraction only
R(" because the ratio itself does not require a mixing-height term. Moreover, our boundary-layer heights come from a single lidar about 40 km "
  "from the urban core, making in-boundary-layer filters potentially unreliable on shallow-mixing days. We therefore fit the full in-box "
  "urban-leg data (section S3).",
  " and instead test the sensitivity of the fossil fraction to them (section S3). Although the enhancement-ratio calculation does not "
  "explicitly require a boundary-layer height, sampling altitude still matters: individual source plumes or air above the mixed layer may "
  "not represent the broader urban source mixture. We did not impose the filters in the primary analysis because our boundary-layer heights "
  "come from a single Doppler lidar about 40 km from central Denver and may not represent mixing conditions over the city, particularly on "
  "days with a shallow boundary layer. We therefore fit all urban-leg samples; with the filters applied, no flight's fossil fraction changes "
  "by more than 3 percentage points (Table S3). That test concerns the fossil fraction only and does not validate the emission estimates.")
# cleanup: data availability
R("All analysis code, including the run-all script that regenerates every number, table and figure in this paper and its Supporting "
  "Information from those files, is at",
  "All analysis code, including the run-all script that regenerates the tables and figures in this paper and its Supporting Information "
  "from the aircraft files together with the ground, lidar and inventory inputs listed below, is at")
R("archived as PDF copies and stored as psco_gas_quality.csv.",
  "archived as PDF copies and stored as psco_gas_quality.csv. The NOAA ARL Air Resources Car data (1 s ICARTT files) and the CDPHE "
  "facility-survey files are not redistributed with the code and are available from NOAA ARL and CDPHE, respectively, on request; the "
  "plume-level tables derived from them (arc_ground_plumes.csv, mobile_survey_summary.csv) and the Doppler-lidar mixing heights used "
  "for each flight are included in the repository's results folder.")

# ---- SI ----------------------------------------------------------------------------------------------
# 1. Table S9 (table 8): Carbon Mapper rates
T("si", 8, 1, 8, cm_dads, expect="68")
T("si", 8, 2, 8, cm_tower, expect="332")
# Table S2 (table 1): 29 June mixing height, PUMAS lidar
T("si", 1, 20, 4, blh_0629, expect="2000")
# 5. S2.2
R("A road plume is a single source sampled at close range,",
  "A road plume may be dominated by a nearby source but can include overlapping emissions, and it is sampled at close range,", "si")
R("while the ethane-bearing plumes inside the box fall within the regional reference range of 0.081 to 0.110 mol mol−1.",
  "while the all-drive median of the ethane-bearing plumes inside the box, 0.084 mol mol−1, falls within the regional reference range of "
  "0.081 to 0.110 mol mol−1 (the flight-window median of 0.070 mol mol−1, from 3 plumes, lies just below it).", "si")
# 2. S3: the bound is per leg
R("those floors exclude an ethane-poor fossil source as the majority of what the tightly correlated legs sampled: with East Coast gas, at most "
  "16 to 27% of the fossil methane on those legs could have been distribution gas. For a basin ratio of 0.0813 the bound rises to 43 to 60%, "
  "so at that ratio the floors no longer exclude it, whereas the Denver composition is unconstrained by the floor.",
  f"the strongest leg-level constraint (13 July L1, floor {float(strong['beta_eff_floor']):.3f}) limits distribution gas to at most {s_lo} to {s_hi}% of the "
  "fossil methane on that leg with East Coast gas, so an ethane-poor fossil source could not have been the majority of what that leg sampled; "
  f"the weaker floors on the other three flights ({o_floor[0]:.3f} to {o_floor[1]:.3f}) allow {o_lo} to {o_hi}% and do not exclude it. For a basin ratio of "
  f"0.0813 the bound on 13 July L1 rises to {s_lo2} to {s_hi2}%, so at that ratio even the strongest floor no longer excludes it, whereas the Denver "
  "composition is unconstrained by the floor.", "si")
# cleanup
R("Per-flight fossil fraction from the full in-box urban-leg data and from the QC-filtered subset",
  "Per-flight fossil fraction from all urban-leg samples (no altitude or time filter) and from the QC-filtered subset", "si")
R(", we retain only observations with methane enhancements above 20 ppb.", " fits, we retain only observations with methane enhancements above 20 ppb.", "si")
R("The dashed gray line rises above 50 ppb only because", "The dashed gray line rises at thresholds above 50 ppb only because", "si")

# ---- anchors must not straddle a hyperlinked citation --------------------------------------------
def sentinel_text(p):
    out = []
    for ch in p._p:
        if ch.tag == qn("w:r"): out.append("".join(t.text or "" for t in ch.iter(qn("w:t"))))
        elif ch.tag == qn("w:hyperlink"): out.append("\x00")
    return "".join(out)
for which, path in (("main", a.main_in), ("si", a.si_in)):
    d = docx.Document(path); txt = "\n".join(sentinel_text(p) for p in all_paragraphs(d))
    for anc, _, doc_, kind, _ in L.E:
        if doc_ == which and kind == "replace" and anc not in txt:
            sys.exit(f"{which}: anchor crosses a hyperlinked citation: {anc[:80]!r}")

L.apply(a.main_in, a.main_out, "main", "main text")
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- GRA2PES: the author-edited documents write the inventory as GRA²PES (superscript 2) throughout; the three
# remaining plain spellings follow suit (the superscript is red so the change is visible) ------------------------
import copy
from docx.text.run import Run
def gra2pes_superscript(d):
    n = 0
    for p in all_paragraphs(d):
        for r in list(p.runs):
            if "GRA2PES" not in r.text: continue
            pieces = []
            for k, part in enumerate(r.text.split("GRA2PES")):
                if k: pieces += [("GRA", False), ("2", True), ("PES", False)]
                pieces.append((part, False))
            prev = r._r
            for txt, sup in pieces:
                if not txt: continue
                el = copy.deepcopy(r._r); prev.addnext(el); prev = el
                nr = Run(el, p); nr.text = txt
                if sup: nr.font.superscript = True; nr.font.color.rgb = L.RED
            r._r.getparent().remove(r._r); n += 1
    return n
for path, label in ((a.main_out, "main text"), (a.si_out, "SI")):
    d = docx.Document(path); n = gra2pes_superscript(d); d.save(path); print(f"{label}: {n} GRA2PES -> GRA²PES")

# ---- equations: the Word math objects carry the symbols the reviewer could not see in pasted text ----
import zipfile, re
x = zipfile.ZipFile(a.main_out).read("word/document.xml").decode()
eqs = ["".join(re.findall(r"<m:t[^>]*>([^<]*)</m:t>", m.group(0))) for m in re.finditer(r"<m:oMath>.*?</m:oMath>", x, re.S)]
assert any("≤" in e for e in eqs) and sum("Δ" in e for e in eqs) >= 2, eqs
print(f"main text: {len(eqs)} equation objects; Δ and ≤ present")

# ---- final checks --------------------------------------------------------------------------------
for path, label, bads in ((a.main_out, "main text", ("fossil-to-biogenic gradient", "ARL car uses", "substantial-,", "150 m end to end", "end-to-end altitude", "in-box urban-leg",
                                                     "two additional, although", "regenerates every number")),
                          (a.si_out, "SI", ("single source sampled", "plumes inside the box fall within", "on those legs could have been",
                                            "full in-box", "rises above 50 ppb"))):
    d = docx.Document(path); t = "\n".join(runs_text(p) for p in all_paragraphs(d))
    for bad in bads: assert bad not in t, f"{label}: {bad!r} remains"
    assert not any("GRA2PES" in r.text for p in all_paragraphs(d) for r in p.runs), f"{label}: plain GRA2PES remains"
    if label == "SI":
        assert [c.text for c in d.tables[8].rows[1].cells][8] == cm_dads and [c.text for c in d.tables[8].rows[2].cells][8] == cm_tower
        assert d.tables[1].rows[20].cells[4].text == blh_0629
print("checks passed")
