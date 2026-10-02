"""78_final_items_edits.py -- the last pre-submission items from the 1 Oct 2026 code review.

  1. Spatial comparison (main 4.3, SI S7.1): the Spearman tests counted repeated inventory cells
     (26 aircraft cells lie in 7 EPA cells and 15 GRA2PES cells). Report the tests with one value
     per inventory cell alongside (script 58, spatial_match_summary.csv).
  2. SI Table S5: single-leg flights printed a zero-width interval ("0.0545 to 0.0545"); with one
     leg there is no leg-block interval (R/ratios.R now returns NA), so the cell reads "n/a (one leg)".
  3. The two ground platforms on 9 July (main 2, SI S2.3): the ARL car drove part of the CSL
     laboratory's route about 11 minutes behind it (script 77, ground_platform_overlap.csv).
  4. Data availability: the two ground datasets, with a bracketed placeholder for the archive
     location and data-use statement, which have to come from the PIs.
  5. SI reference list: entries 16-19, cited in the SI text but missing from its list.

Every number is read from the outputs; edits are exact-anchor and red.

Item 3 and the CSL part of item 4 were later removed again by scripts/79_drop_csl_mobilelab_edits.py
(superseded), when the CSL Mobile Laboratory data were dropped from the paper.

usage: python3 scripts/78_final_items_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, copy, re, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
import docx_edit_lib as L
from docx_edit_lib import R

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.OUTDIR = str(OUT)

# ---- numbers ------------------------------------------------------------------------------
SUM = {r["inventory"].split(" (")[0]: r for r in csv.DictReader(open(OUT / "spatial_match_summary.csv"))}
need = "spearman_rho_per_inventory_cell"
assert all(need in r for r in SUM.values()), "rerun scripts/58_spatial_match.R: summary lacks the per-inventory-cell test"
epa, v11, v2 = SUM["EPA gridded GHGI 2020"], SUM["GRA2PES v1.1"], SUM["GRA2PES v2.0beta"]
def rho(r): return f"{float(r['spearman_rho_per_inventory_cell']):.2f}"
def pv(r): return f"{float(r['p_value_per_inventory_cell']):.2f}"
n_air = int(epa["n_cells"]); n_epa = int(float(epa["n_inventory_cells"]))
n_gra = sorted({int(float(v11["n_inventory_cells"])), int(float(v2["n_inventory_cells"]))})
assert all(float(r["p_value_per_inventory_cell"]) > 0.05 for r in (epa, v11, v2)), "a per-cell test is now significant; rewrite"
gra_txt = f"{n_gra[0]} GRA2PES cells" if len(n_gra) == 1 else f"{n_gra[0]} to {n_gra[1]} GRA2PES cells"

OV = {r["quantity"]: r["value"] for r in csv.DictReader(open(OUT / "ground_platform_overlap.csv"))}
lag_min = round(float(OV["lag_arc_minus_csl_median_s"]) / 60)
lag_iqr = [round(float(x) / 60) for x in OV["lag_iqr_s"].split(" to ")]
side_km = float(OV["min_separation_same_minute_m"]) / 1000
print(f"per-cell Spearman: EPA {rho(epa)} (p {pv(epa)}, n {n_epa}); v1.1 {rho(v11)} ({pv(v11)}); v2 {rho(v2)} ({pv(v2)}); "
      f"overlap {OV['csl_share_shared_pct']}% of CSL route, ARC +{lag_min} min ({lag_iqr[0]}-{lag_iqr[1]}), "
      f"closest same-minute {side_km:.1f} km, ARC metro {OV['arc_metro_seconds_on_csl_route']}/{OV['arc_metro_seconds']} s on route")

# ---- 1. spatial comparison ------------------------------------------------------------------
R("(Spearman ρ = -0.09, p = 0.68)",
  f"(Spearman ρ = -0.09, p = 0.68; {rho(epa)}, p = {pv(epa)}, with one value per EPA cell, since the {n_air} aircraft "
  f"cells lie in only {n_epa} of them)")
R("which says whether the aircraft sampled the fossil-rich or the biogenic-rich part of each inventory.",
  "which says whether the aircraft sampled the fossil-rich or the biogenic-rich part of each inventory. Because "
  f"several aircraft cells can fall in one inventory cell (the {n_air} lie in {n_epa} EPA cells and {gra_txt}), the "
  "rank correlation is also computed with one value per inventory cell, the median aircraft fraction in that cell: "
  f"ρ = {rho(epa)} (p = {pv(epa)}) for the EPA inventory, {rho(v11)} (p = {pv(v11)}) for v1.1 and {rho(v2)} "
  f"(p = {pv(v2)}) for v2.0beta, none of them significant.", "si")

# ---- 3. the two ground platforms on 9 July ----------------------------------------------------
R("drive sampled appreciably inside the analysis box south of 39.80°N,",
  "drive sampled appreciably inside the analysis box south of 39.80°N, partly along the foothills roads the CSL "
  f"laboratory had driven about {lag_min} minutes earlier,")
R("so within the flight window 124,429 of its one-second records fall north of 40.05°N against 29,550 inside the box.",
  "so within the flight window 124,429 of its one-second records fall north of 40.05°N against 29,550 inside the box. "
  "On 9 July the car and the CSL Mobile Laboratory (section S2.2) drove partly the same roads: "
  f"{OV['csl_share_shared_pct']}% of the laboratory's transect was also driven by the car within 30 minutes and 150 m, "
  f"along the foothills side of the metropolitan area between {OV['shared_lat_range'].replace(' to ', ' and ')}°N from "
  f"{OV['shared_window_utc']} UTC, with the car passing a median of {lag_min} minutes after the laboratory "
  f"(interquartile range {lag_iqr[0]} to {lag_iqr[1]} minutes). The two were never side by side; their closest approach "
  f"within the same minute was {side_km:.1f} km. Of the car's {int(OV['arc_metro_seconds']):,} one-second records south "
  f"of 39.80°N that day, {int(OV['arc_metro_seconds_on_csl_route']):,} lie on this shared route (script 77).", "si")

# ---- 4. data availability -------------------------------------------------------------------
R("archived as PDF copies and transcribed in psco_gas_quality.csv.",
  "archived as PDF copies and transcribed in psco_gas_quality.csv. The ground-based data are the NOAA CSL Mobile "
  "Laboratory files for 9 July 2024 and the NOAA Air Resources Laboratory Air Resources Car files (ARC-Suite, 1 s "
  "ICARTT; PI X. Ren) for 24 June to 15 July 2024 [archive location and data-use statement to be confirmed with the PIs].")

L.apply(a.main_in, a.main_out, "main", "main text")
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- 2. SI Table S5: single-leg intervals ---------------------------------------------------
si = docx.Document(a.si_out)
t5 = next((tb for tb in si.tables if " | ".join(c.text for c in tb.rows[0].cells).startswith("Flight | Basin legs | Legs used")), None)
assert t5 is not None, "SI Table S5 not found"
hdr = [c.text for c in t5.rows[0].cells]; i_used, i_ci = hdr.index("Legs used"), hdr.index("95% CI")
n5 = 0
for row in t5.rows[1:]:
    if row.cells[i_used].text.strip() == "1":
        lo_hi = row.cells[i_ci].text.split(" to ")
        if len(lo_hi) == 2 and lo_hi[0].strip() == lo_hi[1].strip():
            cp = row.cells[i_ci].paragraphs[0]
            for r_ in cp.runs[1:]: r_._r.getparent().remove(r_._r)
            cp.runs[0].text = "n/a (one leg)"; cp.runs[0].font.color.rgb = L.RED; n5 += 1
print(f"Table S5: {n5} single-leg interval(s) set to n/a")

# ---- 5. SI references 16-19 -----------------------------------------------------------------
main = docx.Document(a.main_in)
def main_ref(pattern):
    hits = [p.text for p in main.paragraphs if re.match(r"\(\d+\)\s", p.text.strip()) and re.search(pattern, p.text)]
    assert hits, f"main reference matching {pattern!r} not found"
    return re.sub(r"^\(\d+\)\s*", "", hits[0].strip())
SI_REFS = [(16, main_ref(r"Public Service Company of Colorado")), (17, main_ref(r"^\(\d+\)\s+Plant, G\.")),
           (18, main_ref(r"Terracon")), (19, main_ref(r"Carbon Mapper Data Portal"))]
refs = [p for p in si.paragraphs if re.match(r"\(\d+\)\s", p.text.strip())]
have = {int(re.match(r"\((\d+)\)", p.text.strip()).group(1)) for p in refs}
last = refs[-1]; added = 0
for n, txt in SI_REFS:
    if n in have: continue
    newp = copy.deepcopy(last._p); last._p.addnext(newp)
    p = docx.text.paragraph.Paragraph(newp, last._parent)
    for child in list(p._p):
        if child.tag != docx.oxml.ns.qn("w:pPr"): p._p.remove(child)
    run = p.add_run(f"({n})\t{txt}"); run.font.color.rgb = L.RED
    last = p; added += 1
si.save(a.si_out)
print(f"SI references added: {added} ({', '.join(str(n) for n, _ in SI_REFS if n not in have)})")
