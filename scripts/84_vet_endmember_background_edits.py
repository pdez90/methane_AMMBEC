"""84_vet_endmember_background_edits.py -- fixes from the vet of 3 Oct 2026 (main text, SI, figures).

Endmember framing (author decision, 3 Oct 2026)
  * The 0.061 "2021 aircraft flux partition" is not supported by the cited study: Daley et al.
    report 2.6 t/h ethane and 21.4 t/h methane, 63% of it from oil and gas, so ethane against the
    oil-and-gas methane derived from that partition is about 0.10 mol/mol (about 0.065 against total
    methane, which includes livestock methane and is not a fossil endmember). The unsupported
    "2.1 vs 18.4 t/h" derivation is deleted; Daley is cited as supporting about 0.10.
  * The regional reference values are 0.0813 (production-area ground ratio), about 0.10 (2021) and
    0.110 (delivered gas). Medians across them: 27 to 36%. They are presented as regional reference
    values, not as a bound on Denver's effective endmember. 0.061 is kept only as an explicitly
    assumed lean-basin scenario (48%, two flights above 50%), and no sentence claims the fraction
    stays below 50% across all Denver-relevant endmembers.
  * Comparison with earlier cities: the seven-city airborne study (Floerchinger et al.) reports
    natural-gas fractions of 20-80% in summer and 32-100% in winter, so Denver lies toward the lower
    end of published estimates rather than below all of them.
Methods described as the code runs them
  * Aircraft background: the rolling 5th percentile runs over each flight's urban legs joined end to
    end (basin legs for the basin ratio); the window spans neighbouring legs. The along-track
    alternative is reported in SI S4 (median 25% against 29%; two CH4:CO flights would fail r >= 0.7).
  * Aircraft speed about 70 m/s (window about 20 km of flight); level-leg heading test; mixing height
    for the transect rates; NE sector 22.5-90 deg; background scatter over the 11 flights with urban
    legs; facility-separation scan background (flight-wide 5th percentile; 3 of 7 separable under a
    rolling background, conclusion unchanged).
Numbers
  * ARL car urban ambient ratio 0.023 (median of unrounded daily slopes; was rounded twice to 0.024).
  * East Coast break-even 0.52 to 0.66 (unrounded 0.6645; was rounded twice to 0.67).
  * Table S4: 2025 Jan C2H6 10.39 (10.385), 2024 Aug CH4 83.6 (83.55).
  * STILT observed/predicted interquartile range 2.6 to 5.4 (R type-7 quantiles, as elsewhere).
  * Fig 4B / Table S8 aircraft fossil-fraction range 0.31 to 0.42 across 0.0813-0.110.
Other: Fig 4A no longer cited for sub-sectors it does not show; Fig 3 shows 10 July; relaxed Carbon
Mapper encounter of 3 July L1; SI figure/table call-outs; CO2 leave-one-leg-out failures named;
abbreviations defined at first use; ~11% landfill share attributed via Kille et al.; "top-down
atmospheric studies"; references (Petron pages/DOI, Wunch article number, Ngulat preprint entry).
Layout: Arial Unicode MS runs reset to the body font; figures kept with their captions; table rows
not split across pages; Table S7 header unit made white; Limitation 5 sentence un-bolded.
Main references are renumbered in order of first citation (Floerchinger et al. is now cited in 4.1).

usage: python3 scripts/84_vet_endmember_background_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
(inputs: the _2026-10-02 documents written by scripts 79-83)
"""
import sys, csv, re, argparse, copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
from docx_edit_lib import R, I, T, runs_text, all_paragraphs

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.OUTDIR = str(OUT)
rd = lambda f: list(csv.DictReader(open(OUT / f)))
def r2(x, n):  # half-up rounding of a decimal string or float
    return str(Decimal(str(x)).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))

# ---- numbers from the outputs ------------------------------------------------------------------
ARC = {r["quantity"]: r["value"] for r in csv.DictReader(open(OUT / "arc_ground_summary.csv"))}
amb = float(ARC["ambient_slope_median_window_urban"])
assert abs(amb - 0.0235) < 0.0005, amb
AMB = r2(amb, 3)                                                       # 0.023
EC = rd("two_endmember_breakeven_eastcoast.csv")
ec = {(float(r["beta_dist"]), float(r["beta_djb"])): float(r["p_dist_breakeven"]) for r in EC}
ref_be = [v for (bd, bj), v in ec.items() if bj in (0.0813, 0.102)]
lean_be = [v for (bd, bj), v in ec.items() if bj == 0.061]
EC_REF = f"{round(100*min(ref_be))}-{round(100*max(ref_be))}"          # 36-66
EC_LEAN = f"{round(100*min(lean_be))}-{round(100*max(lean_be))}"       # 6-10
EC102 = f"{r2(min(v for (bd,bj),v in ec.items() if bj==0.102),2)} to {r2(max(v for (bd,bj),v in ec.items() if bj==0.102),2)}"
assert EC102 == "0.52 to 0.66", EC102
TE = {r["case"]: r for r in rd("two_endmember_flights.csv")}
m_arc = int(TE["DJB only, 2024 ARC ground ratio"]["median_fossil_pct"])       # 36
m_lean = int(TE["DJB only, assumed lean basin ratio 0.061"]["median_fossil_pct"])  # 48
m_del = int(TE["distribution only, DENVER delivered gas (PSCo Jun-Jul 2024)"]["median_fossil_pct"])  # 27
top_lo = int(TE["distribution only, DENVER delivered gas (PSCo Jun-Jul 2024)"]["20240713_R0_L2"])  # 44
top_hi = int(TE["DJB only, 2024 ARC ground ratio"]["20240713_R0_L2"])         # 59
assert (m_arc, m_lean, m_del, top_lo, top_hi) == (36, 48, 27, 44, 59), (m_arc, m_lean, m_del, top_lo, top_hi)
SM = rd("spatial_match_summary.csv")[0]["aircraft_median_ff_lo_hi"]
assert SM == "0.31 to 0.42", SM
REF = f"{m_del} to {m_arc}%"; REFd = f"{m_del}-{m_arc}%"

# ---- main text: endmember framing ----------------------------------------------------------------
R("At an effective fossil endmember of 0.102 mol mol−1, the median fossil fraction was 29% (2 to 47% across flights). "
  "Across contemporary Denver-relevant endmembers (0.061-0.110 mol mol−1), the median ranged from 27 to 48% and would "
  "exceed 50% only below 0.059 mol mol−1, whereas previous airborne urban studies reported 40 to 88%.",
  "At an effective fossil endmember of 0.102 mol mol-1, the campaign-median fossil fraction was 29% (2 to 47% across "
  "flights), toward the lower end of fractions reported in previous airborne urban studies. Across regional reference "
  f"endmembers (0.081-0.110 mol mol-1) the median was {REF}; it would exceed 50% only for an effective endmember below "
  "0.059 mol mol-1.")
R("Contemporary DJB constraints span approximately 0.061-0.10 mol mol-1, including the 2021 aircraft flux partition "
  "(0.061; section S5), the ground-level ratio of 0.0813 measured in production areas during this campaign, and a 2021 "
  "basin-wide emission ratio",
  "Contemporary DJB reference values span approximately 0.081-0.10 mol mol-1: the ground-level ratio of 0.0813 measured "
  "in production areas during this campaign (an ambient ratio, which bounds the basin source composition from below), "
  "about 0.10 from the ethane and oil-and-gas methane emissions of the 2021 basin mass balance,[[^22]] and a 2021 "
  "basin-wide emission ratio")
R("pipeline gas measured in East Coast cities (0.019-0.037 mol mol-1).",
  "pipeline gas measured in East Coast cities (0.019-0.037 mol mol-1).[[^3]]")
R("because it lies within the contemporary range defined by the DJB and Denver delivered-gas constraints.",
  "because it lies within the range of these regional reference values (0.081-0.110 mol mol-1).")
R("the approximately threefold decline in the basin ethane-to-methane ratio reported between 2015 and 2021,",
  "the lower ratios implied by the decline in basin ethane emissions reported between 2015 and 2021,")
R("  and values below those previously reported.", " and values below those previously reported.")
R("The campaign median therefore remains biogenic-dominated across the Denver-relevant endmembers considered.",
  "The campaign median is therefore biogenic-dominated for every regional reference value considered (0.081-0.110 "
  "mol mol-1). These reference values do not by themselves bound Denver's effective fossil endmember, so leaner values "
  "are evaluated as assumed scenarios (section S3).")
R("With the measured Denver delivered-gas composition, the campaign median ranged from 27 to 48% fossil and never reached 50%.",
  f"With the measured Denver delivered-gas composition, the campaign median ranged from {REF} fossil for the regional "
  f"basin values and reached at most {m_lean}% under an assumed lean basin ratio of 0.061 mol mol-1; it never reached 50%.")
R("with the required distribution-gas contribution ranging from approximately 6% to 67% depending on the assumed basin "
  "and distribution-gas ratios.",
  f"with the required distribution-gas contribution ranging from approximately {EC_REF.replace('-', '% to ')}% for basin "
  f"ratios of 0.0813 and 0.102 ({EC_LEAN.replace('-', '% to ')}% under the assumed 0.061), depending on the "
  "distribution-gas ratio.")
R("indicating that, for a basin gas ratio near 0.061 mol mol−1, a very ethane-poor",
  "indicating that, under an assumed lean basin ratio of 0.061 mol mol-1, a very ethane-poor")
R(" but remains above the 0.059 threshold at which the campaign median becomes majority fossil.",
  " although it rests on the assumed oil-and-gas share and cannot be compared reliably with the ratio of about 0.10 "
  "implied by the 2021 basin emissions (section S5); it remains above the 0.059 threshold at which the campaign median "
  "becomes majority fossil.")
R("Using this lower range would increase", "Using this lower, conditional range would increase")
R("The abstract and Discussion report the median across the Denver-relevant 0.061-0.110 mol mol-1 range, whereas Table 1, "
  "Figure 2, and section 4.1 report values at the adopted 0.102 calibration.",
  "Table 1 and Figure 2 report values at the adopted 0.102 calibration; the abstract, section 4.1, and the Discussion also "
  "report the campaign median across the regional reference values (0.081-0.110 mol mol-1).")
R("Across the Denver-relevant endmember range of 0.061-0.110 mol mol-1, the campaign median ranges from 27 to 48% (sections 3.2, S3).",
  f"Across the regional reference endmembers (0.081-0.110 mol mol-1), the campaign median ranges from {REF}; an assumed "
  f"lean basin ratio of 0.061 mol mol-1 would give {m_lean}% (sections 3.2, S3).")
R("gives an ambient ratio of 0.024 mol mol−1 as a median over 9 days", f"gives an ambient ratio of {AMB} mol mol-1 as a median over 9 days")
R("Overall, this indicates a less fossil-influenced urban methane regime than has been reported in previous airborne "
  "studies of other metropolitan areas.",
  "Our campaign-median fossil fraction of 29% falls within the range of previous airborne urban estimates. Floerchinger "
  "et al. reported natural-gas contributions of 20 to 80% in summer and 32 to 100% in winter across seven U.S. urban "
  "areas,[[^36]] and studies of Los Angeles, Baltimore-Washington, New York and East Coast cities reported 40 to 88%.[[^2,3,5,25]] "
  "Denver therefore lies toward the lower end of published estimates, although differences in season, sampling footprint, "
  "and assumed ethane-to-methane endmembers limit direct comparisons.")
R("the campaign median remains biogenic-dominated across the Denver-relevant endmember range (section 3.2).",
  "the campaign median remains biogenic-dominated across the regional reference endmembers (section 3.2).")
R("the flights that reach a fossil majority under the leanest basin endmember- 9 July and 13 July L2- are",
  "the flights that reach or exceed 50% at the 0.0813 reference value (9 July and 13 July L2) are")
R("ranges from 0.31 to 0.57 across the Denver-relevant endmember range.", f"ranges from {SM} across the regional reference endmembers.")
L.W("Across the fossil endmembers relevant to Denver- from contemporary DJB constraints",
    "Across the regional reference endmembers, from the DJB values of 0.081-0.10 to the measured composition of gas "
    f"delivered to the city (0.110 mol mol-1; sections 3.2, S5), the campaign-median fossil fraction ranges from {REF}. At "
    "the adopted calibration of 0.102 mol mol-1, the median is 29%, with individual flights ranging from 2 to 47%. At "
    f"0.0813, 9 July reaches the 50% line and 13 July L2 exceeds it ({top_hi}%); no flight does at the adopted calibration "
    "or the delivered-gas composition. The campaign median exceeds 50% only for an effective endmember below 0.059 mol "
    f"mol-1 (Figures S4 and S5); an assumed lean basin ratio of 0.061 mol mol-1 gives {m_lean}%. These reference values do "
    "not bound Denver's effective fossil endmember from below, so the absolute fraction remains conditional on it.")
R("The more distinctive difference is in the source mix.", "The source mix also differs.")
R("Denver's median fossil fraction (27 to 48% across the Denver-relevant endmembers, 29% at the adopted 0.102 calibration) "
  "lies below all of them at that calibration and below or within the lower part of the 40-88% range across the "
  "Denver-relevant endmembers (its 48% upper end lies inside the Baltimore-Washington range); its most fossil-influenced "
  "flight (44 to 79% across these calibrations) overlaps the lower end of that range.",
  "Floerchinger et al. reported natural-gas fractions of 20 to 80% in summer and 32 to 100% in winter across seven "
  f"cities.[[^36]] Denver's median fossil fraction (29% at the adopted 0.102 calibration, {REF} across the regional "
  "reference endmembers) is below the Los Angeles, Baltimore-Washington, New York and East Coast values and toward the "
  f"lower end of the seven-city summer range; its most fossil-influenced flight ({top_lo} to {top_hi}% across these "
  "endmembers) falls within the ranges reported elsewhere. The Denver results therefore extend evidence that biogenic "
  "sources can contribute substantially to urban methane emissions, rather than establishing a uniquely low urban fossil "
  "contribution. Comparisons also require care because our fossil category includes the sources represented by the "
  "assumed regional endmember, whereas Floerchinger et al. estimated natural-gas contributions.")
R("compared with an airborne biogenic fraction of 52-73% across the Denver-relevant endmembers.",
  f"compared with an airborne biogenic fraction of {100-m_arc}-{100-m_del}% across the regional reference endmembers.")
R("above the 27-48% campaign median inferred from airborne ethane across the Denver-relevant endmembers",
  f"above the {REFd} campaign median inferred from airborne ethane across the regional reference endmembers")
R("Table 1, Figure 2, and section 4.1 use an endmember of 0.102 mol mol-1, whereas the abstract and Discussion report the "
  "campaign median across the Denver-relevant range of 0.061-0.110 mol mol-1. The lower part of this range is constrained "
  "by the 2021 aircraft flux partition (0.061) and the ground-level production-area ratio measured during this campaign "
  "(0.0813; section S5), which",
  "Table 1 and Figure 2 use an endmember of 0.102 mol mol-1; the abstract, section 4.1, and the Discussion also report the "
  "campaign median across the regional reference values of 0.081-0.110 mol mol-1. The lowest of these is the ground-level "
  "production-area ratio measured during this campaign (0.0813; section S5), which")
R(" and so bounds the fossil source composition from below rather than estimating it (the 2021 value is a flux partition "
  "of ethane against oil-and-gas methane and is not diluted in this way):",
  " and so bounds the basin fossil source composition from below rather than estimating it:")
R("Using these two values as endmembers raises the campaign median to 36-48% fossil and brings two of seven flights to or above 50%.",
  f"Using it as the endmember raises the campaign median to {m_arc}% fossil and brings two of seven flights to or above "
  "50%. The 2021 basin mass balance implies about 0.10 for oil-and-gas emissions (section S5).[[^22]]")
R("while those inside the analysis box (0.084 mol mol−1) fall within the adopted range (section S2.2).",
  "while those inside the analysis box (0.084 mol mol-1) fall within the regional reference range of 0.081-0.110 mol "
  "mol-1 (section S2.2).")
R("the campaign median could exceed 50% once distribution gas contributed roughly 6-67% of urban fossil methane, depending "
  "on the assumed basin and delivered-gas ratios (Figure S5).",
  f"the campaign median could exceed 50% once distribution gas contributed roughly {EC_REF}% of urban fossil methane for "
  f"basin ratios of 0.0813 and 0.102 ({EC_LEAN}% under an assumed lean basin ratio of 0.061), depending on the "
  "delivered-gas ratio (Figure S5).")
R("The absolute fractions therefore remain endmember-dependent, but none of the Denver-relevant source compositions or "
  "mixtures considered produces a fossil-majority campaign median.",
  "The absolute fractions therefore remain endmember-dependent: none of the regional reference compositions, alone or "
  "mixed with Denver's delivered gas, produces a fossil-majority campaign median, but these reference values do not "
  "exclude a leaner effective endmember for the urban fossil mixture, and values below 0.059 mol mol-1 would yield one.")
R("The central result is nevertheless robust across the Denver-relevant endmember range:",
  "The central result holds across the regional reference endmembers:")

# ---- main text: methods described as the code runs them ------------------------------------------
R("identically to every 1 Hz series from all three platforms.",
  "identically to every 1 Hz series from all three platforms (the facility-separation scan of section S12, which only "
  "locates plume samples near one complex, uses a flight-wide 5th percentile instead).")
R("On the aircraft, it produced enhancements in CH4, C2H6, CO, and CO2 that feed the source-apportionment and flux analyses.",
  "On the aircraft, it produced enhancements in CH4, C2H6, CO, and CO2 that feed the source-apportionment and flux "
  "analyses; there the series is each flight's urban level legs joined end to end (basin legs for the basin ratio, "
  "section S5), so near a leg boundary the window spans the adjacent leg and the background is the cleanest air sampled "
  "on neighboring urban legs, which retains the urban-scale enhancement (section S4).")
R("about 5 minutes for the roughly 50 m/s aircraft and", "about 5 minutes (300 samples) for the aircraft and")
R("Because the aircraft moves about five times faster and samples broader, urban-scale enhancements, its window is "
  "correspondingly longer, spanning roughly 15 km of track,",
  "Because the aircraft moves several times faster (about 70 m/s) and samples broader, urban-scale enhancements, its "
  "window is longer, spanning roughly 20 km of flight,")
R("(median 4.3 ppb across flights)", "(median 4.3 ppb, at most 10 ppb, across the 11 flights with urban legs)")
R("and the heading varies by at most 30°,",
  "and the heading stays within 30° of its value at the start of the leg (gaps of up to 10 s in level flight are bridged),")
R("and to scale the single-transect rates downwind of inventory hotspots (section S11).",
  "and, as the median over the urban legs' time window, to scale the single-transect rates downwind of inventory hotspots (section S11).")
R("arriving from the north-east quadrant,", "with wind from the north-east (22.5 to 90°, north-north-east to east),")
R("a fossil share of ~ 26%.", "a fossil share of ~ 26% (Figure 4A).")
R("0.9 t CH4/h to natural-gas distribution (Figure 4A).", "0.9 t CH4/h to natural-gas distribution.")
R(" Only the sub-split within each class is taken from the inventory.", "")
R("when the urban ethane-methane slopes are nearly flat (Figure 3).", "when the urban ethane-methane slopes are nearly flat (10 July in Figure 3).")
R("With only one usable CH4:CO2 flight, however, it should not be interpreted as an independent replication of the CH4:CO range.",
  "With only one usable CH4:CO2 flight, however, this agreement should not be interpreted as an independent replication of the CH4:CO range.")
# abbreviations at first use; attribution wording
R("CDPHE conducted repeated vehicle-based methane surveys",
  "The Colorado Department of Public Health and Environment (CDPHE) conducted repeated vehicle-based methane surveys")
R("Mixing height was taken from a NOAA CSL Doppler-lidar site",
  "Mixing height was taken from a NOAA Chemical Sciences Laboratory (CSL) Doppler-lidar site")
R("and HRRR drives the STILT footprints (section S9).",
  "and the High-Resolution Rapid Refresh (HRRR) model drives the Stochastic Time-Inverted Lagrangian Transport (STILT) "
  "footprints (section S9).")
R("U.S. gridded EPA greenhouse-gas inventory (2020 Express Extension, 0.1 degree)",
  "U.S. gridded EPA Greenhouse Gas Inventory (GHGI; 2020 Express Extension, 0.1 degree)")
R("(4 km, July 2023; NOAA Chemical Sciences Laboratory, beta release of June 2026)", "(4 km, July 2023; NOAA CSL, beta release of June 2026)")
R("The EPA 2020 National Emissions Inventory county-level CO totals", "The EPA 2020 National Emissions Inventory (NEI) county-level CO totals")
R("which top-down airborne studies have repeatedly found", "which top-down atmospheric studies have repeatedly found")
R("consistent with the ~11% landfill and wastewater contribution estimated by Pétron et al.,",
  "which Kille et al. noted is consistent with the ~11% landfill and wastewater contribution derived from Pétron et al.,")

# ---- SI ------------------------------------------------------------------------------------------
R("because the basin ethane-to-methane ratio fell roughly threefold between 2015 and 2021 and because",
  "because basin ethane emissions fell sharply between 2015 and 2021 while methane emissions changed little (section S5), and because", "si")
R("The biogenic-dominated median therefore holds across the Denver-relevant single endmembers considered; only the "
  "classification of the one or two most fossil-influenced flights moves with the choice.",
  "The biogenic-dominated median therefore holds for every regional reference value considered (0.081 to 0.110 mol "
  "mol-1); only the classification of the one or two most fossil-influenced flights moves with the choice. These values "
  "do not by themselves bound the effective endmember of the urban fossil mixture, so leaner values, including an "
  f"assumed lean basin ratio of 0.061 mol mol-1 (median {m_lean}%, with two flights above 50%), are treated as assumed "
  "scenarios rather than as an uncertainty range.", "si")
R("basin ratios of 0.061, 0.0813 and 0.102, and two kinds", "basin ratios of 0.0813 and 0.102 and an assumed lean basin ratio of 0.061, and two kinds", "si")
R("With the measured Denver composition the campaign median stays between 27 and 48% for every mixture and never reaches "
  "50% (inventory-weighted mixture, p = 0.67 with a basin ratio of 0.061: 31%, with the second 13 July flight at 51% and "
  "no other flight above the halfway line; distribution gas alone: 27%).",
  f"With the measured Denver composition the campaign median stays between {m_del} and {m_arc}% for every mixture with "
  "the basin ratios of 0.0813 and 0.102 (inventory-weighted mixture, p = 0.67, basin ratio 0.0813: 29%, no flight above "
  f"the halfway line; distribution gas alone: {m_del}%) and between {m_del} and {m_lean}% with the assumed 0.061 (31% at "
  "the inventory weighting, with the second 13 July flight at 51%); it never reaches 50%.", "si")
R("once p exceeded 0.06 to 0.10 (basin ratio 0.061), 0.36 to 0.51 (0.0813) or 0.52 to 0.67 (0.102), reaching 66 to 90% "
  "at the inventory weighting for a basin ratio of 0.061;",
  f"once p exceeded 0.06 to 0.10 (assumed basin ratio 0.061), 0.36 to 0.51 (0.0813) or {EC102} (0.102), reaching 66 to "
  "90% at the inventory weighting for the assumed basin ratio of 0.061;", "si")
R("for basin ratios of 0.061 (solid), 0.0813 (dashed) and 0.102 (dot-dashed)",
  "for an assumed lean basin ratio of 0.061 (solid) and basin ratios of 0.0813 (dashed) and 0.102 (dot-dashed)", "si")
R("For a basin ratio near 0.061, those floors exclude", "For the assumed lean basin ratio of 0.061, those floors exclude", "si")
R("(at the 0.061 calibration two flights cross 50%, section S3)", "(under the assumed lean ratio of 0.061 two flights cross 50%, section S3)", "si")
R("Its flux partition gives 2.1 t/h of ethane against 18.4 t/h of oil-and-gas methane, an 11% mass ratio, which is 0.061 "
  "mol mol−1. Our 2024 implied source ratio of 0.063 to 0.073 is similar, which is consistent with little additional "
  "decline after 2021, although the two estimates are not methodologically identical: the 2024 range depends on the "
  "assumed oil-and-gas share, the 2021 value on that study's source partition. A basin-wide 2021 emission ratio built "
  "from operator inventories and engine stack tests is higher, about 0.10 mol mol−1,9 because emissions from production "
  "tanks and combustion slip do not carry the wellhead composition; the contemporary basin range is therefore best stated "
  "as about 0.061 to 0.10, and the sensitivity analyses of section S3 use 0.061 as the lower basin case. A decline "
  "continuing at the 2015 to 2021 rate would put the 2024 source ratio near 0.034, below the ambient basin ratio of 0.036 "
  "measured here, which would be difficult to reconcile with the measured ambient ratio, since dilution by biogenic "
  "methane can only lower an ambient ratio below its source value. These contemporary constraints remain above the "
  "break-even value of 0.059 at which the campaign median fossil fraction would reach 50%, although the 0.061 lower case "
  "lies only slightly above it. Adopting an endmember in the 0.063 to 0.073 range would raise the urban median fossil "
  "fraction to roughly 40 to 47% while leaving the biogenic-dominated conclusion intact.",
  "It reports basin emissions of 2.6 t/h of ethane and 21.4 t/h of methane, 63% of the methane from oil and gas. Ethane "
  "against the oil-and-gas methane derived from that partition (about 13.5 t/h) is about 0.10 mol mol-1; against total "
  "methane it is about 0.065, but that ratio includes methane from livestock and other non-oil-and-gas sources and so is "
  "not a fossil endmember. A basin-wide 2021 emission ratio built from operator inventories and engine stack tests[[^9]] "
  "is similar, about 0.10 mol mol-1. Our 2024 implied source ratio of 0.063 to 0.073 is lower; it rests on the assumed "
  "oil-and-gas share near one half, and a share nearer one third would bring it to about 0.10, so the two years cannot be "
  "compared reliably. These values are regional reference points for the basin as a whole, from periods and footprints "
  "that differ from the urban sampling, not a bound on Denver's effective fossil endmember; the sensitivity analyses of "
  "section S3 therefore treat a lean basin ratio of 0.061 mol mol-1 as an assumed scenario only. All of these values lie "
  "above the break-even value of 0.059 at which the campaign median fossil fraction would reach 50%. Adopting an "
  "endmember in the 0.063 to 0.073 range would raise the urban median fossil fraction to roughly 40 to 47%.", "si")
R("Inside the analysis box the same construction gives 0.024 mol mol−1", f"Inside the analysis box the same construction gives {AMB} mol mol-1", "si")
R("fall within the Denver-relevant endmember range of 0.061 to 0.110 mol mol−1 adopted here.",
  "fall within the regional reference range of 0.081 to 0.110 mol mol-1.", "si")
R("(at the aircraft's ~50 m s−1, 2 s corresponds to ~100 m of track)", "(at the aircraft's ~70 m s-1, 2 s corresponds to ~140 m of track)", "si")
R("the basin median lies between 0.027 and 0.037 mol mol−1 for gates from 5 to 75 ppb, against the 0.036 reported there.",
  "the median of the positive basin slopes lies between 0.027 and 0.037 mol mol-1 for gates from 5 to 75 ppb (0.031 at "
  "the adopted 20 ppb gate; the screened median reported there is 0.036).", "si")
R("The median basin ratio over the 11 basin flights with a fit (without the reliability screen of section S5, which gives 0.036 at zero shift)",
  "The median basin ratio over the 11 basin flights with a positive slope (no correlation or leg-count screen; the "
  "screened median of section S5 is 0.036 at zero shift)", "si")
R("and the adopted analysis is left uncorrected.", "and the adopted analysis is left uncorrected (Figure S8).", "si")
R("over the 11 basin flights with a fit (without the reliability screen of section S5)",
  "over the 11 basin flights with a positive slope (without the correlation and leg-count screen of section S5)", "si")
I("The sweep supports the direction of the result, not the precision of that figure.",
  "The background itself is a further choice. For the aircraft, the rolling 5th percentile is computed over each "
  "flight's urban-leg samples joined end to end (basin legs for the basin ratio), so a 300-sample window near a leg "
  "boundary spans the adjacent leg and the background is the cleanest air sampled on neighboring urban legs. Computing "
  "it instead along the continuous 1 Hz flight record with the same window places the background within about 20 km of "
  "flight of each sample and removes part of the broad urban enhancement: the number of gated samples falls on every "
  "flight (from 234 to 95 on 13 July L2 and from 121 to 24 on 8 July), and the fossil fractions become 25, 32, 6, 28, 2, "
  "26 and 0% (3 July L1, 3 July L2, 8 July, 9 July, 10 July L1, 13 July L1 and 13 July L2), a median of 25% against "
  "29%, with no flight above 50%. The CH4:CO correlation also weakens (r from 0.85 to 0.17 on 8 July and from 0.91 to "
  "0.58 on 3 July L1), so two of the four flights passing the CH4:CO screen would fail r ≥ 0.7. We therefore retain the "
  "joined-leg background, which preserves the urban-scale enhancement on which the ratio method relies; the single "
  "flight that changes most, 13 July L2, is also the one whose interval is widest in Table 1.", "si")
R("the per-flight leg counts and leave-one-leg-out ranges are tabulated in Table S7 below.",
  "the per-flight leg counts and the CH4:CO leave-one-leg-out ranges are tabulated in Table S7 below. The CH4:CO2 fits "
  "of 9 July and 13 July L1 (r = 0.95 and 0.82) fail only this last criterion: their leave-one-leg-out ranges are -0.05 "
  "to 6.43 and 3.78 to 6.51 ppb/ppm against slopes of 6.24 and 4.23.", "si")
R("Here the comparison is made cell by cell .", "Here the comparison is made cell by cell (Table S8, Figure S11).", "si")
R("and one 10 July leg 19 km downwind of a site near Brighton (slope 0.07, r = 0.72);",
  "and one 10 July leg 19 km downwind of a site near Brighton (slope 0.07, r = 0.72), and links a 3 July (L1) leg 19 to "
  "20 km downwind to both the Tower Road landfill and two oil-and-gas sites (slopes 0.031 to 0.039, r = 0.70 to 0.73; "
  "crosses in Figure S12);", "si")
R("Fossil fraction versus the fraction of in-box samples arriving from the north-east (the Wattenberg/DJB gas-field sector).",
  "Fossil fraction versus the fraction of in-box boundary-layer samples (below 1.5 km above ground) with wind from 22.5 "
  "to 90° (north-north-east to east; the Wattenberg/DJB gas-field sector).", "si")
R("is consistent with north-easterly advection", "would be consistent with north-easterly advection", "si")
R("At leg level the association is absent.",
  "Across flights the association is weak and not significant (Figure S13). At leg level it is absent (Figure S14).", "si")
R("Basin oil and gas accounts for a median 22%", "In the STILT results (Figure S15), basin oil and gas accounts for a median 22%", "si")
R("(interquartile range 2.5 to 5.4)", "(interquartile range 2.6 to 5.4)", "si")
R("Of the flights that sampled the complex, 6 had a separating cross-wind.",
  "Of the flights that sampled the complex, 6 had a separating cross-wind. With a 5-min rolling background along the "
  "flight record in place of the flight-wide percentile, fewer near-complex samples count as plume (7 to 68 per flight) "
  "and 3 flights meet the separability criteria; the conclusion below holds under either background.", "si")
R("and at least 20 plume points were sampled.",
  "and at least 20 plume points were sampled. Plume points are samples within 4 km of the complex, 150 m to 1.5 km above "
  "ground and more than 20 ppb above the flight's 5th-percentile methane (a flight-wide background, used here only to "
  "locate plume samples near the complex).", "si")
# Table S4 (PSCo composition): two cells rounded half-up from the utility report
T("si", 3, 1, 8, "10.39", expect="10.38")
T("si", 3, 8, 4, "83.6", expect="83.5")

# ---- anchors must not straddle a hyperlink (a citation) ------------------------------------------
def sentinel_text(p):
    out = []
    for ch in p._p:
        if ch.tag == qn("w:r"): out.append("".join(t.text or "" for t in ch.iter(qn("w:t"))))
        elif ch.tag == qn("w:hyperlink"): out.append("\x00")
    return "".join(out)
for which, path in (("main", a.main_in), ("si", a.si_in)):
    d = docx.Document(path); txt = "\n".join(sentinel_text(p) for p in all_paragraphs(d))
    for anc, _, doc_, kind, _ in L.E:
        if doc_ == which and kind == "replace" and anc not in txt:
            sys.exit(f"{which}: anchor crosses a hyperlinked citation: {anc[:80]!r}")

L.apply(a.main_in, a.main_out, "main", "main text")
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- post-passes --------------------------------------------------------------------------------
def drop_hyperlink_after(d, text_before, link_text):
    """remove the hyperlinked citation `link_text` that follows `text_before` in one paragraph"""
    hits = 0
    for p in d.paragraphs:
        kids = list(p._p)
        for i, ch in enumerate(kids):
            if ch.tag == qn("w:hyperlink") and "".join(t.text or "" for t in ch.iter(qn("w:t"))) == link_text:
                prev = "".join(t.text or "" for k in kids[:i] for t in k.iter(qn("w:t")))
                if prev.rstrip().endswith(text_before):
                    p._p.remove(ch); hits += 1
    assert hits == 1, (text_before, link_text, hits)

FIGFONT = "Arial Unicode MS"
def tidy_layout(d, label):
    nf = 0
    for rf in d.element.body.iter(qn("w:rFonts")):
        if rf.get(qn("w:ascii")) == FIGFONT:
            rf.getparent().remove(rf); nf += 1
    nk = 0
    for p in d.paragraphs:
        if p._p.findall(".//" + qn("w:drawing")):
            p.paragraph_format.keep_with_next = True; nk += 1
    for tb in d.tables:                     # table captions sit just above their tables
        prev = tb._tbl.getprevious()
        while prev is not None and prev.tag == qn("w:p") and not "".join(t.text or "" for t in prev.iter(qn("w:t"))).strip():
            prev = prev.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            ppr = prev.find(qn("w:pPr"))
            if ppr is None: ppr = prev.makeelement(qn("w:pPr"), {}); prev.insert(0, ppr)
            if ppr.find(qn("w:keepNext")) is None: ppr.append(ppr.makeelement(qn("w:keepNext"), {}))
        for tr in tb._tbl.iter(qn("w:tr")):
            trpr = tr.find(qn("w:trPr"))
            if trpr is None: trpr = tr.makeelement(qn("w:trPr"), {}); tr.insert(0, trpr)
            if trpr.find(qn("w:cantSplit")) is None: trpr.append(trpr.makeelement(qn("w:cantSplit"), {}))
    print(f"{label}: {nf} runs reset to the body font; {nk} figures kept with their captions; table rows kept whole")

m = docx.Document(a.main_out)
drop_hyperlink_after(m, "assumed ethane-to-methane endmembers limit direct comparisons.", "2,5,25")
drop_hyperlink_after(m, "values below 0.059 mol mol−1 would yield one.", "18")
nb = 0
for p in m.paragraphs:
    for r in p.runs:
        if r.text.startswith("Section S7 (Figure S10) shows") and r.bold: r.bold = False; nb += 1
assert nb == 1, nb

# references: corrected entries (text inside the hyperlinked runs)
def wt_replace(d, old, new, n=1):
    k = 0
    for t in d.element.body.iter(qn("w:t")):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new); k += 1
            rpr = t.getparent().find(qn("w:rPr"))
            if rpr is None: rpr = t.getparent().makeelement(qn("w:rPr"), {}); t.getparent().insert(0, rpr)
            c = rpr.find(qn("w:color"))
            if c is None: c = rpr.makeelement(qn("w:color"), {}); rpr.append(c)
            c.set(qn("w:val"), "C00000")
    assert k == n, (old, k)
def fix_refs(d):
    wt_replace(d, " (11), 2013JD021272. https://doi.org/10/f58ddj", " (11), 6836–6852. https://doi.org/10.1002/2013JD021272", n=1 if d is m else 0)
    wt_replace(d, " (15). https://doi.org/10.1029/2009GL039825", " (15), L15810. https://doi.org/10.1029/2009GL039825")
    wt_replace(d, "Ngulat, M.; Santos", "Ngulat, M. C.; Santos")
    wt_replace(d, "Dickerson, R. R. R.;", "Dickerson, R. R.;")
    wt_replace(d, "ESS Open Archive 2025", "ESS Open Archive")
    wt_replace(d, " (0813). https://doi.org/10.22541", " 2025 (preprint). https://doi.org/10.22541")
fix_refs(m)

# main references renumbered by first citation ----------------------------------------------------
def sup(r):
    rp = r.find(qn("w:rPr")); va = rp.find(qn("w:vertAlign")) if rp is not None else None
    return va is not None and va.get(qn("w:val")) == "superscript"
REFP = re.compile(r"\((\d+)\)\s")
refs = [p for p in m.paragraphs if REFP.match(p.text.strip())]
listed = [int(REFP.match(p.text.strip()).group(1)) for p in refs]
assert listed == list(range(1, len(listed) + 1)), listed
ref_ids = {id(p._p) for p in refs}
groups = []                                   # [(list of w:t, text)] contiguous superscript groups
for p in m.paragraphs:
    if id(p._p) in ref_ids: break
    seen, cur = "", []
    for r in p._p.iter(qn("w:r")):
        ts = list(r.iter(qn("w:t"))); txt = "".join(t.text or "" for t in ts)
        if not txt: continue
        if sup(r):
            cur.append((ts, seen));
        else:
            if cur: groups.append(cur); cur = []
        seen += txt
    if cur: groups.append(cur)
cites = []
for g in groups:
    txt = "".join(t.text or "" for ts, _ in g for t in ts).strip(); before = g[0][1]
    if not re.fullmatch(r"\d+(?:[,–-]\d+)*", txt): continue
    if re.search(r"(GRA|km|δ|m)$", before.rstrip()): continue
    nums = []
    for part in txt.split(","):
        if re.search(r"[–-]", part):
            lo, hi = map(int, re.split(r"[–-]", part)); nums += list(range(lo, hi + 1))
        else: nums.append(int(part))
    cites.append((g, txt, nums))
order = []
for _, _, ns in cites:
    for n in ns:
        if n not in order: order.append(n)
assert sorted(order) == listed, (sorted(order), listed)
new = {o: k + 1 for k, o in enumerate(order)}
def fmt(ns):
    ns = sorted(set(new[n] for n in ns)); out, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1: j += 1
        out.append(f"{ns[i]}–{ns[j]}" if j - i >= 2 else ",".join(str(x) for x in ns[i:j + 1])); i = j + 1
    return ",".join(out)
changed = 0
for g, txt, ns in cites:
    s = fmt(ns)
    if s != txt.replace("-", "–"):
        allt = [t for ts, _ in g for t in ts]
        allt[0].text = s
        for t in allt[1:]: t.text = ""
        changed += 1
anchor = refs[0]._p.getprevious(); parent = refs[0]._p.getparent()
for p in refs: parent.remove(p._p)
for p in sorted(refs, key=lambda p: new[int(REFP.match(p.text.strip()).group(1))]):
    old = int(REFP.match(p.text.strip()).group(1))
    for t in p._p.iter(qn("w:t")):
        if (t.text or "").startswith(f"({old})"):
            t.text = f"({new[old]})" + t.text[len(f"({old})"):]; break
    anchor.addnext(p._p); anchor = p._p
moved = {o: n for o, n in new.items() if o != n}
print(f"main references renumbered by first citation ({changed} citation group(s) changed): " +
      ", ".join(f"{o}->{n}" for o, n in sorted(moved.items())))
tidy_layout(m, "main text")
m.save(a.main_out)

s = docx.Document(a.si_out)
fix_refs(s)
nh = 0
for tb in s.tables:
    hdr = tb.rows[0]
    if "E from CO" in "".join(c.text for c in hdr.cells):
        for c in hdr.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) == "C00000":
                        r.font.color.rgb = docx.shared.RGBColor(0xFF, 0xFF, 0xFF); nh += 1
print(f"SI: {nh} Table S7 header run(s) set to white")
tidy_layout(s, "SI")
s.save(a.si_out)

# ---- regenerated figures -------------------------------------------------------------------------
FIG = OUT / "figures"
for media, png in (("image5.png", "FigS3b_two_endmember.png"), ("image2.png", "FigS8_basin_ethane_trend.png"),
                   ("image3.png", "FigS13_spatial_match.png"), ("image12.png", "SI_wind_fossil.png")):
    L.swap_pictures(a.si_out, [(media, str(FIG / png))], "SI")

# ---- final checks --------------------------------------------------------------------------------
for path, label in ((a.main_out, "main text"), (a.si_out, "SI")):
    t = "\n".join(runs_text(p) for p in all_paragraphs(docx.Document(path)))
    for bad in ("Denver-relevant", "flux partition", "18.4 t/h", "0.024 mol", "0.67 (0.102)", "6-67%", "27 to 48%", "44 to 79"):
        assert bad not in t, f"{label}: {bad!r} remains"
print("checks passed")
