"""101_aloft_background_edits.py -- adds the aloft-background sensitivity (script 100; S. Baidar,
16 Oct 2026) to the manuscript (section 3.2) and SI (section S4). Works on Priyanka's own
final_aircraft files. Numbers are read from background_aloft_sensitivity.csv and asserted.
New text is red.

usage: python3 scripts/101_aloft_background_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, argparse, statistics as st
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
ap.add_argument("--outputs", required=True)
a = ap.parse_args()
R = {r["flight"]: r for r in csv.DictReader(open(Path(a.outputs) / "background_aloft_sensitivity.csv"))}
order = ["20240703_R0_L1", "20240703_R0_L2", "20240705_R0_L1", "20240708_R0_L1", "20240709_R0", "20240710_R0_L1", "20240713_R0_L1", "20240713_R0_L2"]
lab = ["3 July L1", "3 July L2", "5 July L1", "8 July L1", "9 July", "10 July L1", "13 July L1", "13 July L2"]
assert R["20240712_R0_L1"]["fossil_pct_aloft"] in ("", "NA") and int(R["20240712_R0_L1"]["n_aloft"]) == 0
ff = [int(float(R[f]["fossil_pct_aloft"])) for f in order]; f5 = [int(float(R[f]["fossil_pct_aloft_p05"])) for f in order]
fp = [int(float(R[f]["fossil_pct_paper"])) for f in order]
med_a, med_5, med_p = round(st.median(ff)), round(st.median(f5)), st.median(fp)
assert (ff, med_a, med_5) == ([22, 0, 15, 28, 35, 12, 42, 46], 25, 29) and max(ff + f5) < 50, (ff, med_a, med_5)
agl = [float(R[f]["aloft_agl_median_m"]) for f in order]
d32 = float(R["20240703_R0_L2"]["bg_ch4_aloft"]) - float(R["20240703_R0_L2"]["urban_ch4_p05"])
assert round(d32) == 18 and (round(min(agl) / 1000, 1), round(max(agl) / 1000, 1)) == (2.2, 2.5)
CO = ["20240703_R0_L1", "20240703_R0_L2", "20240709_R0", "20240713_R0_L1"]
E = [float(R[f]["E_co_aloft"]) for f in CO]; rco = {f: float(R[f]["r_co_aloft"]) for f in CO}
assert (round(min(E), 1), round(max(E), 1)) == (4.2, 8.3) and rco["20240713_R0_L1"] == 0.68 and all(rco[f] >= 0.7 for f in CO[:3])
outside = [l for f, l in zip(order, lab) if R[f]["aloft_from"] != "over the box"]; assert outside == ["13 July L2"]

MAIN_ADD = (f" Taking instead a single background for each flight from the air sampled above the boundary layer raises the campaign median from 21% to {med_a}% "
            f"({med_5}% with the 5th percentile of that air), and no flight reaches 50% (section S4).")
SI_PARA = (
    "Because a rolling background could also remove part of a spatially extended urban enhancement, such as the diffuse emissions of an area source, "
    "we tested a background taken from the air above the boundary layer. For each flight we took the constant-altitude transect samples at least 100 m above the "
    "boundary-layer height measured by the airborne lidar at that sample and at least 1 km above ground (median heights of "
    f"{min(agl) / 1000:.1f} to {max(agl) / 1000:.1f} km above ground; over the analysis box except on 13 July L2, whose aloft samples lie outside it) "
    "and used their median CH4, C2H6 and CO as one background for all of the flight's urban-transect samples, followed by the same 20 ppb gate and within-transect fit. "
    "The 12 July L1 flight had no transect above its boundary layer and is omitted. "
    f"The fossil fractions become {', '.join(str(x) for x in ff[:-1])} and {ff[-1]}% for {', '.join(lab[:-1])} and {lab[-1]}, respectively; "
    f"the median of these eight flights rises from 21% to {med_a}% ({med_5}% with the 5th percentile of the aloft air instead of its median), "
    "and no flight reaches 50%. The lowest fractions change little (3 July L2 stays at 0% and 10 July L1 moves from 8 to 12%). "
    "Because each transect has its own intercept in the within-transect fit, an enhancement that raises a whole transect does not change the slope under either background; "
    "the aloft background changes the fit mainly through which samples pass the 20 ppb gate and through structure on scales longer than the 300-sample window. "
    f"The air above the boundary layer is a different air mass, however: on 3 July L2 its methane exceeded the cleanest urban-transect air by {round(d32)} ppb, "
    "and a constant background from aloft counts methane advected into the boundary layer from the basin as urban enhancement, which raises the inferred fossil fraction. "
    f"We therefore treat the two backgrounds as bracketing the campaign median (21 to {med_5}%) rather than adopting the aloft one. "
    f"With the aloft background the four CH4:CO emission estimates become {min(E):.1f} to {max(E):.1f} t CH4/h, "
    f"although 13 July L1 then falls just below the 0.7 correlation criterion (r = {rco['20240713_R0_L1']:.2f}).")

m = docx.Document(a.main_in)
O.do_replace(m, "which retains the urban-scale enhancement (section S4).", "which retains the urban-scale enhancement (section S4)." + MAIN_ADD, 1, "main")
L.subscript_chem(L.all_paragraphs(m)); m.save(a.main_out)

s = docx.Document(a.si_in)
anchor = [p for p in s.paragraphs if O.vis_text(p).startswith("The alternative background leaves the CH4:CO correlations")]
assert len(anchor) == 1
L.insert_after(anchor[0], SI_PARA, s)
L.subscript_chem(L.all_paragraphs(s)); s.save(a.si_out)

tm = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.main_out)))
ts = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.si_out)))
assert tm.count("from the air sampled above the boundary layer") == 1 and ts.count("bracketing the campaign median") == 1
print("aloft background added: main 3.2, SI S4")
