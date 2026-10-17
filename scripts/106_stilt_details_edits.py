"""106_stilt_details_edits.py -- states the HRRR-STILT configuration of script 64 in full
(SI section S9, one new paragraph after the footprint-methods paragraph; "3-km" in main 2.4).
Numbers (receptors, transects, flights, receptor heights) are read from the script-64 outputs
and asserted against the existing text. Works on the output of script 105. New text is red.

usage: python3 scripts/106_stilt_details_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, argparse, re
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx, docx.oxml.ns
import docx_edit_lib as L
from docx_edit_lib import I, R, all_paragraphs
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs)

RR = list(csv.DictReader(open(OUT / "stilt_exposure_receptors.csv")))
LG = list(csv.DictReader(open(OUT / "stilt_exposure_legs.csv")))
n_rec, n_leg, n_fl = len(RR), len(LG), len(set(r["flight"] for r in LG))
assert (n_rec, n_leg, n_fl) == (112, 28, 7), (n_rec, n_leg, n_fl)
z = [float(r["zagl"]) for r in RR]; assert 150 <= min(z) and max(z) <= 2500
assert all(float(r["foot_sum"]) > 0 for r in RR)
per_leg = {}
for r in RR: per_leg.setdefault((r["flight"], r["leg_id"]), 0); per_leg[(r["flight"], r["leg_id"])] += 1
assert max(per_leg.values()) == 4 and n_leg == len(per_leg)
# settings fixed in scripts/64_stilt_djb_exposure.R (write_setup / write_control / plume_dilution / DOMAIN)
src = open(Path(__file__).resolve().parent / "64_stilt_djb_exposure.R").read()
for needle in ('NUMPAR  <- as.integer(Sys.getenv("METHANE_STILT_NUMPAR", "500"))', 'N_HOURS <- as.numeric(Sys.getenv("METHANE_STILT_HOURS", "-10"))',
               'N_REC   <- as.integer(Sys.getenv("METHANE_STILT_NREC", "4"))', "DELT = 1", "KBLT = 5", "KMIXD = 3", "KMIX0 = 150",
               "VEGHT = 0.5", "TKERD = 0.18", "DOMAIN  <- c(ymn = 37.5, xmn = -108.5, ymx = 42.5, xmx = -101.5)", "zagl = max(5, round(q$ALTAGL))",
               "by = 0.02)", "basin_rows <- gy >= BASIN_LAT", "30 min back", "plume_dilution <- function(p, r_zagl, veght = 0.5)"):
    assert needle in src, needle

SI_PARA = (
    "The configuration follows the STILT v2 defaults.[[^21]] The four receptors per transect are gated samples spread evenly along the gated part of the transect, "
    "each released at the aircraft's position and height above ground at that sample (receptor heights "
    f"{min(z):.0f} to {max(z):.0f} m), with the release time rounded to the minute. The particles were followed at a "
    "1-min time step, with the Hanna turbulence parameterization, the mixed-layer depth diagnosed from a modified Richardson-number profile with a "
    "150 m minimum, and the surface sensitivity (footprint) accumulated while a particle is in the lower half of the mixed layer; the near-field "
    "plume-dilution correction of STILT v2 was applied to the particle output. The meteorology is the NOAA ARL archive of HRRR analyses (3 km, hourly, "
    "in 6-h files), cut to 37.5 to 42.5°N and 108.5 to 101.5°W before the runs. The footprints (ppm per μmol m-2 s-1, summed over the 10 h) were "
    "convolved on each inventory's own grid: the 0.1° EPA gridded GHGI (2020 Express Extension), and the GRA2PES v2.0beta 4-km cells assigned to a 0.02° "
    "grid by nearest cell. In the GHGI, basin-type oil and gas comprises the petroleum and natural-gas exploration, production, processing and "
    "transmission and storage sectors; distribution and post-meter the natural-gas distribution and post-meter sectors; waste the municipal and "
    "industrial landfills, composting and domestic and industrial wastewater; livestock enteric fermentation and manure management; and all other "
    "sectors the remainder. The share from the basin north of 40.05°N counts only basin-type oil-and-gas flux in grid cells at or north of that "
    "latitude. The HRRR transport direction at a receptor is the bearing from the receptor to the center of its particle cloud 30 min back, averaged "
    f"over the transect's receptors with the speed as weight. Every one of the {n_rec} receptors has a non-zero footprint. The predicted enhancement "
    "carries no background of its own, whereas the observed enhancement is relative to the rolling background of section 3.2, which is one reason "
    "the observed-to-predicted ratios below are not read as emission scalings. The run files, particle outputs and convolution are produced by "
    "scripts/64_stilt_djb_exposure.R in the code repository.")

if any("The configuration follows the STILT v2 defaults" in p.text for p in docx.Document(a.si_in).paragraphs):
    sys.exit("the STILT details are already in this SI; run on the output of script 105")
I("and 0.110 for distribution and post-meter gas.", SI_PARA, "si")
R("The High-Resolution Rapid Refresh (HRRR) model drives the Stochastic Time-Inverted Lagrangian Transport (STILT) footprints (section S9).",
  "The 3-km High-Resolution Rapid Refresh (HRRR) model drives the Stochastic Time-Inverted Lagrangian Transport (STILT) footprints (section S9).", "main")
L.apply(a.main_in, a.main_out, "main", "main text")
L.apply(a.si_in, a.si_out, "si", "SI")
# the SI separates blocks with empty paragraphs: put one before the new paragraph
import copy as _copy
_si = docx.Document(a.si_out)
for _p in _si.paragraphs:
    if _p.text.startswith("The configuration follows the STILT v2 defaults"):
        sp = _copy.deepcopy(_p._p)
        for ch in list(sp):
            if ch.tag != docx.oxml.ns.qn("w:pPr"): sp.remove(ch)
        _p._p.addprevious(sp); break
_si.save(a.si_out)

ts = "\n".join(O.vis_text(p) for p in all_paragraphs(docx.Document(a.si_out)))
tm = "\n".join(O.vis_text(p) for p in all_paragraphs(docx.Document(a.main_out)))
assert ts.count("The configuration follows the STILT v2 defaults") == 1 and tm.count("The 3-km High-Resolution Rapid Refresh") == 1
print("STILT configuration added to SI S9; 3-km HRRR in main 2.4")
