"""98_final_vet_fixes.py -- last consistency pass (14 Oct 2026), applied to the output of script 97.

Main text
  * 2.4: the ground-vs-airborne mixing-height comparison is over the 13 flights of Table S1 with
    both values; adds what the lidar team reported for the two Table 1 flights without an
    airborne retrieval (12 July: lidar not operating; 3 July L2: retrieval not archived because
    the boundary layer was poorly defined, best estimate puts the low-level transects inside it).
  * "leg" -> "transect" where the text means the constant-altitude transects (3.3).
  * 3.3: the basin ratios are those of the four flights passing the screen.
  * 4.2: 3 July L2 wind "about 300°" (Table S1, from the transect samples, gives 298°; the
    wind-sector analysis, from all in-box boundary-layer samples, gives 299°).
  * 3.2: missing full stop at the end of the background paragraph; US spelling.
SI
  * S1.1.1 and Table S1 caption: same mixing-height wording; Table S1 shows three flights without
    an airborne retrieval, not five.
  * S3: US spelling ("labeled").
New text is red; spacing and punctuation fixes keep the formatting of the text they touch.

usage: python3 scripts/98_final_vet_fixes.py <main97.docx> <si97.docx> <main_out> <si_out>
"""
import sys, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
a = ap.parse_args()
W_ = lambda t: qn("w:" + t)

def t_fix(d, old, new, count, label):
    n = 0
    for p in L.all_paragraphs(d):
        for t in p._p.iter(W_("t")):
            if t.text and old in t.text:
                n += t.text.count(old); t.text = t.text.replace(old, new)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    if n != count: sys.exit(f"{label}: {old!r} fixed {n}x, expected {count}x")

LIDAR = ("For 3 July L2 the airborne retrieval was not archived because the boundary layer was poorly defined that afternoon; "
         "the lidar team's best estimate places that flight's low-level transects inside the boundary layer. The airborne lidar did not operate on 12 July.")

# ================================ main ================================
m = docx.Document(a.main_in); LM = "main text"
O.do_replace(m, "on 10 of the 13 flights covered by both", "on 10 of the 13 flights in Table S1 covered by both", 1, LM)
O.do_replace(m, "whereas 7 of the 10 transects flown above 1.5 km lay above it.",
             "whereas 7 of the 10 transects flown above 1.5 km lay above it. " + LIDAR, 1, LM)
O.do_replace(m, "Pooling all legs in a single regression", "Pooling all transects in a single regression", 1, LM)
O.do_replace(m, "we also fitted the pooled leg data using", "we also fitted the pooled transect data using", 1, LM)
O.do_replace(m, "Among legs with strong ethane-methane correlations", "Among transects with strong ethane-methane correlations", 1, LM)
O.do_replace(m, "the DJB basin legs yielded ambient", "the DJB basin transects yielded ambient", 1, LM)
O.do_replace(m, "(the four best-sampled flights and all flights passing the screen)", "(the four flights passing the screen)", 1, LM)
O.do_replace(m, "(10.4 m/s from 299°)", "(10.4 m/s from about 300°)", 1, LM)
t_fix(m, "within a few hundred metres shifts", "within a few hundred meters shifts", 1, LM)
p76 = O.find_para(m, "even though the regional background varies slowly", LM)
last = [t for t in p76._p.iter(W_("t")) if t.text][-1]
assert last.text.endswith("varies slowly"), last.text[-40:]
last.text += "."
m.save(a.main_out)

# ================================ SI ================================
s = docx.Document(a.si_in); LS = "SI"
O.do_replace(s, "on 10 of the 13 flights covered by both", "on 10 of the 13 flights in Table S1 covered by both", 1, LS)
O.do_replace(s, "Five flights have no airborne retrieval (28 June, 3 July L2, 11 July L2, and 12 July L1 and L2).",
             "Five flights have no airborne retrieval (28 June, 3 July L2, 11 July L2, and 12 July L1 and L2). " + LIDAR, 1, LS)
O.do_replace(s, "a dash marks the five flights without an airborne retrieval, for which the ground value is used.",
             "a dash marks the three flights here without an airborne retrieval (3 July L2 and both 12 July flights), for which the ground value is used.", 1, LS)
t_fix(s, "labelled", "labeled", 2, LS)
s.save(a.si_out)

# ---- checks ----------------------------------------------------------------------------------
tm = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.main_out)))
ts = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.si_out)))
for b in ("Pooling all legs", "pooled leg data", "Among legs", "basin legs", "metres", "from 299°", "varies slowly\n"):
    assert b not in tm, f"main: {b!r} remains"
for b in ("labelled", "dash marks the five flights"):
    assert b not in ts, f"SI: {b!r} remains"
print("checks passed")
