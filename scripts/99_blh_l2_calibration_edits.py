"""99_blh_l2_calibration_edits.py -- 3 July L2 airborne mixing height, calibration wording from the file
headers, CDPHE QA/QC (14 Oct 2026; applied to the output of script 98).

  * The lidar team supplied the airborne boundary-layer file for 3 July L2
    (AMMBEC-AirborneDL-BLH_TwinOtter_20240703_R1_L2.ict; not in the public archive because the
    boundary layer was poorly defined that afternoon). Scripts 13, 22, 33, 59, 90, 91 and 21 were
    re-run with it: 18 of 22 flights now have an airborne retrieval; 3 July L2 gets 1502 m in the
    box (ground lidar 3352 m); its QC-filtered fossil fraction becomes 9% (full 0%). Main 2.4, 3.3,
    3.4, limitations, data availability; SI S1.1, Table S1 (+ caption), Table S2, S3 text and Table S3.
  * 2.1: the calibration wording follows the ICARTT headers (pre- and post-campaign calibration on
    the eight files that say so; pre-campaign laboratory calibration on the other 14).
  * SI S2.1: CDPHE applied no QA/QC to the surveys (replaces the placeholder).
Numbers are read from the outputs and asserted. New text is red.

usage: python3 scripts/99_blh_l2_calibration_edits.py <main98.docx> <si98.docx> <main_out> <si_out>
         --outputs <OUT_DIR>
"""
import sys, csv, re, argparse, statistics as st
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
from docx.text.run import Run
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs)
rd = lambda f: list(csv.DictReader(open(OUT / f)))
W_ = lambda t: qn("w:" + t)
f2 = lambda v, d=2: "" if v in ("NA", "", None) else f"{float(v):.{d}f}"

# ---- numbers -------------------------------------------------------------------------------------
T1 = rd("table1.csv"); FL = [r["Flight"] for r in T1 if r["Fossil_pct_CI"] != "NA"]; assert len(FL) == 9
FS = {r["flight"]: r for r in rd("flight_summary.csv")}; assert len(FS) == 16
FB = {r["flight"]: r for r in rd("airborne_blh_flights.csv")}; assert len(FB) == 22
n_air = sum(r["airborne_file"] == "TRUE" for r in FB.values()); no_air = [f for f, r in FB.items() if r["airborne_file"] != "TRUE"]
assert n_air == 18 and no_air == ["20240628_R0", "20240711_R0_L2", "20240712_R0_L1", "20240712_R0_L2"], no_air
assert FS["20240703_R0_L2"]["blh_src"].startswith("airborne") and FS["20240703_R0_L2"]["blh_m"] == "1502" and FS["20240703_R0_L2"]["blh_ground_m"] == "3352"
both = [f for f in FS if FB[f]["airborne_file"] == "TRUE" and "DSRC" in FB[f]["ground_src"]]
air_of = lambda f: float(FB[f]["blh_airborne_inbox_m"] or "nan") if FB[f]["blh_airborne_inbox_m"] not in ("", "NA") else float(FB[f]["blh_airborne_flight_m"])
diff = [float(FB[f]["blh_ground_m"]) - air_of(f) for f in both]
n_both, n_higher, dmax, dmed = len(both), sum(x > 0 for x in diff), max(diff), st.median(diff)
assert (n_both, n_higher, round(dmax / 1000, 1), round(dmed / 1000, 2)) == (14, 11, 2.9, 0.84), (n_both, n_higher, dmax, dmed)
below = sorted(f for f, x in zip(both, diff) if x < 0); assert below == ["20240710_R0_L1", "20240713_R0_L1", "20240713_R0_L2"], below
air9 = [f for f in FL if FS[f]["blh_src"].startswith("airborne")]; assert len(air9) == 8 and "20240712_R0_L1" not in air9
span9 = [int(FS[f]["blh_m"]) for f in FL]; span8 = [int(FS[f]["blh_m"]) for f in air9]
assert (min(span9), max(span9), min(span8), max(span8), FS["20240712_R0_L1"]["blh_m"]) == (630, 2752, 630, 2661, "2752")
TB = [r for r in rd("transects_vs_blh.csv") if r["flight"] in FL and r["blh_airborne_m"] not in ("", "NA")]
lo = [r for r in TB if int(r["agl_m"]) < 1000]; hi = [r for r in TB if int(r["agl_m"]) >= 1500]
n_lo_in, n_lo = sum(r["inside_bl"] == "TRUE" for r in lo), len(lo); n_hi_out, n_hi = sum(r["inside_bl"] == "FALSE" for r in hi), len(hi)
assert (n_lo_in, n_lo, n_hi_out, n_hi) == (38, 38, 8, 11), (n_lo_in, n_lo, n_hi_out, n_hi)
g_pct = round(100 * sum(int(r["gated_inside_bl"]) for r in TB) / sum(int(r["gated_with_blh"]) for r in TB)); assert g_pct == 96
QC = {r["flight"]: r for r in rd("qc_robustness.csv")}
qc_a = [int(float(QC[f]["fossil_QC"])) for f in FL]; qc_g = [int(float(QC[f]["fossil_QC_ground"])) for f in FL]; qc_f = [int(float(QC[f]["fossil_full"])) for f in FL]
assert qc_a == [24, 9, 8, 0, 34, 7, 21, 36, 34] and qc_g == [24, 0, 8, 0, 34, 9, 21, 27, 43] and qc_f == [24, 0, 8, 17, 34, 8, 21, 35, 34], (qc_a, qc_g, qc_f)
assert st.median(qc_a) == 21 and st.median(qc_g) == 21 and st.median(qc_f) == 21
sh = [float(QC[f]["share_above_blh"]) for f in air9]; assert (round(100 * min(sh)), round(100 * max(sh))) == (2, 48), sh
assert float(QC["20240703_R0_L2"]["share_above_blh"]) == 0.259
CS = rd("denver_curtain_screen.csv")
cnt = {k: sum(r[k] == "TRUE" for r in CS) for k in ("city_upwind", "width_ok", "vertical_ok", "wind_ok")}
assert cnt == {"city_upwind": 11, "width_ok": 9, "vertical_ok": 5, "wind_ok": 14} and max(int(r["conditions_met"]) for r in CS) == 2 and len(CS) == 48
ENC, seen = [], set()
for r in rd("v2_pointsource_encounters.csv"):
    k = (r["flight"], r["leg_id"], r["hotspot_id"])
    if r["hotspot_in_fetch"] == "TRUE" and k not in seen: seen.add(k); ENC.append(r)
assert len(ENC) == 15 and len({r["flight"] for r in ENC}) == 7 and "20240703_R0_L2" not in {r["flight"] for r in ENC}
print(f"airborne files {n_air}; ground > airborne on {n_higher} of {n_both} (max {dmax:.0f}, median {dmed:.0f} m); transects {n_lo_in}/{n_lo}, {n_hi_out}/{n_hi}, gated {g_pct}%")

def span_fix(d, old, new, label, red=True):
    """Replace `old` in the visible text of one paragraph even when it spans several w:t elements
    (plain runs and tracked insertions alike): the first element overlapped takes `new`."""
    hits = 0
    for p in L.all_paragraphs(d):
        ts = list(p._p.iter(W_("t"))); txt = "".join(t.text or "" for t in ts)
        i = txt.find(old)
        if i < 0: continue
        hits += 1; j = i + len(old); pos = 0; first = True
        for t in ts:
            s = t.text or ""; e = pos + len(s)
            if e > i and pos < j:
                a0, a1 = max(i, pos) - pos, min(j, e) - pos
                t.text = s[:a0] + (new if first else "") + s[a1:]; first = False
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                if red and t.text:
                    r = t.getparent(); Run(r, p).font.color.rgb = L.RED
            pos = e
    if hits != 1: sys.exit(f"{label}: {old[:70]!r} found {hits}x, expected 1x")

L2 = ("The 3 July L2 retrieval is not in the public archive, because the boundary layer was poorly defined that afternoon; "
      "the lidar team provided it for this analysis, and it is used like the others. The airborne lidar did not operate on 12 July.")

# ================================ main ================================
m = docx.Document(a.main_in); LM = "main text"
span_fix(m, "Both analyzers were calibrated before, during, and after the field study", "Both analyzers were calibrated before and after the field study", LM)
span_fix(m, "The archived files (revision R0) are final data based on the calibrations conducted before, during, and after the campaign.",
         "The archived files (revision R0, labeled final data) state in their headers that the data of eight flights (5, 8, 9 and 10 July and 11 July L1) are based on the pre- and post-campaign calibration results, "
         "and that those of the other 14 flights, including 3 July, 12 July L1 and both 13 July flights, are based on the pre-campaign laboratory calibration.", LM)
O.do_replace(m, "17 of the 22 flights).", "18 of the 22 flights).", 1, LM)
O.do_replace(m, "supplies the mixing height for four of the five flights without an airborne retrieval", "supplies the mixing height for three of the four flights without an airborne retrieval", 1, LM)
O.do_replace(m, "It spans 0.63 to 2.66 km on the seven of the nine flights with estimable urban slopes that have an airborne retrieval; for the other two, 3 July L2 and 12 July L1, the ground lidar gives 3.35 and 2.75 km (Table S1).",
             "It spans 0.63 to 2.66 km on the eight of the nine flights with estimable urban slopes that have an airborne retrieval; for the ninth, 12 July L1, the ground lidar gives 2.75 km (Table S1).", 1, LM)
O.do_replace(m, "On those seven flights all 34 urban transects flown below 1 km above ground lay inside the boundary layer measured along the track, as did 97% of the samples entering the ethane fits, whereas 7 of the 10 transects flown above 1.5 km lay above it.",
             f"On those eight flights all {n_lo} urban transects flown below 1 km above ground lay inside the boundary layer measured along the track, as did {g_pct}% of the samples entering the ethane fits, whereas {n_hi_out} of the {n_hi} transects flown above 1.5 km lay above it.", 1, LM)
O.do_replace(m, "For 3 July L2 the airborne retrieval was not archived because the boundary layer was poorly defined that afternoon; the lidar team's best estimate places that flight's low-level transects inside the boundary layer. The airborne lidar did not operate on 12 July.", L2, 1, LM)
O.do_replace(m, "exceeds the airborne value on 10 of the 13 flights in Table S1 covered by both, by up to 2.9 km (median difference 0.75 km)",
             f"exceeds the airborne value on {n_higher} of the {n_both} flights in Table S1 covered by both, by up to {dmax / 1000:.1f} km (median difference {dmed / 1000:.2f} km)", 1, LM)
O.do_replace(m, "leaves the campaign median unchanged at 21% and changes only one flight by more than 1 percentage point (section S3).",
             "leaves the campaign median unchanged at 21% and changes two flights by more than 1 percentage point (section S3).", 1, LM)
O.do_replace(m, "only 8 July L1 changes by more than 1 percentage point (17 to 0%; Table S3).",
             "only 3 July L2 (0 to 9%) and 8 July L1 (17 to 0%) change by more than 1 percentage point (Table S3).", 1, LM)
O.do_replace(m, "airborne Doppler lidar on 17 of the 22 flights; for the other five, including 3 July L2 and 12 July L1 in Table 1, it comes from",
             "airborne Doppler lidar on 18 of the 22 flights; for the other four, including 12 July L1 in Table 1, it comes from", 1, LM)
O.do_replace(m, "and airborne Doppler-lidar boundary-layer heights are available from the NOAA Chemical Sciences Laboratory AMMBEC 2024 data archive",
             "and airborne Doppler-lidar boundary-layer heights are available from the NOAA Chemical Sciences Laboratory AMMBEC 2024 data archive (the 3 July L2 boundary-layer file, which is not posted there, was provided by the lidar team)", 1, LM)
m.save(a.main_out)

# ================================ SI ================================
s = docx.Document(a.si_in); LS = "SI"
O.do_replace(s, "[Authors to complete with CDPHE: calibration standards and frequency for the G2204.]",
             "CDPHE applied no quality-assurance or quality-control procedures to these survey data; the processing described here is the only quality control applied, which is a further reason the surveys are used only as context.", 1, LS)
O.do_replace(s, "Five flights have no airborne retrieval (28 June, 3 July L2, 11 July L2, and 12 July L1 and L2).",
             "Four flights have no airborne retrieval (28 June, 11 July L2, and 12 July L1 and L2).", 1, LS)
O.do_replace(s, "For 3 July L2 the airborne retrieval was not archived because the boundary layer was poorly defined that afternoon; the lidar team's best estimate places that flight's low-level transects inside the boundary layer. The airborne lidar did not operate on 12 July.", L2, 1, LS)
O.do_replace(s, "spans 0.63 to 3.35 km.", "spans 0.63 to 2.75 km.", 1, LS)
O.do_replace(s, "exceeds the airborne value on 10 of the 13 flights in Table S1 covered by both, by up to 2.9 km (median difference over the 13 flights 0.75 km)",
             f"exceeds the airborne value on {n_higher} of the {n_both} flights in Table S1 covered by both, by up to {dmax / 1000:.1f} km (median difference over the {n_both} flights {dmed / 1000:.2f} km)", 1, LS)
O.do_replace(s, "a dash marks the three flights here without an airborne retrieval (3 July L2 and both 12 July flights), for which the ground value is used.",
             "a dash marks the two flights here without an airborne retrieval (both 12 July flights), for which the ground value is used.", 1, LS)
O.do_replace(s, "(the ground lidar for 3 July L2 and 12 July L1, which have no airborne retrieval)", "(the ground lidar for 12 July L1, which has no airborne retrieval)", 1, LS)
O.do_replace(s, "It removes 2 to 48% of the urban-transect samples on the seven flights with airborne retrievals, mostly on the transects flown above the mixed layer, but few of the samples that enter the ethane fits: 97% of those lie inside the boundary layer (section S1.1).",
             f"It removes 2 to 48% of the urban-transect samples on the eight flights with airborne retrievals, mostly on the transects flown above the mixed layer, but few of the samples that enter the ethane fits: {g_pct}% of those lie inside the boundary layer (section S1.1).", 1, LS)
O.do_replace(s, "Eight of the nine flights change by at most 1 percentage point (Table S3); 8 July L1 falls from 17 to 0% under either mixing height,",
             "Seven of the nine flights change by at most 1 percentage point (Table S3). 3 July L2 rises from 0 to 9% because the filter removes its transect flown at 2.4 km, above the 1.5 km boundary layer measured that afternoon (26% of its urban-transect samples), and the remaining transects give a small positive slope; 8 July L1 falls from 17 to 0% under either mixing height,", 1, LS)
O.do_replace(s, "With the ground-lidar mixing height instead, the median is also unchanged, but the two 13 July flights move to 27 and 43%:",
             "With the ground-lidar mixing height instead, the median is also unchanged and 3 July L2 stays at 0% (the ground value of 3.35 km removes none of its transects), but the two 13 July flights move to 27 and 43%:", 1, LS)
O.do_replace(s, "and only one flight changed by more than 1 percentage point (Table S3).", "and only two flights changed by more than 1 percentage point (Table S3).", 1, LS)
O.do_replace(s, "(the ground lidar for 3 July L2 and 12 July L1)", "(the ground lidar for 12 July L1)", 1, LS)

# ---- tables ----
lab = lambda f: (lambda g: f"{g.group(1)}-{g.group(2)}-{g.group(3)}" + (f" ({g.group(4)})" if g.group(4) else ""))(re.match(r"(\d{4})(\d\d)(\d\d)_R0_?(L\d)?", f))
def blh_cell(f):
    r = FS[f]; air = r["blh_src"].startswith("airborne")
    return (r["blh_m"] if air else "-") + " / " + (r["blh_ground_m"] if r["blh_ground_m"] not in ("", "NA") else "-")
S1rows = []
for r in T1:
    f = r["Flight"]; x = FS[f]
    S1rows.append([lab(f), x["flight_utc"], x["inbox_utc"], x["n_box_transects"], f"{int(x['inbox_transect_s']):,}", x["agl_range"],
                   x["bl_wind"].replace(" m/s from ", " m/s, "), x["above_wind"].replace(" m/s from ", " m/s, "), blh_cell(f), x["n_alt_levels"], x["encircles"]])
c1 = O.do_table(s, 0, S1rows, None, LS, widths=[1.2, 1.0, 1.0, 0.85, 0.8, 0.9, 1.2, 1.2, 0.95, 0.6, 0.75])
rnd = lambda x: int(abs(x) + 0.5) * (1 if x >= 0 else -1)
S2rows = [[lab(r["flight"]), f"{'N-S' if r['axis'] == 'NS' else 'E-W'}, {rnd(float(r['offset_km_from_box_center']))} km", f"{r['n_legs']} / {r['n_alt_levels']}",
           f"{r['lowest_agl_m']}-{r['highest_agl_m']}", r["blh_m"], f"{float(r['width_frac']):.2f}", f"{float(r['U_perp_ms']):.1f} / {float(r['constancy']):.2f}", f"{r['conditions_met']} of 4"]
          for r in sorted(CS, key=lambda r: (-int(r["conditions_met"]), r["flight"]))]
c2 = O.do_table(s, 1, S2rows, None, LS)
S3rows = [[lab(f), qf, qa, abs(qa - qf), qg] for f, qf, qa, qg in zip(FL, qc_f, qc_a, qc_g)]
c3 = O.do_table(s, 2, S3rows, None, LS)
S10rows = [[lab(r["flight"]), r["leg_id"], r["hotspot_id"], f2(r["dist_km"], 1), r["agl_m"], r["blh_m"], r["dch4_max_ppb"], f2(r["Q_t_hr"]),
            f2(r["fossil_frac"]), f2(r["Q_bio_t_hr"]), f2(r["v2_fetch_waste_t_hr"]), f2(r["v2_fetch_total_t_hr"])]
           for r in sorted(ENC, key=lambda r: (r["hotspot_id"], r["flight"], int(r["leg_id"])))]
c10 = O.do_table(s, 9, S10rows, None, LS)
print(f"tables: S1 {c1} cells changed, S2 {c2}, S3 {c3}, S10 {c10} (expect S10 0)")
assert c10 == 0, "Table S10 changed: crossings differ from the SI"
for p in L.all_paragraphs(s): pass
L.subscript_chem(L.all_paragraphs(s)); L.exponent_fix(L.all_paragraphs(s))
s.save(a.si_out)

# ---- checks ----------------------------------------------------------------------------------
tm = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.main_out)))
ts = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.si_out)))
for b in ("before, during, and after", "17 of the 22", "3.35 and 2.75", "all 34 urban", "97% of the samples", "10 of the 13", "best estimate", "other five"):
    assert b not in tm, f"main: {b!r} remains"
for b in ("Authors to complete", "Five flights have no airborne", "3 July L2 and 12 July L1", "0.63 to 3.35", "10 of the 13", "best estimate", "seven flights with airborne", "three flights here"):
    assert b not in ts, f"SI: {b!r} remains"
assert "1502 / 3352" in ts and "- / 2752" in ts and "- / 4152" in ts
print("checks passed")
