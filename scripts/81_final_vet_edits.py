"""81_final_vet_edits.py -- fixes from the last full vet of the main text and SI (2 Oct 2026).

Numbers checked against the outputs (values read from the CSVs, not typed):
  * Section 4.3 allocation: 2.4 t/h to municipal solid-waste landfills, 2.1 t/h to domestic wastewater
    (source_attribution_representative.csv; the text said 2.5 and 2.0).
  * Discussion: exceedance factors 1.5-6 (lowest estimate 4.1 t/h CH4:CO2 against the 2.8 t/h EPA
    total; the 1.6 predated the EPA area weighting); EPA biogenic 1.0 t/h of which landfills and
    wastewater 0.9 (inventory_comparison.csv).
  * GRA2PES comparison: 24 of the 26 aircraft cells fall inside the GRA2PES domain
    (spatial_match_summary.csv n_cells), so the 45% / 13% shares are over 24 cells.
  * DADS: Carbon Mapper has 9 plumes on 6 dates, not "nine overpasses".
  * Endmember range written 0.06 in one place (0.061 everywhere else).
  * Lag sensitivity: 29 vs 30% is a 1.6-point difference ("within 1 percentage point" -> 2).
Claims brought in line with the results:
  * Discussion: the Denver median is below the 40-88% of other cities at the adopted calibration
    and at or below its low end across the endmembers (48% at 0.061 is inside 40-60%).
  * 4.1 ground check: similar central value and spread, and the car sampled mainly the box's
    northern edge (the SI wording), rather than "the same ratio".
  * 3.2 / SI S3 leg floors: they exclude an ethane-poor fossil majority for a basin ratio near
    0.061, not at 0.0813 (up to 60% distribution gas).
  * 2.5 Carbon Mapper: persistence-weighted emission rates, not a lower bound (4.3 uses them to
    compare against v2.0beta).
  * "clearest low-fossil case" (3 July L2, 20%) -> clearest wind-linked case (8 and 10 July are lower).
References and cross-references:
  * 3.2: the 0.0813 ratio carried citation 16 (Rousseeuw & Croux, the MAD paper); it now points to
    section S5, where the campaign report is cited.
  * Discussion: "Section S10 maps" -> section S7 (Figure S10); "Figure 4B and Table S11 ... these
    maps" -> Table S11 is the separation scan, not a map.
  * Reference 11 author field ("Laboratory (CSL), N. C. S.").
  * SI S9: four author-year citations (STILT, STILT-R, HYSPLIT, HRRR) become numbered SI refs 20-23.
  * SI Figure S2 caption: the Vulcan basemap is in Figures 1 and S16 (not S1).
  * SI Figure S5 caption: 0.0377 -> 0.037.
Method descriptions:
  * SI S2.2: the ARL car's calibration cycles also ran while driving (transits of 24 June and
    15 July), and the screen removes them in moving and parked stretches alike.
  * SI S11: the crossing rates use the lidar mixing height over the flight's urban legs, which can
    differ from the whole-flight value in Table S1.
  * Main 3.1: three platforms, not two.
  * SI S3: the daytime statement applies to the eight analysis flights.
  * SI S1.1.1: drop the pointer to an "accompanying equations reference" that does not exist.
  * Table S1 caption wording; Table S6 (12 July L1 CH4:CO2 "NA | 1") and Table S7 (one-leg CO2
    intervals) set to n/a.
Typos: "ta slope is considerable usable", "very meaningful", "it's most fossil-influence",
"an explicitly boundary-layer-height", "interval.This", "Julesberg", "increases the CH4:CO ratio".

usage: python3 scripts/81_final_vet_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, copy, re, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
from docx_edit_lib import R, P, runs_text, all_paragraphs

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.OUTDIR = str(OUT)
rd = lambda f: list(csv.DictReader(open(OUT / f)))

# ---- numbers ------------------------------------------------------------------------------------
ATT = {r["sector"]: float(r["top_down_t_hr"]) for r in rd("source_attribution_representative.csv")}
landfill_msw, ww_dom = ATT["5A1_Landfills_MSW"], ATT["5D_Wastewater_Treatment_Domestic"]
INV = {r["sector"]: float(r["t_hr"]) for r in rd("inventory_comparison.csv")}
epa_total = sum(INV.values())
epa_bio = sum(v for k, v in INV.items() if k.startswith(("3", "5")))
epa_lw = sum(v for k, v in INV.items() if k.startswith(("5A1", "5D")))
EE = rd("emission_estimates.csv")
ests = [float(r[c]) for r in EE for c in ("CH4CO_GRA2PES_t_hr", "CH4CO2_Vulcan_t_hr") if r[c] not in ("", "NA")]
PV = open(OUT / "paper_values.json").read()
gra11 = float(re.search(r'"E_CH4_GRA2PES"\s*:\s*([0-9.]+)', PV).group(1)) if "E_CH4_GRA2PES" in PV else 1.69
f_lo, f_hi = min(ests) / epa_total, max(ests) / min(epa_total, gra11)
SM = {r["inventory"].split(" (")[0]: r for r in rd("spatial_match_summary.csv")}
n_gra = int(SM["GRA2PES v2.0beta"]["n_cells"]); n_air = int(SM["EPA gridded GHGI 2020"]["n_cells"])
assert int(SM["GRA2PES v1.1"]["n_cells"]) == n_gra
DADS_LL = (39.662, -104.690)                                       # as in scripts/54
CMS = [r for r in rd("carbonmapper_denver_sources.csv") if r["gas"] == "CH4"
       and abs(float(r["geom_lat"]) - DADS_LL[0]) < 0.01 and abs(float(r["geom_lon"]) - DADS_LL[1]) < 0.01]
assert len(CMS) == 1, f"DADS source: {len(CMS)} rows within 0.01 deg in carbonmapper_denver_sources.csv"
dads_pl, dads_dates = int(float(CMS[0]["plume_count"])), int(float(CMS[0]["detection_date_count"]))
LAG = {int(r["lag_s"]): float(r["median_fossil_pct"]) for r in rd("lag_sensitivity_summary.csv")}
lag_gap = abs(LAG[-1] - LAG[0])                                   # best-aligned (-1 s) vs uncorrected
lag_pts = int(-(-lag_gap // 1))                                    # ceiling
ARC = {r["quantity"]: r["value"] for r in rd("arc_ground_summary.csv")}
words = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine"}
print(f"landfills {landfill_msw:.2f}, wastewater {ww_dom:.2f}; EPA total {epa_total:.2f}, biogenic {epa_bio:.2f}, "
      f"landfill+ww {epa_lw:.2f}; factors {f_lo:.2f}-{f_hi:.1f}; GRA2PES cells {n_gra}/{n_air}; DADS {dads_pl} plumes "
      f"{dads_dates} dates; lag gap {lag_gap:.1f} -> {lag_pts}")

# ---- main text ----------------------------------------------------------------------------------
R("Denver-Julesberg", "Denver-Julesburg")
R("do not require an explicitly boundary-layer-height or wind terms", "do not require explicit boundary-layer-height or wind terms")
R("identically to every 1 Hz series from both platforms.", "identically to every 1 Hz series from all three platforms.")
R("interval.This procedure", "interval. This procedure")
R("indicating that it is unlikely a very ethane-poor fossil source dominated those legs",
  "indicating that, for a basin gas ratio near 0.061 mol mol-1, a very ethane-poor fossil source is unlikely to have dominated those legs")
R("Specifically, ta slope is considerable usable", "Specifically, a slope is considered usable")
R("and provide an independent lower bound on emissions from detected facilities.",
  "and provide independent, persistence-weighted emission rates for the detected facilities.")
R("Denver-relevant endmember range of 0.06-0.110", "Denver-relevant endmember range of 0.061-0.110")
R("so a street-level platform and the aircraft return the same urban enhancement ratio within their day-to-day spread",
  "so a street-level platform and the aircraft return urban enhancement ratios of similar central value and spread, "
  "although the car sampled mainly the northern edge of the box")
R("Carbon Mapper detected the DADS landfill on nine overpasses",
  f"Carbon Mapper detected plumes from the DADS landfill {words.get(dads_pl, dads_pl)} times on {words.get(dads_dates, dads_dates)} dates")
R("very meaningful day-to-day", "vary meaningfully from day to day")
R("increases the CH4:CO ratio to", "raises the CH4:CO-based estimates to")
R("The clearest low-fossil case, 3 July L2", "The clearest wind-linked low-fossil case, 3 July L2")
R("assigns approximately 2.5 t CH4 h-1 to landfills, 2.0 t CH4 h-1 to domestic wastewater",
  f"assigns approximately {landfill_msw:.1f} t CH4 h-1 to municipal solid-waste landfills, {ww_dom:.1f} t CH4 h-1 to domestic wastewater")
R("wastewater estimate of approximately 2.0 t CH4 h-1", f"wastewater estimate of approximately {ww_dom:.1f} t CH4 h-1")
R(f"The {n_air} sampled cells contain", f"The {n_gra} sampled cells inside the GRA2PES domain contain")
R("Figure 4B and Table S11; importantly, these maps depict",
  "Figure 4B, with the per-flight separation scan in Table S11; importantly, the map depicts")
R("by factors of roughly 1.6-6 across flights and anchors", f"by factors of roughly {f_lo:.1f}-{round(f_hi)} across flights and anchors")
R("Biogenic sources contribute approximately 1.0 t h-1 from landfills and wastewater,",
  f"Biogenic sources contribute approximately {epa_bio:.1f} t h-1, {epa_lw:.1f} t h-1 of it from landfills and wastewater,")
R("Section S10 maps this spatial separation directly", "Section S7 (Figure S10) maps this spatial separation directly")
R("Denver’s median fossil fraction is lower than the 40-88% values",
  "Denver’s median fossil fraction is at or below the low end of the 40-88% values")
R("lies below all of them, though it’s most fossil-influence flight",
  "lies below all of them at the adopted calibration and at or below their low end across the Denver-relevant "
  "endmembers, although its most fossil-influenced flight")
P("main", "(11)\tLaboratory (CSL), N. C. S.",
  "(11)\tNOAA Chemical Sciences Laboratory (CSL). NOAA CSL Projects: AMMBEC 2024. "
  "https://csl.noaa.gov/groups/csl3/measurements/2024ammbec/ (accessed 2026-08-21)")

# ---- SI -----------------------------------------------------------------------------------------
R(" The typeset equations for the density, the loop integral, and the mixing height are in the accompanying equations reference.",
  "", "si")
R("None of the eight flights encircle the box; the rest cover only part of the perimeter around the city, or sample only two altitude levels",
  "None of the eight flights encircles the box: each covers only part of the perimeter around the city, or samples only two altitude levels", "si")
R("the same offline layer used in Figures 1 and S1)", "the same offline layer used in Figures 1 and S16)", "si")
R("Zero-air and span-gas cycles run while the car was parked were removed first, from the first cylinder sample of each "
  "parked period to the last plus two minutes either side (",
  "Zero-air and span-gas cycles were removed first. The analyzer ran them while parked and also while driving, on the "
  "transits of 24 June and 15 July, so in every stretch of parking or driving that held a cylinder sample the record was "
  "removed from the first cylinder sample to the last plus two minutes either side (", "si")
R("all flights were flown during the day (local", "all eight analysis flights were flown during the day (local", "si")
R("Those floors exclude an ethane-poor fossil source as the majority of what the tightly correlated legs sampled: with "
  "East Coast gas, at most 16 to 27% of the fossil methane on those legs could have been distribution gas for a basin "
  "ratio of 0.061 (43 to 60% for 0.0813)",
  "For a basin ratio near 0.061, those floors exclude an ethane-poor fossil source as the majority of what the tightly "
  "correlated legs sampled: with East Coast gas, at most 16 to 27% of the fossil methane on those legs could have been "
  "distribution gas. For a basin ratio of 0.0813 the bound rises to 43 to 60%, so at that ratio the floors no longer "
  "exclude it", "si")
R("East Coast delivered-gas range of 0.019 to 0.0377)", "East Coast delivered-gas range of 0.019 to 0.037)", "si")
R("The uncorrected alignment therefore lies within 1 percentage point of the best-aligned result",
  f"The uncorrected alignment therefore lies within {lag_pts} percentage points of the best-aligned result "
  f"({LAG[0]:.0f} against {LAG[-1]:.0f}%)", "si")
R(f"(the {n_air} lie in 7 EPA cells and 15 GRA2PES cells)",
  f"(the {n_air} lie in 7 EPA cells, and the {n_gra} inside the GRA2PES domain lie in 15 GRA2PES cells)", "si")
R("STILT (HYSPLIT in STILT emulation; Lin et al., 2003; Fasoli et al., 2018; Stein et al., 2015)",
  "STILT[[^20,21]] (HYSPLIT in STILT emulation[[^22]])", "si")
R("archived by NOAA ARL (Dowell et al., 2022)", "archived by NOAA ARL[[^23]]", "si")
R("multiplied by the lidar mixing height of the flight (uniform mixing to that height).",
  "multiplied by the lidar mixing height over that flight's urban legs (uniform mixing to that height; this median "
  "covers the period of the crossings and can differ from the whole-flight value in Table S1).", "si")

# ---- main: citation 16 after the 0.0813 ratio --------------------------------------------------------
def drop_cite(path_in, path_out, anchor, num):
    d = docx.Document(path_in)
    hits = [p for p in d.paragraphs if anchor in p.text]
    assert len(hits) == 1, f"{anchor!r}: {len(hits)} paragraphs"
    p = hits[0]; seen = ""; done = False
    for t in list(p._p.iter(qn("w:t"))):
        if not done and seen.endswith(anchor) and (t.text or "").strip() == num:
            r = t.getparent(); rp = r.find(qn("w:rPr")); va = rp.find(qn("w:vertAlign")) if rp is not None else None
            assert va is not None and va.get(qn("w:val")) == "superscript", "citation run is not superscript"
            par = r.getparent(); par.remove(r)
            if par.tag == qn("w:hyperlink") and not list(par.iter(qn("w:t"))): par.getparent().remove(par)
            done = True; break
        seen += t.text or ""
    assert done, f"citation {num} after {anchor!r} not found"
    d.save(path_out)
tmp_main = Path(a.main_out).with_suffix(".tmp.docx")
drop_cite(a.main_in, tmp_main, "contemporary production-area ratio of 0.0813 mol mol-1.", "16")
R("contemporary production-area ratio of 0.0813 mol mol-1. The campaign",
  "contemporary production-area ratio of 0.0813 mol mol-1 (section S5). The campaign")

L.apply(tmp_main, a.main_out, "main", "main text"); tmp_main.unlink()
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- SI tables S6 and S7: artefacts for single-leg / undefined fits -----------------------------------
si = docx.Document(a.si_out)
def table_by_caption(prefix):
    body = list(si.element.body.iterchildren())
    for k, el in enumerate(body):
        if el.tag == qn("w:p") and "".join(x.text or "" for x in el.iter(qn("w:t"))).strip().startswith(prefix):
            for el2 in body[k + 1:]:
                if el2.tag == qn("w:tbl"): return docx.table.Table(el2, si)
    raise SystemExit(f"table {prefix} not found")
def set_cell(cell, text):
    cp = cell.paragraphs[0]
    for r_ in cp.runs[1:]: r_._r.getparent().remove(r_._r)
    cp.runs[0].text = text; cp.runs[0].font.color.rgb = L.RED
n6 = n7 = 0
t6 = table_by_caption("Table S6:"); h6 = [c.text for c in t6.rows[0].cells]
for row in t6.rows[1:]:
    cells = row.cells
    for j in range(len(h6) - 1):
        if h6[j].startswith("CH4:CO2") and cells[j].text.strip() == "NA":
            set_cell(cells[j], "n/a"); n6 += 1
            if cells[j + 1].text.strip() in ("1", "-1", "NA"): set_cell(cells[j + 1], "n/a"); n6 += 1
t7 = table_by_caption("Table S7:"); h7 = [c.text for c in t7.rows[0].cells]
i_legs, i_ci = h7.index("Legs (CO2)"), next(i for i, h in enumerate(h7) if h.startswith("E from CO2"))
for row in t7.rows[1:]:
    if row.cells[i_legs].text.strip() == "1":
        lo_hi = row.cells[i_ci].text.split(" to ")
        if len(lo_hi) == 2 and lo_hi[0].strip() == lo_hi[1].strip(): set_cell(row.cells[i_ci], "n/a (one leg)"); n7 += 1
print(f"Table S6: {n6} cell(s) set to n/a; Table S7: {n7} single-leg interval(s) set to n/a")

# ---- SI references 20-23 ------------------------------------------------------------------------------
NEW_REFS = [
  (20, "Lin, J. C.; Gerbig, C.; Wofsy, S. C.; Andrews, A. E.; Daube, B. C.; Davis, K. J.; Grainger, C. A. A Near-Field "
       "Tool for Simulating the Upstream Influence of Atmospheric Observations: The Stochastic Time-Inverted Lagrangian "
       "Transport (STILT) Model. J. Geophys. Res. Atmos. 2003, 108 (D16), 4493. https://doi.org/10.1029/2002JD003161"),
  (21, "Fasoli, B.; Lin, J. C.; Bowling, D. R.; Mitchell, L.; Mendoza, D. Simulating Atmospheric Tracer Concentrations "
       "for Spatially Distributed Receptors: Updates to the Stochastic Time-Inverted Lagrangian Transport Model's R "
       "Interface (STILT-R Version 2). Geosci. Model Dev. 2018, 11 (7), 2813–2824. https://doi.org/10.5194/gmd-11-2813-2018"),
  (22, "Stein, A. F.; Draxler, R. R.; Rolph, G. D.; Stunder, B. J. B.; Cohen, M. D.; Ngan, F. NOAA's HYSPLIT Atmospheric "
       "Transport and Dispersion Modeling System. Bull. Am. Meteorol. Soc. 2015, 96 (12), 2059–2077. "
       "https://doi.org/10.1175/BAMS-D-14-00110.1"),
  (23, "Dowell, D. C.; Alexander, C. R.; James, E. P.; Weygandt, S. S.; Benjamin, S. G.; Manikin, G. S.; Blake, B. T.; "
       "Brown, J. M.; Olson, J. B.; Hu, M.; Smirnova, T. G.; Ladwig, T.; Kenyon, J. S.; Ahmadov, R.; Turner, D. D.; "
       "Duda, J. D.; Alcott, T. I. The High-Resolution Rapid Refresh (HRRR): An Hourly Updating Convection-Allowing "
       "Forecast Model. Part I: Motivation and System Description. Weather Forecast. 2022, 37 (8), 1371–1395. "
       "https://doi.org/10.1175/WAF-D-21-0151.1"),
]
refs = [p for p in si.paragraphs if re.match(r"\(\d+\)\s", p.text.strip())]
have = [int(re.match(r"\((\d+)\)", p.text.strip()).group(1)) for p in refs]
assert have == list(range(1, len(have) + 1)) and have[-1] == 19, f"SI reference list is {have}"
last = refs[-1]
for n, txt in NEW_REFS:
    newp = copy.deepcopy(last._p); last._p.addnext(newp)
    p = docx.text.paragraph.Paragraph(newp, last._parent)
    for child in list(p._p):
        if child.tag != qn("w:pPr"): p._p.remove(child)
    run = p.add_run(f"({n})\t{txt}"); run.font.color.rgb = L.RED; last = p
si.save(a.si_out)
print("SI references 20-23 added")
