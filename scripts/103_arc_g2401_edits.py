"""103_arc_g2401_edits.py -- ARL car ratios with the Picarro G2401-m methane (S. Baidar / X. Ren,
16 Oct 2026): main 2.3 and 4.1, SI S2.2, S3 and S5, Figure S3 (image and caption). Works on
Priyanka's final_aircraft files after script 101. Numbers are read from the outputs of scripts 71
and 102 and asserted. New text is red.

usage: python3 scripts/103_arc_g2401_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, re, argparse, shutil
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs)
rd = lambda f: list(csv.DictReader(open(OUT / f)))
S = {r["quantity"]: r["value"] for r in rd("arc_ground_summary.csv")}
CV = {r["date"]: r for r in rd("arc_ground_coverage.csv")}
CMP = {r["region"]: r for r in rd("arc_analyzer_comparison.csv")}
PL = rd("arc_ground_plumes.csv")
f3 = lambda k: f"{float(S[k]):.3f}"
ci = lambda k: " to ".join(f"{float(x):.3f}" for x in S[k].split(" to "))
assert S["ch4_analyzer"] == "g2401", S["ch4_analyzer"]
lags = [int(v.split(":")[1]) for v in S["g2401_lag_s_by_day"].split()]
win = [d for d in CV if "20240628" <= d <= "20240713"]
n_in = sum(int(CV[d]["n_urban"]) for d in win); n_bas = sum(int(CV[d]["n_basin"]) for d in win)
m9, m11 = int(CV["20240709"]["n_urban_metro"]), int(CV["20240711"]["n_urban_metro"])
others = sum(int(CV[d]["n_urban_metro"]) for d in win if d not in ("20240709", "20240711"))
fo_u = [p for p in PL if p["region"] == "urban" and p["class"] == "fossil-type"]
xl, yl = (-105.70, -104.05), (39.20, 40.85)
shown = sum(1 for p in PL if xl[0] <= float(p["lon"]) <= xl[1] and yl[0] <= float(p["lat"]) <= yl[1])
amb_u = float(S["ambient_slope_median_window_urban"]); amb_b = float(S["ambient_slope_median_window_basin"])
rng = lambda k: " to ".join(("−" if float(x) < 0 else "") + f"{abs(float(x)):.3f}" for x in S[k].split(" to "))
N = dict(pu=S["plumes_urban"], fu=S["plumes_fossil_type_urban"], eu=S["plumes_ethane_free_urban"], mu=f3("fossil_type_ratio_median_urban"),
         ciu=ci("fossil_type_ratio_ci_urban"), iqu=ci("fossil_type_ratio_iqr_urban"), wu=f3("fossil_type_ratio_median_window_urban"), nwu=S["plumes_fossil_type_window_urban"],
         pb=S["plumes_basin"], fb=S["plumes_fossil_type_basin"], eb=S["plumes_ethane_free_basin"], mb=f3("fossil_type_ratio_median_basin"),
         cib=ci("fossil_type_ratio_ci_basin"), iqb=ci("fossil_type_ratio_iqr_basin"), wb=f3("fossil_type_ratio_median_window_basin"), nwb=S["plumes_fossil_type_window_basin"],
         au=f"{amb_u:.3f}", aur=rng("ambient_slope_range_window_urban"), ab=f"{amb_b:.3f}", abr=rng("ambient_slope_range_window_basin"),
         nd_u=S["ambient_n_days_urban"])
below_pct = round(100 * (1 - amb_b / 0.0813))
fmed = round(21 * 0.102 / float(N["mu"]))                 # campaign median if the urban plume ratio were the endmember (fossil fraction scales as 1/beta)
assert (N["pu"], N["fu"], N["eu"], N["mu"], N["pb"], N["fb"], N["eb"], N["mb"], N["au"], N["ab"]) == ("22", "5", "5", "0.057", "174", "43", "59", "0.107", "0.020", "0.059"), N
assert all(float(p["lat"]) >= 39.79 for p in fo_u) and len(fo_u) == 5
assert (m9, m11, others, n_in, n_bas, shown, len(PL)) == (1566, 42, 0, 25514, 124365, 222, 233), (m9, m11, others, n_in, n_bas, shown, len(PL))
assert (min(lags), max(lags), below_pct, fmed) == (-1, 2, 27, 38), (lags, below_pct, fmed)
ig_b, ig_u = float(CMP["basin"]["int_g_over_a"]), float(CMP["urban"]["int_g_over_a"]); yg_b = float(CMP["basin"]["york_g_over_a"])
assert (round(100 * (1 - ig_b)), round(100 * (1 - yg_b))) == (7, 15)

# ================================ main ================================
m = docx.Document(a.main_in); LM = "main"
O.do_replace(m, "the archived files report CH4 from the Aeris, the G2401 and a Picarro G2201i isotope analyzer, and the ethane-to-methane ratios here use the Aeris CH4 and C2H6, measured by the same analyzer.",
             "the archived files report CH4 from the Aeris, the G2401 and a Picarro G2201i isotope analyzer, and the ethane-to-methane ratios here use the G2401 CH4, which has the best CH4 precision of the three, with the Aeris C2H6 aligned to it in time (section S2.2).", 1, LM)
O.do_replace(m, "ethane-bearing plumes give median ratios of 0.084 mol mol−1 inside the box (4 plumes, all north of 39.79°N) and 0.137 mol mol−1 in the basin (47 plumes)",
             f"ethane-bearing plumes give median ratios of {N['mu']} mol mol−1 inside the box ({N['fu']} plumes, all north of 39.79°N) and {N['mb']} mol mol−1 in the basin ({N['fb']} plumes)", 1, LM)
O.do_replace(m, "gives an ambient ratio of 0.023 mol mol−1 as a median over 9 days (-0.004 to 0.062)",
             f"gives an ambient ratio of {N['au']} mol mol−1 as a median over {N['nd_u']} days ({N['aur']})", 1, LM)
m.save(a.main_out)

# ================================ SI ================================
s = docx.Document(a.si_in); LS = "SI"
O.do_replace(s, "(1,567 one-second records; 42 on 11 July, and none on the other days)", f"({m9:,} one-second records; {m11} on 11 July, and none on the other days)", 1, LS)
O.do_replace(s, "within the flight window 124,429 of its one-second records fall north of 40.05°N against 29,550 inside the box.",
             f"within the flight window {n_bas:,} of its one-second records with valid methane fall north of 40.05°N against {n_in:,} inside the box.", 1, LS)
# methods: analyzer and alignment, after the calibration-screen sentence
anchor = [p for p in s.paragraphs if O.vis_text(p).startswith("We treated each drive as the mobile analogue of a flight.")]
assert len(anchor) == 1
O.do_replace(s, "We treated each drive as the mobile analogue of a flight.",
             "We treated each drive as the mobile analogue of a flight. Methane is taken from the Picarro G2401-m, which has the best CH4 precision of the car's three methane analyzers, and ethane from the Aeris; "
             f"because the two sample through separate lines, each day's G2401 record was shifted onto the Aeris time base by the lag ({str(min(lags)).replace('-', '\u2212')} to +{max(lags)} s) that maximizes the correlation of the two instruments' 61-s high-passed methane records, "
             "and samples without a valid G2401 methane value were dropped.", 1, LS)
O.do_replace(s, "Over all drives (Figure S3), the car crossed 23 plumes inside the box, of which 4 are ethane-bearing, all north of 39.79°N, with a median ethane-to-methane ratio of 0.084 mol mol−1 (95% confidence interval 0.069 to 0.100; interquartile range 0.069 to 0.098), and 4 are ethane-free. In the basin 47 of 181 plumes are ethane-bearing, with a median of 0.137 mol mol−1 (0.103 to 0.167; interquartile range 0.095 to 0.183), and 65 are ethane-free. Restricted to the flight window the medians are 0.070 (3 plumes) and 0.138 (46 plumes) mol mol−1.",
             f"Over all drives (Figure S3), the car crossed {N['pu']} plumes inside the box, of which {N['fu']} are ethane-bearing, all north of 39.79°N, with a median ethane-to-methane ratio of {N['mu']} mol mol−1 (95% confidence interval {N['ciu']}; interquartile range {N['iqu']}), and {N['eu']} are ethane-free. "
             f"In the basin {N['fb']} of {N['pb']} plumes are ethane-bearing, with a median of {N['mb']} mol mol−1 ({N['cib']}; interquartile range {N['iqb']}), and {N['eb']} are ethane-free. "
             f"Restricted to the flight window the medians are {N['wu']} ({N['nwu']} plumes) and {N['wb']} ({N['nwb']} plumes) mol mol−1. "
             f"The choice of methane analyzer matters at this level: on the plumes that are ethane-bearing with either instrument, the plume-integrated ethane-to-methane ratio is {round(100 * (1 - ig_b))}% lower in the basin "
             f"({round(100 * (1 - ig_u))}% inside the box) with the G2401 methane than with the Aeris methane, and the York slope {round(100 * (1 - yg_b))}% lower, because the G2401 records slightly larger methane enhancements and responds more slowly than the Aeris.", 1, LS)
O.do_replace(s, "The median daily ratio was 0.074 mol mol-1 (range: 0.000-0.137), within 10% of the production-area ratio of 0.0813 mol mol-1 reported in the AMMBEC report to CDPHE.",
             f"The median daily ratio was {N['ab']} mol mol-1 (range: {N['abr']}), {below_pct}% below the production-area ratio of 0.0813 mol mol-1 reported in the AMMBEC report to CDPHE; part of that difference reflects the methane analyzer (see above).", 1, LS)
O.do_replace(s, "the median daily ethane-to-methane ratio measured by the car was 0.023 mol mol-1 (range: -0.004 to 0.062 across nine days)",
             f"the median daily ethane-to-methane ratio measured by the car was {N['au']} mol mol-1 (range: {N['aur']} across {'nine' if N['nd_u'] == '9' else N['nd_u']} days)", 1, LS)
O.do_replace(s, "while the all-drive median of the ethane-bearing plumes inside the box, 0.084 mol mol−1, falls within the regional reference range of 0.081 to 0.110 mol mol−1 (the flight-window median of 0.070 mol mol−1, from 3 plumes, lies just below it).",
             f"while the all-drive median of the {N['fu']} ethane-bearing plumes inside the box, {N['mu']} mol mol−1 ({N['wu']} in the flight window, {N['nwu']} plumes), lies below the regional reference range of 0.081 to 0.110 mol mol−1 but above the 0.042 mol mol−1 break-even of section S3. "
             f"These few plumes, all on the industrial northern edge of the box, do not define the effective endmember of the urban fossil mixture; if they did, the campaign median would be about {fmed}%, still biogenic-dominated.", 1, LS)
O.do_replace(s, "231 of the 242 plumes fall in the plotted area", f"{shown} of the {len(PL)} plumes fall in the plotted area", 1, LS)
O.do_replace(s, "Resolving the same ground record plume by plume puts it higher still: 47 of the 181 basin plumes crossed by the ground car carry ethane tracking methane, at a median of 0.137 mol mol−1 (95% confidence interval 0.103 to 0.167; section S2.2), while 65 show methane with no ethane response.",
             f"Resolving the same ground record plume by plume puts it higher still: {N['fb']} of the {N['pb']} basin plumes crossed by the ground car carry ethane tracking methane, at a median of {N['mb']} mol mol−1 (95% confidence interval {N['cib']}; section S2.2), while {N['eb']} show methane with no ethane response.", 1, LS)
O.do_replace(s, "The car-data analysis in section S2.2 gives a ratio within 10% of that production-area value.",
             f"The car's daily basin ratio in section S2.2 ({N['ab']} mol mol−1 with the G2401 methane) lies below that production-area value.", 1, LS)
L.subscript_chem(L.all_paragraphs(s)); s.save(a.si_out)

# Figure S3 image: the picture just before its caption
d = docx.Document(a.si_out); body = list(d.element.body.iterchildren()); target = None
for k, el in enumerate(body):
    if "".join(x.text or "" for x in el.iter(qn("w:t"))).startswith("Figure S3:"):
        for j in range(k - 1, max(0, k - 4), -1):
            emb = [b.get(qn("r:embed")) for b in body[j].iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")]
            if emb: target = d.part.rels[emb[0]].target_ref.split("/")[-1]; break
assert target, "Figure S3 image not found"
O.swap_figures(a.si_out, {target: OUT / "figures" / "FigS15_arc_ground.png"}, set(), "SI")

tm = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.main_out)))
ts = "\n".join(O.vis_text(p) for p in L.all_paragraphs(docx.Document(a.si_out)))
for b in ("0.084", "0.137", "Aeris CH4 and C2H6, measured by the same analyzer"):
    assert b not in tm, f"main: {b!r} remains"
for b in ("0.084", "0.137", "within 10%", "1,567", "29,550", "124,429", "231 of the 242", "47 of 181", "47 of the 181"):
    assert b not in ts, f"SI: {b!r} remains"
print("ARL car (G2401 CH4) edits applied; Figure S3 image", target, "replaced")
