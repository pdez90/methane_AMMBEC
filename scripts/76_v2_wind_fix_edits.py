"""76_v2_wind_fix_edits.py -- carry the corrected GRA2PES v2.0beta point-source test (script 59) into
the manuscript after the wind-direction fix of 1 Oct 2026.

Script 59 used to average wind directions arithmetically to orient each crossing's upwind fetch.
For a segment whose winds straddle north that mean points the wrong way: the 10 July L1 leg-1
crossing averaged to 112 deg, whereas its vector-mean wind is 25 deg, and with the fetch pointed
correctly the Suncor/Metro waste cell falls inside it. The test therefore has one more crossing
than the text reported. This script rewrites the affected sentences (main 4.3, SI S11), rebuilds
Table S10 from the encounters CSV, and replaces the Figure S17 picture. Every number is read from
v2_pointsource_encounters.csv; edits are exact-anchor and red.

usage: python3 scripts/76_v2_wind_fix_edits.py <main.docx> <si.docx> <main_out> <si_out>
           --outputs <OUT_DIR> --figures <OUT_DIR>/figures
"""
import sys, csv, copy, argparse, statistics
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
from docx_edit_lib import R, FIGSWAP

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True); ap.add_argument("--figures", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.FIGDIR = a.figures; L.OUTDIR = str(OUT)

E = [r for r in csv.DictReader(open(OUT / "v2_pointsource_encounters.csv")) if r["hotspot_in_fetch"] == "TRUE"]
H = {r["hotspot_id"]: r for r in csv.DictReader(open(OUT / "v2_pointsource_hotspots.csv"))}
num = lambda r, k: float(r[k])
WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine"}
def w(n, cap=False): s = WORDS.get(n, str(n)); return s.capitalize() if cap else s
def med(v): return statistics.median(v)

n_enc = len(E); n_fl = len({r["flight"] for r in E}); n_site = len({r["hotspot_id"] for r in E})
dads = [r for r in E if H[r["hotspot_id"]]["nearest_facility"].startswith("Denver Arapahoe")]
sm = [r for r in E if H[r["hotspot_id"]]["nearest_facility"].startswith("Suncor/Metro")]
assert len(dads) == 1, "the DADS sentence assumes a single crossing"
sm_bio = med([num(r, "Q_bio_t_hr") for r in sm]); sm_ratio = med([num(r, "ratio_bio_over_v2waste") for r in sm])
sm_lo, sm_hi = min(num(r, "Q_t_hr") for r in sm), max(num(r, "Q_t_hr") for r in sm)
cell_waste = num(H[sm[0]["hotspot_id"]], "waste_t_hr")
tot_lo, tot_hi = min(num(r, "Q_t_hr") for r in E), max(num(r, "Q_t_hr") for r in E)
bio_lo, bio_hi = min(num(r, "Q_bio_t_hr") for r in E), max(num(r, "Q_bio_t_hr") for r in E)
rb = [num(r, "ratio_bio_over_v2waste") for r in E]; rt = [num(r, "ratio_total_over_v2total") for r in E]
# Sensitivity: the crossings flown in light winds, where a mean direction is least well defined.
LIGHT = 2.0
sm_steady = [r for r in sm if num(r, "ws_ms") >= LIGHT]
light = [r for r in sm if num(r, "ws_ms") < LIGHT]
sens = ""
if light and sm_steady:
    sens = (f" Excluding the {w(len(light))} crossing{'s' if len(light) > 1 else ''} flown in winds below "
            f"{LIGHT:.0f} m s-1 ({', '.join(str(int(r['flight'][6:8])) + ' July, leg ' + r['leg_id'] for r in light)}), where the mean wind direction is "
            f"least well defined, gives {med([num(r, 'Q_bio_t_hr') for r in sm_steady]):.2f} t/hr and "
            f"{med([num(r, 'ratio_bio_over_v2waste') for r in sm_steady]):.2f} times.")
print(f"{n_enc} crossings on {n_fl} flights at {n_site} sites; Suncor/Metro {len(sm)}: bio median {sm_bio:.2f}, "
      f"ratio {sm_ratio:.2f}, total {sm_lo:.2f}-{sm_hi:.2f}; all: bio/v2waste {med(rb):.2f} ({min(rb):.2f}-{max(rb):.2f}), "
      f"total/v2total {med(rt):.2f}{'; ' + sens.strip() if sens else ''}")

# ---- main text 4.3 ------------------------------------------------------------------------
R("yielding four crossings on three flights at two of the three sites",
  f"yielding {w(n_enc)} crossings on {w(n_fl)} flights at {w(n_site)} of the three sites")
R("Three crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery complex, "
  "assigned 0.39 t CH4 h-1 of waste emissions in v2.0beta, yield a median biogenic rate of 1.10 t CH4 h-1, or "
  "2.82 times the inventory value; total methane rates across those crossings range from 0.36 to 3.59 t CH4 h-1.",
  f"{w(len(sm), True)} crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery "
  f"complex, assigned {cell_waste:.2f} t CH4 h-1 of waste emissions in v2.0beta, yield a median biogenic rate of "
  f"{sm_bio:.2f} t CH4 h-1, or {sm_ratio:.2f} times the inventory value in the median; total methane rates across "
  f"those crossings range from {sm_lo:.2f} to {sm_hi:.2f} t CH4 h-1.")

# ---- SI S11 -------------------------------------------------------------------------------
R("There are four crossings on three flights at two of the hotspot sites (Table S10, Figure S17). Their total rates "
  "are 0.4 to 3.6 t/hr and their biogenic parts 0.3 to 2.4 t/hr. Relative to the v2.0beta waste emissions in the "
  "fetch, the aircraft biogenic rate is 1.77 in the median (range 0.07 to 6.03); relative to the v2.0beta total in "
  "the fetch, the aircraft total is 2.02 in the median.",
  f"There are {w(n_enc)} crossings on {w(n_fl)} flights at {w(n_site)} of the hotspot sites (Table S10, Figure S17). "
  f"Their total rates are {tot_lo:.2f} to {tot_hi:.2f} t/hr and their biogenic parts {bio_lo:.2f} to {bio_hi:.2f} t/hr. "
  f"Relative to the v2.0beta waste emissions in the fetch, the aircraft biogenic rate is {med(rb):.2f} in the median "
  f"(range {min(rb):.2f} to {max(rb):.2f}); relative to the v2.0beta total in the fetch, the aircraft total is "
  f"{med(rt):.2f} in the median. Pooling across sites mixes a landfill and a wastewater cell, so the per-site "
  "values below are the informative ones. The fetch of each crossing is oriented by the vector mean of the "
  "aircraft wind along the crossed segment.", "si")
R("the three crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery complex "
  "(v2.0beta waste 0.39 t/hr in that cell) give a biogenic rate of 1.10 t/hr in the median (total 0.36 to 3.59 "
  "t/hr), 2.82 times the v2.0beta waste emissions in the upwind fetch.",
  f"the {w(len(sm))} crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery "
  f"complex (v2.0beta waste {cell_waste:.2f} t/hr in that cell) give a biogenic rate of {sm_bio:.2f} t/hr in the "
  f"median (total {sm_lo:.2f} to {sm_hi:.2f} t/hr), {sm_ratio:.2f} times the v2.0beta waste emissions in the upwind "
  f"fetch.{sens}", "si")

# Figure S17 is media image7.png in the SI (by name: the ARC figure has the same pixel size).
FIGSWAP("si", "image7.png", str(Path(a.figures) / "FigS14_v2_pointsource.png"))
L.apply(a.main_in, a.main_out, "main", "main text")
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- Table S10: rebuild the data rows from the CSV ---------------------------------------
si = docx.Document(a.si_out)
body = si.element.body; ti = -1; target = None; last = ""
for el in body.iterchildren():
    if el.tag == qn("w:p"):
        t = "".join(x.text or "" for x in el.iter(qn("w:t")))
        if t.strip(): last = t
    elif el.tag == qn("w:tbl"):
        ti += 1
        hdr = " | ".join(c.text for c in si.tables[ti].rows[0].cells)
        if hdr.startswith("Flight | Leg | Site | Distance (km)"): target = si.tables[ti]
assert target is not None, "Table S10 (crossings) not found"
cols = ["flight", "leg_id", "hotspot_id", "dist_km", "agl_m", "blh_m", "dch4_max_ppb", "Q_t_hr", "fossil_frac",
        "Q_bio_t_hr", "v2_fetch_waste_t_hr", "v2_fetch_total_t_hr", "v1_fetch_total_t_hr"]
fmt = [str, str, str, "{:.1f}", "{:.0f}", "{:.0f}", "{:.0f}", "{:.2f}", "{:.2f}", "{:.2f}", "{:.2f}", "{:.2f}", "{:.2f}"]
assert len(target.rows[0].cells) == len(cols), "Table S10 has a different column layout than expected"
template = copy.deepcopy(target.rows[1]._tr)
old = [[c.text for c in r.cells] for r in target.rows[1:]]
for r in list(target.rows[1:]): target._tbl.remove(r._tr)
changed = 0
for e in E:
    tr = copy.deepcopy(template); target._tbl.append(tr)
    row = target.rows[-1]
    vals = [f(e[k]) if f is str else f.format(float(e[k])) for k, f in zip(cols, fmt)]
    for cell, v in zip(row.cells, vals):
        p = cell.paragraphs[0]
        for r_ in p.runs[1:]: r_._r.getparent().remove(r_._r)
        p.runs[0].text = v
        if vals not in old: p.runs[0].font.color.rgb = L.RED; changed += 1
print(f"Table S10: {len(old)} -> {len(E)} rows ({changed} cells marked red)")

# ---- caption placement --------------------------------------------------------------------
# Every SI table caption sits above its table except the three added on 15 Sep (S8-S10), whose
# captions were inserted below. Move any caption that directly follows a table to directly above it.
moved = 0
for el in list(si.element.body.iterchildren()):
    if el.tag != qn("w:p"): continue
    t = "".join(x.text or "" for x in el.iter(qn("w:t")))
    if not t.startswith("Table S"): continue
    prev = el.getprevious()
    while prev is not None and prev.tag == qn("w:p") and not "".join(x.text or "" for x in prev.iter(qn("w:t"))).strip():
        prev = prev.getprevious()                         # skip blank spacer paragraphs
    if prev is not None and prev.tag == qn("w:tbl"):
        prev.addprevious(el); moved += 1
si.save(a.si_out)
print(f"table captions moved above their tables: {moved}")
