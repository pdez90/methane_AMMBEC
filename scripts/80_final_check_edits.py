"""80_final_check_edits.py -- fixes from the last consistency check before submission (2 Oct 2026).

  Main text
    * Figure 1 is now cited in section 2.1, so the main figures are first cited in numerical order
      (Figure 2 used to be cited in section 3.2, before Figure 1).
    * Section 3.1 background windows: name the platforms. The CDPHE surveys use a 120 s window
      (scripts 03 and 06), and the ARL car uses the same window as the aircraft (script 71).
    * The Introduction no longer says that the platforms are not compared (the ARL car's ambient
      ratio is compared with the aircraft in section 4.1).
    * Typos: "field..", "New York ," and a double space.
  SI
    * Figure S7, Table S3 and Table S6 were not cited anywhere; each is now cited in its section,
      and the opening of section S6 (a sentence fragment) becomes a sentence that cites Table S6.
    * SI references are renumbered in order of first citation, and the list is reordered to match
      (section reordering and the late additions of references 16-19 had left them out of order).
      Superscripts after a delta sign or sigma-w (the 13 of delta-13C, the 2 of sigma-w squared) or a unit (m2)
      are not citations and are left alone.

usage: python3 scripts/80_final_check_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, re, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
from docx_edit_lib import R

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.OUTDIR = str(OUT)
ARC = {r["quantity"]: r["value"] for r in csv.DictReader(open(OUT / "arc_ground_summary.csv"))}
arc_min = round(float(ARC["bg_window_s"]) / 60)

# ---- main text ----------------------------------------------------------------------------------
R("over the DJB basin and the Denver metropolitan area. A Picarro",
  "over the DJB basin and the Denver metropolitan area (Figure 1). A Picarro")
R("in the context of independent multi-year ground-based facility surveys, without comparing the platforms directly.",
  "in the context of independent ground-based measurements: multi-year facility surveys and a car that measured "
  "ethane at street level during the campaign.")
R("On the ground, the mobile surveys used it to isolate facility plumes (section S2).",
  "On the ground, the CDPHE facility surveys and the ARL car used it to isolate plumes (sections S2.1 and S2.2).")
R("Only the window width w differed by platform speed:", "Only the window width w differed between platforms:")
R("about 2 minutes for the slower vehicle.",
  f"about 2 minutes for the CDPHE survey vehicle, and {arc_min} minutes for the ARL car, whose drives were mostly "
  "long transits through the basin rather than circuits of individual facilities.")
R("whereas the vehicle works next to individual facilities", "whereas the survey vehicle works next to individual facilities")
R("field..", "field.")
R("New York ,", "New York,")
R("NOy.  None", "NOy. None")

# ---- SI text ------------------------------------------------------------------------------------
R("On that fixed panel, the median fossil fraction", "On that fixed panel (Figure S7), the median fossil fraction", "si")
R("Most flights are unchanged; the largest change", "Most flights are unchanged (Table S3); the largest change", "si")
R("For every flight with urban legs, the enhancement-ratio slopes and their Pearson correlations, and the resulting "
  "relative fossil fraction.",
  "Table S6 lists, for every flight with urban legs, the enhancement-ratio slopes and their Pearson correlations, and "
  "the resulting relative fossil fraction.", "si")

L.apply(a.main_in, a.main_out, "main", "main text")
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- SI references: renumber by first citation ---------------------------------------------------
si = docx.Document(a.si_out)
REFP = re.compile(r"\((\d+)\)\s")
refs = [p for p in si.paragraphs if REFP.match(p.text.strip())]
listed = [int(REFP.match(p.text.strip()).group(1)) for p in refs]
assert listed == list(range(1, len(listed) + 1)), f"SI reference list is not 1..N: {listed}"
ref_ids = {id(p._p) for p in refs}

def is_sup(r):
    rp = r.find(qn("w:rPr")); va = rp.find(qn("w:vertAlign")) if rp is not None else None
    return va is not None and va.get(qn("w:val")) == "superscript"

cites = []                                   # (w:t element, [numbers]) in document order
for p in si.paragraphs:
    if id(p._p) in ref_ids: break
    seen = ""
    for t in p._p.iter(qn("w:t")):
        txt = t.text or ""
        if is_sup(t.getparent()) and re.fullmatch(r"\d+(?:,\d+)*", txt.strip()) and not re.search(r"(δ|σw|\bm|\bs|\bkm)$", seen):
            cites.append((t, [int(x) for x in txt.strip().split(",")]))
        seen += txt
order = []
for _, ns in cites:
    for n in ns:
        if n not in order: order.append(n)
assert sorted(order) == listed, f"cited {sorted(order)} vs listed {listed}"
new = {old: k + 1 for k, old in enumerate(order)}
changed = 0
for t, ns in cites:
    s = ",".join(str(new[n]) for n in ns)
    if s != t.text.strip():
        t.text = s; changed += 1
        r = t.getparent(); rp = r.find(qn("w:rPr"))
        c = rp.find(qn("w:color"))
        if c is None: c = rp.makeelement(qn("w:color"), {}); rp.append(c)
        c.set(qn("w:val"), "C00000")
# reorder and renumber the list
anchor = refs[0]._p.getprevious(); parent = refs[0]._p.getparent()
for p in refs: parent.remove(p._p)
for p in sorted(refs, key=lambda p: new[int(REFP.match(p.text.strip()).group(1))]):
    old = int(REFP.match(p.text.strip()).group(1))
    for t in p._p.iter(qn("w:t")):
        if (t.text or "").startswith(f"({old})"):
            t.text = f"({new[old]})" + t.text[len(f"({old})"):]; break
    else:
        sys.exit(f"reference ({old}): number not found in a single text run")
    anchor.addnext(p._p); anchor = p._p
si.save(a.si_out)
moved = {o: n for o, n in new.items() if o != n}
print(f"SI references renumbered by first citation ({changed} citation(s) changed): " +
      ", ".join(f"{o}->{n}" for o, n in sorted(moved.items())))
