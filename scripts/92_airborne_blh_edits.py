"""92_airborne_blh_edits.py -- the mixing height measured along the Twin Otter track by the CSL
airborne Doppler lidar (MD02; Brewer and Baidar, R1 release) replaces the ground lidar at the DSRC
in Boulder (S. Baidar, email of 9 Oct 2026: under strong westerlies the DSRC variance picks up
turbulence aloft that is not connected to the surface). Applied to the 9 Oct manuscript and SI
(output of script 89).

Changes: section 2.4 (source, coverage, transects inside the boundary layer), 3.3 and 3.4 (QC
filter result), 4.3 (hotspot crossing rates), limitations; SI S1.1 (mixing-height definition),
S1.2 (curtain screen), S3 (QC filter, Table S3 with a ground-lidar column), S11 and Tables S1,
S2, S10; Figure S17. Numbers are read from the outputs of scripts 13, 22, 33, 59, 90 and 91 and
asserted. All inserted text is red.

usage: python3 scripts/92_airborne_blh_edits.py <main.docx> <si.docx> <main_out> <si_out>
         --outputs <OUT_DIR> --figures <OUT_DIR>/figures
"""
import sys, csv, re, argparse, statistics as st
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
ap.add_argument("--outputs", required=True); ap.add_argument("--figures", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); FIG = Path(a.figures)
rd = lambda f: list(csv.DictReader(open(OUT / f)))
f2 = lambda v, d=2: "" if v in ("NA", "", None) else f"{float(v):.{d}f}"

# ---- numbers -------------------------------------------------------------------------------
T1 = rd("table1.csv"); FL = [r["Flight"] for r in T1 if r["Fossil_pct_CI"] != "NA"]; assert len(FL) == 9
CL = {r["flight"]: r for r in rd("closeloop_diagnostic.csv")}
air9 = [f for f in FL if CL[f]["blh_src"].startswith("airborne")]
assert sorted(set(FL) - set(air9)) == ["20240703_R0_L2", "20240712_R0_L1"]
b9 = [int(CL[f]["blh_m"]) for f in air9]; assert (min(b9), max(b9)) == (630, 2661), b9
assert (CL["20240703_R0_L2"]["blh_m"], CL["20240712_R0_L1"]["blh_m"]) == ("3352", "2752")
both = [f for f in CL if CL[f]["blh_src"].startswith("airborne") and CL[f]["blh_ground_m"] not in ("", "NA")]
diff = [int(CL[f]["blh_ground_m"]) - int(CL[f]["blh_m"]) for f in both]
assert sorted(f for f, x in zip(both, diff) if x < 0) == ['20240710_R0_L1', '20240713_R0_L1', '20240713_R0_L2']
n_both, n_higher, dmax, dmed = len(both), sum(x > 0 for x in diff), max(diff), st.median(diff)
assert (n_both, n_higher, round(dmax / 1000, 1), round(dmed / 1000, 2)) == (13, 10, 2.9, 0.75), (n_both, n_higher, dmax, dmed)
TB = [r for r in rd("transects_vs_blh.csv") if r["flight"] in FL and r["blh_airborne_m"] not in ("", "NA")]
lo = [r for r in TB if int(r["agl_m"]) < 1000]; hi = [r for r in TB if int(r["agl_m"]) >= 1500]
assert (sum(r["inside_bl"] == "TRUE" for r in lo), len(lo)) == (34, 34)
assert (sum(r["inside_bl"] == "FALSE" for r in hi), len(hi)) == (7, 10)
g_in = sum(int(r["gated_inside_bl"]) for r in TB); g_all = sum(int(r["gated_with_blh"]) for r in TB)
assert round(100 * g_in / g_all) == 97
QC = {r["flight"]: r for r in rd("qc_robustness.csv")}
assert [int(QC[f]["fossil_QC"]) for f in FL] == [24, 0, 8, 0, 34, 7, 21, 36, 34]
assert [int(QC[f]["fossil_QC_ground"]) for f in FL] == [24, 0, 8, 0, 34, 9, 21, 27, 43]
assert st.median(int(QC[f]["fossil_QC"]) for f in FL) == 21 and st.median(int(QC[f]["fossil_QC_ground"]) for f in FL) == 21
sh = [float(QC[f]["share_above_blh"]) for f in air9]; assert (round(100 * min(sh)), round(100 * max(sh))) == (2, 48), sh
CS = rd("denver_curtain_screen.csv")
cnt = {k: sum(r[k] == "TRUE" for r in CS) for k in ("city_upwind", "width_ok", "vertical_ok", "wind_ok")}
assert cnt == {"city_upwind": 11, "width_ok": 9, "vertical_ok": 5, "wind_ok": 14} and max(int(r["conditions_met"]) for r in CS) == 2
s13 = [r for r in CS if r["flight"] == "20240713_R0_L2" and r["axis"] == "NS" and r["offset_km_from_box_center"] == "-12.5"][0]
assert (s13["blh_m"], s13["highest_agl_m"]) == ("2661", "2461")
s09 = [r for r in CS if r["flight"] == "20240709_R0" and r["axis"] == "EW"][0]; assert s09["blh_m"] == "1826"
ENC, seen = [], set()
for r in rd("v2_pointsource_encounters.csv"):
    k = (r["flight"], r["leg_id"], r["hotspot_id"])
    if r["hotspot_in_fetch"] == "TRUE" and k not in seen: seen.add(k); ENC.append(r)
assert len(ENC) == 15 and len({r["flight"] for r in ENC}) == 7
Q = [float(r["Q_t_hr"]) for r in ENC]; QB = [float(r["Q_bio_t_hr"]) for r in ENC if r["Q_bio_t_hr"] != "NA"]
RB = sorted(float(r["ratio_bio_over_v2waste"]) for r in ENC if r["ratio_bio_over_v2waste"] not in ("NA", "Inf"))
RT = sorted(float(r["ratio_total_over_v2total"]) for r in ENC)
SUM = {r["hotspot_id"]: r for r in rd("v2_pointsource_summary.csv")}
d1, d3 = SUM["1"], SUM["3"]
assert (d1["Q_bio_median"], d1["Q_total_min"], d1["Q_total_max"], d1["ratio_bio_over_v2waste_median"]) == ("0.06", "0.06", "0.84", "0.01")
assert (d3["Q_bio_median"], d3["Q_total_min"], d3["Q_total_max"], d3["ratio_bio_over_v2waste_median"]) == ("0.5", "0.39", "3.68", "1.3")
ww = [r for r in ENC if r["hotspot_id"] == "3" and float(r["ws_ms"]) >= 2]
ww_b = st.median(float(r["Q_bio_t_hr"]) for r in ww if r["Q_bio_t_hr"] != "NA")
ww_r = st.median(float(r["ratio_bio_over_v2waste"]) for r in ww if r["ratio_bio_over_v2waste"] not in ("NA", "Inf"))
FS = {r["flight"]: r for r in rd("flight_summary.csv")}

OPS = {"main": [], "si": []}
def R(doc, old, new, count=1): OPS[doc].append(("R", old, new, count))
def RW(doc, anchor, text): OPS[doc].append(("RW", anchor, text))
def TAB(doc, ti, rows, header=None, widths=None): OPS[doc].append(("TAB", ti, rows, header, widths))
M, S = "main", "si"

# ================================ main text ================================
RW(M, "Mixing heights were obtained from a NOAA Chemical Sciences Laboratory (CSL) Doppler-lidar site",
   "Mixing heights come from the NOAA Chemical Sciences Laboratory (CSL) airborne Doppler lidar (MD02) on the Twin Otter, which retrieved the boundary-layer height below the aircraft about every 30 s along the flight track "
   "from the lidar's fit variance, range-corrected intensity gradient, wind-speed gradient and wind shear (final R1 release; 17 of the 22 flights).[[^11]] "
   "The ground-based CSL Doppler lidar at the Boulder DSRC (39.99°N, 105.26°W), which reports vertical-velocity variance, supplies the mixing height for the five flights without an airborne retrieval and is kept as a sensitivity case. "
   "{{c17}}The winds in the wind-direction, encounter and transect analyses are the aircraft's own measurements (wind roses for the flights and for June and July 2024 are in Figure S13).{{/c17}} "
   "The High-Resolution Rapid Refresh (HRRR) model drives the Stochastic Time-Inverted Lagrangian Transport (STILT) footprints (section S9).")
RW(M, "Each flight's mixing height is the median of the lidar's retrievals",
   "Each flight's mixing height is the median of the airborne retrievals inside the analysis box during that flight (the whole track where fewer than three lie in the box). "
   "It spans 0.63 to 2.66 km on the seven of the {{c18}}nine flights with estimable urban slopes{{/c18}} that have an airborne retrieval; for the other two, 3 July L2 and 12 July L1, the ground lidar gives 3.35 and 2.75 km (Table S1). "
   "On those seven flights all 34 urban transects flown below 1 km above ground lay inside the boundary layer measured along the track, as did 97% of the samples entering the ethane fits, whereas 7 of the 10 transects flown above 1.5 km lay above it. "
   f"The ground lidar is a single site about 40 km north-west of the urban core; under strong westerly flow over the mountains its vertical-velocity variance can register turbulence aloft that is not connected to the surface, and its mixing height exceeds the airborne value on {n_higher} of the {n_both} flights covered by both, by up to 2.9 km (median difference 0.75 km). "
   "The mixing height is used to document the daytime boundary layer, in the sensitivity analysis that applies an in-boundary-layer filter (section S3), in the vertical-coverage test for curtains (section S1.2), and to scale the single-transect rates downwind of inventory hotspots (section S11).")
R(M, "leaves the campaign median unchanged at 21%, although it shifts three flights by 8 to 17 percentage points (section S3).",
     "leaves the campaign median unchanged at 21% and changes only one flight by more than 1 percentage point (section S3).")
R(M, "We did not impose the filters in the primary analysis because our boundary-layer heights come from a single Doppler lidar about 40 km from central Denver and may not represent mixing conditions over the city, particularly on days with a shallow boundary layer. We therefore fit all urban-transect samples; with the filters applied, the campaign median is unchanged (21%), although three flights change by 8 to 17 percentage points (Table S3).",
     "We did not impose the filters in the primary analysis, so that the fits keep every urban transect, including those flown above the mixed layer. We therefore fit all urban-transect samples; with the filters applied, using the boundary-layer height measured along the track by the airborne lidar, the campaign median is unchanged (21%) and only 8 July L1 changes by more than 1 percentage point (17 to 0%; Table S3).")
RW(M, "Boundary-layer height is constrained at only one location.",
   "Boundary-layer height is measured along the flight track for most flights. The airborne Doppler lidar resolves the boundary layer over the air actually sampled on 17 of the 22 flights, but not on the other five, two of which (3 July L2 and 12 July L1) are in Table 1; "
   "for those the mixing height comes from the single ground lidar about 40 km from the urban core, whose values exceed the airborne ones on most flights covered by both. "
   "We do not impose an in-boundary-layer filter on the primary enhancement-ratio analysis; applying it leaves the campaign median unchanged with either mixing height (section S3).")
RW(M, "The three waste-source cells contain 12.0 of v2.0beta",
   "The three waste-source cells contain 12.0 of v2.0beta's 12.2 t CH4/h of box-wide waste emissions. "
   "We identified 15 aircraft crossings on seven flights within 12 km and 25° of the downwind direction from two of these cells (section S11, Figure S17, Table S10). "
   "The three crossings downwind of the DADS landfill (two flights) give a median biogenic methane rate of 0.06 t CH4/h- about 1% of the inventory's waste emissions in the upwind area, dominated by the landfill cell's 6.07 t CH4/h. "
   "Twelve crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery complex (seven flights) give a median biogenic rate of 0.50 t CH4/h, a median 1.3 times the v2.0beta waste emissions in their upwind fetch, "
   "and total methane rates across these crossings range from 0.39 to 3.68 t CH4/h. "
   "These estimates scale directly with the mixing height, measured by the airborne lidar at each crossing, and assume uniform vertical mixing, so they provide approximate constraints rather than reliable facility emission rates.")

# ================================ SI ================================
RW(S, "The signed wind crossing, w·nout, times the air density",
   "The signed wind crossing, w·nout, times the air density and methane enhancement is summed around the loop by trapezoidal integration along the flown path, and the altitude levels are integrated vertically from the surface to the mixing height. "
   "The mixing height is the boundary-layer height measured below the aircraft by the NOAA CSL airborne Doppler lidar (MD02) on the Twin Otter, about every 30 s along the track (final R1 release; retrieved from the lidar's fit variance, range-corrected intensity gradient, wind-speed gradient and wind shear). "
   "A per-flight value is the median of the retrievals inside the analysis box, or along the whole track where fewer than three lie in the box; the in-boundary-layer filter of section S3 uses the retrieval at each sample, interpolated along the track. "
   "Five flights have no airborne retrieval (28 June, 3 July L2, 11 July L2, and 12 July L1 and L2). For them the mixing height comes from the ground Doppler lidar at the Boulder DSRC: the highest range gate, contiguous from the surface, "
   "at which the noise-corrected vertical-velocity variance (the variable wVar of the lidar's velStats files, on 20 min profiles with 50 m range gates starting 102 m above the lidar) is at least 0.1 m² s⁻², tolerating one non-turbulent gate, as the median of the profiles within the flight's UTC window. "
   "Across the nine flights with estimable urban slopes the mixing height spans 0.63 to 3.35 km. "
   f"The ground lidar exceeds the airborne value on {n_higher} of the {n_both} flights covered by both, by up to 2.9 km (median 0.75 km): under strong westerly flow over the mountains its variance registers turbulence aloft that is not connected to the surface. "
   "It falls below the airborne value on 10 July L1 and both 13 July flights (for example 0.35 against 0.63 km on 13 July L1). Table S1 lists both values; the ground values are used here only as a sensitivity case.")
R(S, "two provided adequate vertical coverage, and 14 had adequate and sufficiently steady transport; no screen passed more than three criteria.",
     "five provided adequate vertical coverage, and 14 had adequate and sufficiently steady transport; no screen passed more than two criteria (with the ground-lidar mixing height one screen passes three).")
RW(S, "The strongest candidates illustrate the limiting conditions.",
   "The strongest candidates illustrate the limiting conditions. The most complete wall, north-south about 13 km west of the box center on the second 13 July flight, stacked 11 transects on three levels (231 to 2461 m above ground), "
   "spanned 1.18 times the city width and carried 3.1 m/s of through-screen flow with a constancy of 0.86; but it bisects the metropolitan area, so the city is not entirely upwind of it, and its highest transect lies below the 2661 m mixing height measured by the airborne lidar. "
   "Screens fully beyond the city met at most two conditions: on 9 July, an east-west wall 35 km south of the box center had the city upwind and steady transport (4.6 m/s, constancy 0.94), but a single altitude level (272 to 421 m above ground, far below the 1826 m mixing height) spanning 41% of the city width, "
   "and on the second 13 July flight a similar wall 28 km south had one level at 232 to 297 m.")
RW(S, "Candidate curtain screens near the Denver metropolitan area, one row per screen",
   "**Table S2:** Candidate curtain screens near the Denver metropolitan area, one row per screen, ordered by the number of conditions met. A screen is a wall of two or more constant-altitude transects sharing an axis and a common offset. "
   "Wall gives the axis and the offset of the wall from the center of the analysis box, positive east for north-south walls and positive north for east-west walls. Width is the wall's cross-wind span as a fraction of the box's. "
   "U⊥ is the mean wind component through the wall, and constancy is the ratio of the vector-mean to the scalar-mean wind speed; both come from the screen's own samples below 2 km above ground. None of the 48 screens meets all four conditions, so no curtain flux is computed. "
   "{{c1}}Mixing heights are from the NOAA CSL airborne Doppler lidar on the Twin Otter (median of the in-box retrievals during the flight; section S1.1), or, for flights without an airborne retrieval, from the ground Doppler lidar at the Boulder DSRC "
   "(the highest 50 m range gate, contiguous from the surface, at which the noise-corrected vertical-velocity variance wVar is at least 0.1 m² s⁻²).{{/c1}}")
RW(S, "For the AMMBEC urban flights, these filters are mild",
   "For the AMMBEC urban flights, these filters are mild: all nine analysis flights were flown during the day (local sampling windows within 09:36 to 18:00), and only 4 to 12% of all in-box samples (fewer on the urban transects used in the fits) fall below 200 m AGL. "
   "The in-boundary-layer filter uses the boundary-layer height measured by the airborne lidar at each sample (the ground lidar for 3 July L2 and 12 July L1, which have no airborne retrieval). "
   "It removes 2 to 48% of the urban-transect samples on the seven flights with airborne retrievals, mostly on the transects flown above the mixed layer, but few of the samples that enter the ethane fits: 97% of those lie inside the boundary layer (section S1.1). "
   "Eight of the nine flights change by at most 1 percentage point (Table S3); 8 July L1 falls from 17 to 0% under either mixing height, because its four weakly correlated transects (r = -0.35 to 0.23) leave the fit sensitive to which samples are kept. "
   "The campaign median is unchanged at 21% and no flight exceeds 50%, so the biogenic-dominated source-mix conclusion is robust to data selection. "
   "With the ground-lidar mixing height instead, the median is also unchanged, but the two 13 July flights move to 27 and 43%: the ground lidar's 0.35 km on 13 July L1 removes most of that flight's samples, and its 1.75 km on 13 July L2 removes transects that lie inside the 2.7 km boundary layer measured by the airborne lidar.")
R(S, "the campaign median did not change, although three flights changed by 8 to 17 percentage points (Table S3). We did not impose the filters in the primary analysis because our boundary-layer heights come from a single Doppler lidar approximately 40 km from central Denver and may not represent mixing conditions over the city, particularly on days with a shallow boundary layer.",
     "the campaign median did not change, and only one flight changed by more than 1 percentage point (Table S3). We did not impose the filters in the primary analysis, so that the fits keep every urban transect, including those flown above the mixed layer.")
RW(S, "Per-flight fossil fraction from all urban-transect samples (no altitude or time filter)",
   "**Table S3:** Per-flight fossil fraction from all urban-transect samples (no altitude or time filter) and from the QC-filtered subset, with the absolute difference between them (percentage points). "
   "The QC subset keeps only samples inside the boundary layer measured by the airborne lidar at each sample (the ground lidar for 3 July L2 and 12 July L1), above 200 m AGL and during daytime hours (10:00 to 17:00 MDT); "
   "the last column repeats the filter with the ground-lidar mixing height. The campaign median is 21% in all three columns.")
R(S, "multiplied by the lidar mixing height over that flight's urban transects (uniform mixing to that height; this median covers the period of the crossings and can differ from the whole-flight value in Table S1)",
     "multiplied by the mixing height measured by the airborne lidar within 5 minutes of the crossing (uniform mixing to that height; the flight's in-box value where the lidar has no retrieval that close, and the ground lidar for flights without an airborne retrieval)")
R(S, "The rate scales with the mixing height, which comes from a lidar 40 km away, and the fetch", "The rate scales with the mixing height, and the fetch")
RW(S, "There are 15 crossings on seven flights at two of the hotspot sites",
   f"There are 15 crossings on seven flights at two of the hotspot sites (Table S10, Figure S17). Their total rates are {min(Q):.2f} to {max(Q):.2f} t/h and their biogenic parts 0 to {max(QB):.2f} t/h. "
   f"Relative to the v2.0beta waste emissions in the fetch, the aircraft biogenic rate is {st.median(RB):.2f} in the median (range 0 to {RB[-1]:.1f}); relative to the v2.0beta total in the fetch, the aircraft total is {st.median(RT):.2f} in the median. "
   "Pooling across sites mixes a landfill and a wastewater cell, so the per-site values below are the informative ones. The fetch of each crossing is oriented by the vector mean of the aircraft wind along the crossed segment.")
RW(S, "By site: the three crossings downwind of the DADS landfill",
   "By site: the three crossings downwind of the DADS landfill (v2.0beta waste 6.07 t/h in that cell; 8 and 10 July L1) give a median biogenic rate of 0.06 t/h (total 0.06 to 0.84 t/h), 0.01 times the v2.0beta waste emissions in the upwind fetch; "
   "the largest, on 8 July, carries an ethane signal (slope 0.13) that assigns it entirely to fossil methane. "
   "The twelve crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery complex (v2.0beta waste 0.39 t/h in that cell; seven flights) give a median biogenic rate of 0.50 t/h (total 0.39 to 3.68 t/h), 1.30 times the v2.0beta waste emissions in the upwind fetch. "
   f"Excluding the two crossings there flown in winds below 2 m s−1 (3 July L1 and 10 July L1) gives {ww_b:.2f} t/h and {ww_r:.2f} times. "
   "With the ground-lidar mixing height instead, the medians are 0.11 t/h at the landfill and 0.85 t/h (2.14 times) at the wastewater complex, because the ground lidar's mixing height exceeds the airborne value on most of these flights. "
   "Which facility the third cell represents cannot be determined from the gridded product: it lies immediately north of the Metro Water Recovery plant, and the cell that contains the plant itself carries no waste methane in v2.0beta, so the cell may be the plant with a displaced location or another waste facility. "
   "Either way, the aircraft find far less methane than v2.0beta downwind of the landfill they could test, and more biogenic methane around the wastewater plant than v2.0beta places in that cell and the plant's own cell together: "
   "the v2.0beta waste total is high, and its split between the landfills and the wastewater plant appears to be the reverse of what the aircraft, and Carbon Mapper's persistence-weighted landfill rates (Table S9), indicate.")
R(S, "BLH is the lidar mixing height (median over the flight window).",
     "BLH gives the mixing height from the airborne Doppler lidar (median of the in-box retrievals) and, for comparison, from the ground lidar at the Boulder DSRC (median over the flight window); a dash marks the five flights without an airborne retrieval, for which the ground value is used.")

# ---- tables ----
lab = lambda f: (lambda m: f"{m.group(1)}-{m.group(2)}-{m.group(3)}" + (f" ({m.group(4)})" if m.group(4) else ""))(re.match(r"(\d{4})(\d\d)(\d\d)_R0_?(L\d)?", f))
def blh_cell(f):
    s = FS[f]; air = s["blh_src"].startswith("airborne")
    return (s["blh_m"] if air else "-") + " / " + (s["blh_ground_m"] if s["blh_ground_m"] not in ("", "NA") else "-")
S1rows = []
for r in T1:
    f = r["Flight"]; s = FS[f]
    S1rows.append([lab(f), s["flight_utc"], s["inbox_utc"], s["n_box_transects"], f"{int(s['inbox_transect_s']):,}", s["agl_range"],
                   s["bl_wind"].replace(" m/s from ", " m/s, "), s["above_wind"].replace(" m/s from ", " m/s, "), blh_cell(f), s["n_alt_levels"], s["encircles"]])
TAB(S, 0, S1rows, header=[None] * 8 + ["BLH airborne / ground (m)", None, None], widths=[1.2, 1.0, 1.0, 0.85, 0.8, 0.9, 1.2, 1.2, 0.95, 0.6, 0.75])
def rnd(x): return int(abs(x) + 0.5) * (1 if x >= 0 else -1)
S2rows = []
for r in sorted(CS, key=lambda r: (-int(r["conditions_met"]), r["flight"])):
    S2rows.append([lab(r["flight"]), f"{'N-S' if r['axis'] == 'NS' else 'E-W'}, {rnd(float(r['offset_km_from_box_center']))} km", f"{r['n_legs']} / {r['n_alt_levels']}",
                   f"{r['lowest_agl_m']}-{r['highest_agl_m']}", r["blh_m"], f"{float(r['width_frac']):.2f}", f"{float(r['U_perp_ms']):.1f} / {float(r['constancy']):.2f}", f"{r['conditions_met']} of 4"])
TAB(S, 1, S2rows)
TAB(S, 2, [[lab(f), QC[f]["fossil_full"], QC[f]["fossil_QC"], QC[f]["delta"], QC[f]["fossil_QC_ground"]] for f in FL],
    header=["Flight", "Fossil % (full)", "Fossil % (QC, airborne BLH)", "|Δ| (pts)", "Fossil % (QC, ground BLH)"])
S10rows = []
for r in sorted(ENC, key=lambda r: (r["hotspot_id"], r["flight"], int(r["leg_id"]))):
    S10rows.append([lab(r["flight"]), r["leg_id"], r["hotspot_id"], f2(r["dist_km"], 1), r["agl_m"], r["blh_m"], r["dch4_max_ppb"], f2(r["Q_t_hr"]),
                    f2(r["fossil_frac"]), f2(r["Q_bio_t_hr"]), f2(r["v2_fetch_waste_t_hr"]), f2(r["v2_fetch_total_t_hr"]), f2(r["v1_fetch_total_t_hr"])])
TAB(S, 9, S10rows)

# ================================ apply ================================
def run_doc(src, dst, which, label):
    d = docx.Document(src); n = 0
    for op in OPS[which]:
        if op[0] == "R": O.do_replace(d, op[1], op[2], op[3], label); n += 1
        elif op[0] == "RW": O.do_rewrite(d, op[1], op[2], label); n += 1
        elif op[0] == "TAB": c = O.do_table(d, op[1], op[2], op[3], label, op[4]); print(f"{label}: table {op[1]} rebuilt ({len(op[2])} rows, {c} cells changed)")
    allp = L.all_paragraphs(d)
    m = L.subscript_chem(allp); e = L.exponent_fix(L.all_paragraphs(d)); c = L.cite_fix(L.all_paragraphs(d)); g = O.gra2pes_superscript(d)
    d.save(dst); print(f"{label}: {n} edit(s) -> {dst} ({m} subscripts, {e} exponents, {c} citations)")
run_doc(a.main_in, a.main_out, "main", "main text")
run_doc(a.si_in, a.si_out, "si", "SI")
O.swap_figures(a.si_out, {"image4.png": FIG / "FigS14_v2_pointsource.png"}, set(), "SI")

# ---- checks ----
import zipfile
for path, bads, ncom in ((a.main_out, ("40 km from central Denver", "constrained at only one location", "0.35 to 3.35 km", "0.85 t CH4/h", "three flights change by 8 to 17"), 33),
                         (a.si_out, ("comes from a lidar 40 km away", "no screen passed more than three", "Six of the nine flights", "three flights changed by 8 to 17", "Mixing heights are retrievals from the NOAA CSL Doppler lidar near Boulder"), 2)):
    d = docx.Document(path); t = "\n".join(O.vis_text(p) for p in d.paragraphs)
    for b in bads: assert b not in t, f"{path}: {b!r} remains"
    x = zipfile.ZipFile(path).read("word/document.xml").decode()
    assert len(set(re.findall(r'<w:commentReference w:id="(\d+)"', x))) == ncom
ds = docx.Document(a.si_out)
assert len(ds.tables[2].columns) == 5 and ds.tables[0].rows[0].cells[8].text == "BLH airborne / ground (m)"
print("checks passed")
