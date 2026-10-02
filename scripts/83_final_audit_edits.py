"""83_final_audit_edits.py -- fixes from the independent audit of 2 Oct 2026 (main text, SI, figures, code).

Corrections of statements made by script 82 the same day:
  * SI S2.1: the CDPHE trend slopes were rounded twice (0.0055 -> 0.006); script 06 now writes
    unrounded slopes and the range is set from them (0.005 to 0.008).
  * Table S10 caption: the DADS distances differ because they average different samples, not
    because the cell centre is 2 km from the landfill (it is 0.4 km from the Carbon Mapper source;
    the 2 km is to the geocoded street address in Table S9).
Numbers and statements checked against the outputs:
  * NEI: the 2.4-fold increase is about 1.3 from the larger footprint and about 1.9 from the NEI
    giving more CO than GRA2PES over the same counties (the box-scaled NEI gives 8.7-20.1 t/h).
  * Basin ratios 0.031-0.036 behind the 0.063-0.073 source ratio; weekday-to-weekday GRA2PES CO.
  * Discussion: 27-48% is below or within the lower part of the 40-88% range (48% lies inside
    Baltimore-Washington's 40-60%); 29% is at the adopted, not the lowest published, calibration.
  * The STILT result of SI S9 is now reflected in 4.3 (the transport-weighted inventory slope is
    0.041 against an observed 0.034).
  * Wind/fossil association r = 0.35 is weak and not significant (p computed here).
  * Smaller items: 2.4 (lidar role, winds from the aircraft, PUMAS on 29 June), 2.2 (cadence),
    2.5 (DADS spelled out; 0.1 t/h), 3.1, 3.2 (precision source, York et al., subject), 4.1, 4.2,
    Discussion (LA oil and gas, crossings, Section S7, EPA total below every estimate), the
    Limitations endmember sentence, figure captions 1-4 and S2-S4, S6, S8, S9, S13, S16, S17.
  * SI: code names and script numbers removed from the prose; S3 duplicated paragraph removed;
    S3, S4, S4.1, S5, S8, S9, S12 wording; Table S4 ratios recomputed from the PSCo table;
    Table S9 Carbon Mapper rates in t/h; Table S10 flight labels.

usage: python3 scripts/83_final_audit_edits.py <main.docx> <si.docx> <main_out> <si_out> --outputs <OUT_DIR>
"""
import sys, csv, re, math, argparse, statistics as st
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
import docx_edit_lib as L
from docx_edit_lib import R, runs_text, all_paragraphs, replace_in_paragraph

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); L.OUTDIR = str(OUT)
REPO = Path(__file__).resolve().parent.parent
rd = lambda f: list(csv.DictReader(open(OUT / f)))
hu = lambda x, q: str(Decimal(repr(float(x))).quantize(Decimal(q), rounding=ROUND_HALF_UP))
def q7(v, p):                                    # R's default (type 7) quantile
    v = sorted(v); h = (len(v) - 1) * p; lo = math.floor(h); return v[lo] + (h - lo) * (v[min(lo + 1, len(v) - 1)] - v[lo])
def t_p_two_sided(r, n):                          # p of a Pearson r, by Simpson integration of Student t
    df = n - 2; t = abs(r) * math.sqrt(df / (1 - r * r))
    c = math.exp(math.lgamma((df + 1) / 2) - math.lgamma(df / 2)) / math.sqrt(df * math.pi)
    f = lambda x: c * (1 + x * x / df) ** (-(df + 1) / 2)
    N = 2000; h = t / N; s = f(0) + f(t) + sum((4 if k % 2 else 2) * f(k * h) for k in range(1, N))
    return 1 - 2 * s * h / 3

# ---- numbers ----------------------------------------------------------------------------------
TR = rd("mobile_trend_fits.csv")
sl = [abs(float(r["slope_ppmv_per_yr"])) for r in TR]; pv = [float(r["p_value"]) for r in TR]
assert all(len(r["slope_ppmv_per_yr"].split(".")[-1]) > 4 for r in TR), "rerun scripts/06 (unrounded slopes)"
slope_txt = f"{hu(min(sl), '0.001')} to {hu(max(sl), '0.001')}"; p_txt = f"{hu(min(pv), '0.01')} to {hu(max(pv), '0.01')}"
CB = rd("gra2pes_CO_202307_v1.1_county_box.csv")[0]
box_frac = float(CB["box_over_county"]); nei7 = float(CB["NEI_7county_Gg_yr"]); gra_box = float(CB["GRA2PES_box_Gg_yr"])
gra_county = gra_box / box_frac; foot = 1 / box_frac; mag = nei7 / gra_county
PV = open(OUT / "paper_values.json").read(); pvj = lambda k: float(re.search(rf'"{k}"\s*:\s*([0-9.]+)', PV).group(1))
nei_box_lo, nei_box_hi = pvj("nei_box_lo"), pvj("nei_box_hi")
BS = rd("basin_summary.csv")[0]
b4 = float(BS["median_min3_legs"])
C11 = rd("gra2pes_cells_v1.1_summary.csv")[0]; C2 = rd("gra2pes_cells_v2.0beta_summary.csv")[0]
assert C11["daytype"] == C2["daytype"] == "weekdy"
co11, co2 = float(C11["co_Gg_yr"]), float(C2["co_Gg_yr"])
WF = [r for r in rd("wind_fossil.csv") if r["fossil_pct"] not in ("", "NA")]
xs = [float(r["pct_from_NE"]) for r in WF]; ys = [float(r["fossil_pct"]) for r in WF]
mx, my = st.mean(xs), st.mean(ys)
r_w = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / math.sqrt(sum((x - mx) ** 2 for x in xs) * sum((y - my) ** 2 for y in ys))
p_w = t_p_two_sided(r_w, len(xs))
NE = {r["flight"]: r["pct_from_NE"] for r in rd("wind_fossil.csv")}
PL = [r for r in rd("arc_ground_plumes.csv") if r["region"] == "basin" and r["class"] == "fossil-type"]
bs = [float(r["slope"]) for r in PL]; iqr_txt = f"{hu(q7(bs, .25), '0.001')} to {hu(q7(bs, .75), '0.001')}"
SL = [float(r["slope"]) for r in rd("stilt_exposure_legs.csv") if r["slope"] not in ("", "NA")]
stilt_obs = hu(st.median(SL), "0.001")
HS = {r["hotspot_id"]: r for r in rd("v2_pointsource_hotspots.csv")}
CMS = [r for r in rd("carbonmapper_denver_sources.csv") if r["gas"] == "CH4"
       and abs(float(r["geom_lat"]) - 39.662) < 0.01 and abs(float(r["geom_lon"]) + 104.690) < 0.01]
assert len(CMS) == 1
def km(la1, lo1, la2, lo2): return math.hypot((lo2 - lo1) * 111.32 * math.cos(math.radians((la1 + la2) / 2)), (la2 - la1) * 110.57)
cell_cm_km = km(float(HS["1"]["lat"]), float(HS["1"]["lon"]), float(CMS[0]["geom_lat"]), float(CMS[0]["geom_lon"]))
CE = [r for r in rd("carbonmapper_encounters.csv") if r["criterion"] == "strict" and "DADS" in r["source"]][0]
EN = [r for r in rd("v2_pointsource_encounters.csv") if r["hotspot_id"] == "1"][0]
LS = {int(r["lag_s"]): r for r in rd("lag_sensitivity_summary.csv")}
n_basin0 = int(LS[0]["n_basin"])
BE = rd("basin_ethane_methane.csv"); neg = sorted(float(r["basin_ratio"]) for r in BE if float(r["basin_ratio"]) < -0.005)
fmt_neg = lambda v, d: f"{v:.{d}f}".replace("-", "\u2212")
SEP = int(re.search(r"CROSS_SEP_MIN_M\s*<-\s*(\d+)", open(REPO / "scripts/24_facility_separation_scan.R").read()).group(1))
NPM = int(re.search(r"N_PLUME_MIN\s*<-\s*(\d+)", open(REPO / "scripts/24_facility_separation_scan.R").read()).group(1))
print(f"trend |slope| {slope_txt}, p {p_txt}; NEI footprint x{foot:.2f}, magnitude x{mag:.2f}, box {nei_box_lo}-{nei_box_hi}; "
      f"basin 4-flight {b4:.3f}; CO weekday v2 {co2:.1f} v1.1 {co11:.1f}; wind r {r_w:.2f} p {p_w:.2f}; ARC basin IQR {iqr_txt}; "
      f"STILT obs median {stilt_obs}; DADS cell-CM {cell_cm_km:.2f} km; negative basin {neg}; sep {SEP} m, {NPM} pts")
assert abs(r_w - 0.35) < 0.005

# ---- main text ----------------------------------------------------------------------------------
R("(1.7 and 2.8 t CH4/h), by about 2.7 to 4.5-fold", "(1.7 and 2.8 t CH4/h) by about 2.7 to 4.5-fold")
R("and the single screen-passing CH4:CO2 flight 4.1 t CH4/h;", "and the single screen-passing CH4:CO2 flight gave 4.1 t CH4/h;")
R("The surveys measured CH4 in ppmv with differential GPS at about 1 Hz, from 2023 to 2025.",
  "The surveys measured CH4 in ppmv with differential GPS, delivered on a 1 s grid (the analyzer acquires about every 5 s), from 2023 to 2025.")
R("Mixing height and winds aloft were used from a NOAA CSL Doppler-lidar site near Boulder (39.99°N, 105.26°W), which reports "
  "vertical-velocity variance and wind profiles.",
  "Mixing height was taken from a NOAA CSL Doppler-lidar site near Boulder (39.99°N, 105.26°W), which reports vertical-velocity "
  "variance; the winds in the wind-direction, encounter and transect analyses are the aircraft's own measurements, and HRRR "
  "drives the STILT footprints (section S9).")
R("It is used to document the daytime boundary layer and in the sensitivity analysis that applies an in-boundary layer filter (section S3).",
  "It is used to document the daytime boundary layer, in the sensitivity analysis that applies an in-boundary-layer filter "
  "(section S3), in the vertical-coverage test for curtains (section S1.2), and to scale the single-transect rates downwind of "
  "inventory hotspots (section S11). On 29 June, when this lidar has no data, the curtain test uses the campaign's second "
  "Doppler lidar (PUMAS), then in the DJ Basin.")
R("two landfills- Tower Road (18 plumes on 15 dates) and DADS (9 plumes on 6 dates)-",
  "two landfills- Tower Road (18 plumes on 15 dates) and the Denver Arapahoe Disposal Site (DADS; 9 plumes on 6 dates)-")
R("CarbonMapper observations are episodic detections above instrument-specific plume thresholds, roughly 100 kg CH4/h for the satellite observations",
  "Carbon Mapper observations are episodic detections above instrument-specific plume thresholds, roughly 0.1 t CH4/h for the satellite observations")
R("Only the window width w differed between platforms:", "Apart from the ground car's plume criteria (section S2.2), only the window width w differed between platforms:")
R("using the York (2004) bivariate estimator", "using the York et al. (2004) bivariate estimator")
R("C2H6 ±0.2 ppb from the instrument specifications", "C2H6 ±0.2 ppb, the values reported in the data files")
# (two edits, because the citation between them is a hyperlink outside the run text)
R("Second, the DJB basin legs yielded an ambient C2H6:CH4 enhancement ratio of 0.036 mol mol-1. Assuming that oil and gas account for "
  "approximately half of the basin methane enhancement,",
  f"Second, the DJB basin legs yielded ambient C2H6:CH4 enhancement ratios of {b4:.3f} to 0.036 mol mol-1 (the four best-sampled "
  "flights and all flights passing the screen). If oil and gas account for approximately half of the basin methane enhancement,")
R(" implies an oil-and-gas source ratio of approximately 0.063-0.073", " these imply an oil-and-gas source ratio of approximately 0.063-0.073")
R("The abstract and Discussion therefore report the median across", "The abstract and Discussion report the median across")
R("GRA2PES v2.0beta provides a nearly identical box total of 123.4 Gg CO yr-1, within 3% of the v1.1 value on the day-type-weighted "
  "July basis used for the anchor.",
  f"On a common July-weekday basis, GRA2PES v2.0beta gives a box total of {co2:.1f} Gg CO yr-1 against {co11:.1f} for v1.1 "
  f"({abs(co2 / co11 - 1) * 100:.0f}% lower).")
R("eight contained at least three urban legs (Figure 1)", "eight contained at least three urban legs (Figures 1 and S16)")
R("the campaign median ranges from 27 to 48% (sections 3.2, S5).", "the campaign median ranges from 27 to 48% (sections 3.2, S3).")
R("show modest positive slopes of 0.02-0.03, corresponding to fossil fractions of 22-34%",
  "show modest positive slopes of 0.023-0.035, corresponding to fossil fractions of 22-34%")
R("across the four flights that pass the reliability screen (r ≥ 0.7).", "across the four flights that pass the reliability screen (section 3.4).")
R("is moderately associated with the fossil fraction (Pearson r = 0.35, n = 7), although the small number of flights makes this an exploratory relationship",
  f"is weakly and not significantly associated with the fossil fraction (Pearson r = {r_w:.2f}, p = {p_w:.2f}, n = {len(xs)}), so with "
  "this few flights the relationship is exploratory")
R("On 8 and 10 July, the north-easterly exposure was comparable, yet fossil fractions were near zero.",
  f"On 8 and 10 July, {NE['20240708_R0_L1']}% and {NE['20240710_R0_L1']}% of samples came from the north-east, comparable to "
  f"9 July ({NE['20240709_R0']}%), yet fossil fractions were near zero.")
R("so the discrepancy cannot be explained by the aircraft preferentially sampling unusually biogenic parts of the domain.",
  "so the discrepancy cannot be explained by the aircraft preferentially sampling unusually biogenic parts of the domain. "
  "Transport narrows it: convolving HRRR-STILT footprints with the same inventory predicts a median leg slope of 0.041 mol "
  f"mol−1 against {stilt_obs} observed (section S9), because the air reaching the legs also carries waste methane from upwind "
  "cells, including a large one just north of the box. The inventory's box-wide fossil share therefore overstates the "
  "contrast in the air actually sampled, although the transport-weighted inventory is still more fossil than the observations.")
R("That difference reflects a spatial-footprint mismatch rather than a change in Denver’s methane emissions. The NEI total integrates "
  "CO over approximately 11,700 km2, about 4.2 times the area of the analysis box, so applying it to a box-scale enhancement ratio "
  "assigns emissions from a substantially larger region to the sampled domain. The resulting increase is smaller than the fourfold "
  "area difference because CO emissions are spatially concentrated: the urban core within the box contains a disproportionately "
  "large share of the seven-county CO total.",
  "That difference has two parts, neither of which is a change in Denver’s methane emissions. The NEI total integrates CO over "
  "approximately 11,700 km2, about 4.2 times the area of the analysis box, but CO emissions are concentrated in the urban core: "
  f"GRA2PES places {box_frac * 100:.0f}% of the seven-county CO inside the box, so the larger footprint alone accounts for a factor "
  f"of about {foot:.1f}. The rest, a factor of about {mag:.1f}, is a difference in magnitude between the two CO inventories: the NEI "
  f"seven-county total is {nei7:.1f} Gg CO yr−1 against {gra_county:.0f} Gg CO yr−1 in GRA2PES over the same counties, and "
  f"scaling the NEI to the box with the GRA2PES spatial pattern still gives {nei_box_lo:.1f} to {nei_box_hi:.1f} t CH4/h.")
R("roughly four times the analysis box- it over-scales emissions for our domain (Figure 1).",
  "roughly four times the analysis box- it over-scales emissions for our domain (Figure 1), and the NEI also gives about "
  f"{mag:.1f} times more CO than GRA2PES over the same counties (section 5).")
R("We treat these as an upper bound sensitivity case because the county inventory covers a substantially larger area than the aircraft analysis domain.",
  "We treat these as an upper bound sensitivity case because the county inventory covers a substantially larger area than the "
  "aircraft analysis domain and gives more CO than GRA2PES over the same counties (section 5).")
R("Denver’s median fossil fraction is at or below the low end of the 40-88% values",
  "Denver’s median fossil fraction is below or within the lower part of the 40-88% range of values")
R("(27 to 48% across the Denver-relevant endmembers, 29% under the lowest published calibration) lies below all of them at the adopted "
  "calibration and at or below their low end across the Denver-relevant endmembers, although its most fossil-influenced flight",
  "(27 to 48% across the Denver-relevant endmembers, 29% at the adopted 0.102 calibration) lies below all of them at that "
  "calibration and below or within the lower part of the 40-88% range across the Denver-relevant endmembers (its 48% upper "
  "end lies inside the Baltimore-Washington range); its most fossil-influenced flight")
R("and the facility-scale crossings in section 4.3 point in the same direction.",
  "and the crossing downwind of the DADS landfill in section 4.3 points the same way, although near the "
  "wastewater plant the median crossing exceeds the v2.0beta value there.")
R("Los Angeles also contains oil-and-gas production within its urban basin, unlike the Denver analysis box.",
  "Los Angeles also contains substantial oil-and-gas production within its urban basin, whereas the Denver analysis box contains "
  "only a few small oil-and-gas sites near its northern edge.")
R("even though its total emissions are below those inferred on the higher-flux flight days.",
  "even though its total emissions are below every airborne estimate.")
R("Section S7 (Figure S10) maps this spatial separation directly: major waste sources lie within and near the analysis box, whereas "
  "livestock sources are concentrated largely outside it.",
  "Section S7 (Figure S10) shows that the major waste sources lie within and near the analysis box, whereas livestock sources are "
  "concentrated largely outside it; their offset from traffic CO is not mapped.")
R("Both are ambient enhancement ratios, diluted by whatever biogenic methane was co-sampled, and so bound the fossil source "
  "composition from below rather than estimating it:",
  "The production-area value is an ambient enhancement ratio, diluted by whatever biogenic methane was co-sampled, and so bounds "
  "the fossil source composition from below rather than estimating it (the 2021 value is a flux partition of ethane against "
  "oil-and-gas methane and is not diluted in this way):")
R("Using the two ambient values as endmembers raises", "Using these two values as endmembers raises")
# figure captions, main
R("Right: the EPA gridded GHGI fossil fraction for the same domain.",
  "Right: the EPA gridded GHGI fossil fraction on the whole 0.1° cells that overlap the same domain; the black square marks the DADS landfill.")
R("Figure 1: AMMBEC flight legs over the Denver Front Range, colored by leg CH4 enhancement,",
  "Figure 1: AMMBEC flight tracks (red lines) and level legs (circles, colored by leg CH4 enhancement) over the Denver Front Range,")
R("All flight tracks are drawn in red, with the 13 July first flight traced in blue.",
  "The first 13 July flight, which shows the largest enhancements of the campaign, is traced in blue.")
R("no flight reaches an even split at this calibration, and the highest, the second 13 July flight, is 47%.",
  "no flight reaches an even split at this calibration, and the highest, the second 13 July flight, is 47%. The 3 July L2 flight has a single "
  "contributing leg and so no interval; the 8 July interval (0 to 2%) is hidden by its marker.")
R("for a biogenic-dominated flight (07-10, slope near 0) and a fossil-influenced flight (07-13, tight positive slope).",
  "for a biogenic-dominated flight (10 July L1, slope near 0) and a fossil-influenced flight (13 July L2, tight positive slope).")
R("The airborne-derived estimate is larger and less fossil than the inventory, and only the sub-split",
  "The airborne-derived estimate is larger and less fossil than the EPA inventory and GRA2PES v1.1; GRA2PES v2.0beta is larger "
  "still, about three-quarters of it waste. Percentages are fossil/(fossil + biogenic); gray is combustion and other sectors. Only the sub-split")
R("drawn only where at least 12 in-plume boundary-layer samples (below about 1.5 km), spanning at least 50 ppb of methane enhancement "
  "so that a local slope can be fitted (26 of the 106 sampled cells meet both conditions), fall in the cell, over a Vulcan",
  "drawn only where a cell holds at least 12 in-plume boundary-layer samples (below 1.5 km) spanning at least 50 ppb of methane "
  "enhancement, so that a local slope can be fitted (26 of the 106 sampled cells), over a Vulcan")

# ---- SI -----------------------------------------------------------------------------------------
R("Mixing heights are Dalek 2 lidar retrievals (Boulder); that lidar has no data on 29 June, so the flight uses the PUMAS lidar in the DJ Basin.",
  "Mixing heights are retrievals from the NOAA CSL Doppler lidar near Boulder; that lidar has no data on 29 June, so for that "
  "flight they come from the campaign's second Doppler lidar (PUMAS), then in the DJ Basin.", "si")
R("analogue of a flight (script 71 of the analysis repository).", "analogue of a flight.", "si")
R("interquartile range 0.096 to 0.183", f"interquartile range {iqr_txt}", "si")
R("showing that the Suncor refinery has the largest and most repeatable plumes.",
  "showing that the Suncor refinery has the highest median per-survey enhancement.", "si")
R("and the production-area ratio of 0.0813 mol mol−1 (dotted).",
  "and the AMMBEC production-area ratio of 0.0813 mol mol−1 (dotted); in the right panel, red marks plumes inside the box and "
  "blue plumes in the basin.", "si")
R("we do not use the largest observed atmospheric ethane-to-methane slope as a lower bound on the endmember,",
  "we do not use the largest observed atmospheric ethane-to-methane slope as an estimate of the endmember,", "si")
R("and only 4 to 12% of in-box samples fall below 200 m AGL.",
  "and only 4 to 12% of all in-box samples (fewer on the urban legs used in the fits) fall below 200 m AGL.", "si")
R("so the number of individual flights above the halfway line reaches zero at the top of the range,",
  "so the number of individual flights above the halfway line reaches zero for βsource above about 0.096 mol mol−1,", "si")
R(" The per-flight leg counts are reported alongside the fossil fractions, so this can be checked directly.", "", "si")
R("on 13 July the first flight holds all eleven legs", "on 13 July the first flight keeps all eleven of its contributing legs", "si")
R("The median DJB basin ratio behaves the same way, holding at 0.030 to 0.031",
  f"The median basin ratio over the {n_basin0} basin flights with a fit (without the reliability screen of section S5, which "
  "gives 0.036 at zero shift) behaves the same way, holding at 0.030 to 0.031", "si")
R("Right: median DJB basin ethane-to-methane ratio over all basin flights against the same shift. The fossil fraction peaks at the "
  "estimated offset, and both quantities decline as the records are misaligned further.",
  f"Right: median DJB basin ethane-to-methane ratio over the {n_basin0} basin flights with a fit (without the reliability screen of "
  "section S5) against the same shift; it is flat from −3 to +2 s and falls beyond, and the blue line marks the estimated −1 s "
  "offset. The fossil fraction peaks at the estimated offset.", "si")
R("which extend to 0.187, which is consistent with", "which extend to 0.187; this is consistent with", "si")
R("Four flights with strongly negative slopes fall below the axis and are therefore not drawn, although their labels remain on the axis.",
  f"Four flights with slopes of {fmt_neg(neg[-1], 3)} to {fmt_neg(neg[0], 2)} mol mol−1 lie below or at the lower edge of the axis and are not "
  "drawn, or only partly; their labels remain on the axis.", "si")
R("Each flight axis label carries the Pearson correlation for that panel species pair.",
  "Each flight axis label carries the Pearson correlation for that panel species pair; flights with no evaluable ethane fit are blank.", "si")
R("(scripts 57 and 58 of the analysis repository)", "", "si")
R("13 to 19 km downwind of the oil-and-gas sites near Brighton, with slopes of 0.03 to 0.07 and r of 0.8 to 0.9;",
  "13 to 19 km downwind of oil-and-gas sites near Brighton and Westminster, with slopes of 0.03 to 0.07 and r of 0.8 to 0.9, and one "
  "10 July leg 19 km downwind of a site near Brighton (slope 0.07, r = 0.72);", "si")
R("The moderate positive correlation (Pearson r = 0.35, n = 7)",
  f"The weak positive correlation (Pearson r = {r_w:.2f}, n = {len(xs)}, not significant, p = {p_w:.2f})", "si")
R("close to the observed median of 0.035.", f"close to the observed median of {stilt_obs}.", "si")
R("The model reproduces the contrast between legs within and above the boundary layer",
  "The model reproduces the contrast between legs below and above 1 km", "si")
R("but not the variation among boundary-layer legs", "but not the variation among legs below 1 km", "si")
R("is independent of the CO and CO2 anchors and is consistent with section 4.2,",
  "is independent of the CO and CO2 anchors and is consistent with section 5,", "si")
R("Each panel title gives the flight and its York fossil fraction.",
  "Each panel title gives the flight and its York fossil fraction (5 July L1 has none; Table 1).", "si")
R("Distance is the mean distance of the crossing's linked samples from the centre of the 4-km hotspot cell; the centre of the DADS "
  f"cell (site 1) lies {float(HS['1']['facility_km']):.0f} km from the landfill, so this distance differs from the plume-peak distance (3.2 km, "
  f"section S3) and from the mean distance of the linked samples to the landfill itself ({CE['dist_km']} km, section S8).",
  "Distance is the mean distance of the crossing's samples from the center of the 4-km hotspot cell. For site 1 the cell center "
  f"lies {cell_cm_km:.1f} km from the Carbon Mapper DADS source ({float(HS['1']['facility_km']):.1f} km from the facility's "
  f"street address, Table S9); the {EN['dist_km']} km here, the {CE['dist_km']} km of section S8 and the 3.2 km of section S3 differ "
  "because they average different samples (the crossed segment, the samples linked to the source, the plume peak).", "si")
R("the nearest named facility, and the Carbon Mapper", "the nearest named facility (km: from the cell center to its mapped location), and the Carbon Mapper", "si")
R("Dashed line 1:1, dotted lines a factor of two either side; colours identify the sites of Table S9.",
  "Dashed line 1:1, dotted lines a factor of two either side; colors identify the sites of Table S9, and the legend values are the "
  "v2.0beta waste emissions of each hotspot cell. No crossing qualified at site 2.", "si")
R("A flight is flagged separable when the wind is cross-complex (easterly or westerly) and enough plume points were sampled.",
  f"A flight is flagged separable when the wind places the two facilities at least {SEP} m apart across the flow (winds within "
  f"about 50° of east or west) and at least {NPM} plume points were sampled.", "si")
R("sampling within 4 km of the Suncor / Robert Hite complex,", "sampling within 4 km of the Suncor/Metro Water Recovery complex,", "si")
R("Carbon Mapper persistence-weighted rate (kg/h)", "Carbon Mapper persistence-weighted rate (t/h)", "si")

# regenerated figures (labels, legends, crop, CDPHE positions; see the commit of 2 Oct 2026)
FIG = OUT / "figures"
for doc, media, png in [("main", "image3.png", "Fig1_study_map.png"), ("main", "image1.png", "Fig4_source_combined.png"),
                        ("si", "image6.png", "FigS4_mobile_plumes.png"), ("si", "image17.png", "FigS15_arc_ground.png"),
                        ("si", "image13.png", "FigS5_beta_sensitivity.png"), ("si", "image2.png", "FigS8_basin_ethane_trend.png"),
                        ("si", "image10.png", "FigS9_biogenic_sources.png"), ("si", "image3.png", "FigS13_spatial_match.png"),
                        ("si", "image11.png", "FigS12_carbonmapper_encounters.png"), ("si", "image8.png", "FigS11_wind_exposure.png"),
                        ("si", "image14.png", "SI_perflight_maps.png")]:
    L.FIGSWAP(doc, media, str(FIG / png))

tmp_m = Path(a.main_out).with_suffix(".t1.docx"); tmp_s = Path(a.si_out).with_suffix(".t1.docx")
L.apply(a.main_in, tmp_m, "main", "main text")
L.apply(a.si_in, tmp_s, "si", "SI")
main = docx.Document(tmp_m); si = docx.Document(tmp_s)

def set_cell(cell, text):
    cp = cell.paragraphs[0]
    for r_ in cp.runs[1:]: r_._r.getparent().remove(r_._r)
    if not cp.runs: cp.add_run("")
    cp.runs[0].text = text; cp.runs[0].font.color.rgb = L.RED
def table_by_header(d, prefix):
    hits = [t for t in d.tables if " | ".join(c.text for c in t.rows[0].cells).startswith(prefix)]
    assert len(hits) == 1, f"table {prefix!r}: {len(hits)}"
    return hits[0]
def drop_para(d, startswith):
    hits = [p for p in d.paragraphs if p.text.strip().startswith(startswith)]
    assert len(hits) == 1, f"{startswith!r}: {len(hits)}"
    hits[0]._p.getparent().remove(hits[0]._p)

# SI S2.1 trend range (script 82 wrote it from double-rounded slopes)
hit = [p for p in all_paragraphs(si) if re.search(r"slopes of 0\.\d{3} to 0\.\d{3} ppmv per year in magnitude, p = 0\.\d\d to 0\.\d\d", runs_text(p))]
assert len(hit) == 1
m = re.search(r"slopes of 0\.\d{3} to 0\.\d{3} ppmv per year in magnitude, p = 0\.\d\d to 0\.\d\d", runs_text(hit[0]))
new = f"slopes of {slope_txt} ppmv per year in magnitude, p = {p_txt}"
if m.group(0) != new: replace_in_paragraph(hit[0], m.group(0), new, m.start())
# duplicated paragraph in S3 and code constants in S6
drop_para(si, "The ratio method does not mathematically require a boundary-layer height.")
drop_para(si, "The four thresholds are the constants R_MIN")
# Table 1 header
t1 = table_by_header(main, "Flight (leg) | Urban legs"); set_cell(t1.rows[0].cells[0], "Flight")
# Table S4: C2H6:CH4 from the transcribed mol% (the ratio column of the CSV is itself rounded)
PS = {(r["year"], r["month"]): r for r in csv.DictReader(open(REPO / "psco_gas_quality.csv")) if r["zone"] == "DENVER"}
t4 = table_by_header(si, "Month | 2023 CH4 %"); h4 = [c.text for c in t4.rows[0].cells]; n4 = 0
for row in t4.rows[1:]:
    mon = row.cells[0].text.strip()
    for j, h in enumerate(h4):
        mm = re.match(r"(\d{4}) C2H6:CH4", h)
        if not mm or (mm.group(1), mon.upper()[:3]) not in PS or not row.cells[j].text.strip(): continue
        r_ = PS[(mm.group(1), mon.upper()[:3])]
        v = hu(float(r_["ethane_pct"]) / float(r_["methane_pct"]), "0.001")
        if row.cells[j].text.strip() != v: set_cell(row.cells[j], v); n4 += 1
print(f"Table S4: {n4} ratio cell(s) corrected")
# Table S9: Carbon Mapper in t/h, facility km to one decimal
t9 = table_by_header(si, "Site | Nearest named facility | km"); h9 = [c.text for c in t9.rows[0].cells]
icm = next(i for i, h in enumerate(h9) if h.startswith("Carbon Mapper")); ikm = h9.index("km")
for row in t9.rows[1:]:
    hs = HS[row.cells[0].text.strip()]
    if hs["carbonmapper_kg_hr"] not in ("", "NA"): set_cell(row.cells[icm], f"{float(hs['carbonmapper_kg_hr']) / 1000:.2f}")
    set_cell(row.cells[ikm], f"{float(hs['facility_km']):.1f}")
# Table S10: flight labels
t10 = table_by_header(si, "Flight | Leg | Site");
for row in t10.rows[1:]:
    f = row.cells[0].text.strip(); mm = re.match(r"(\d{4})(\d{2})(\d{2})_R\d(?:_(L\d))?$", f)
    if mm: set_cell(row.cells[0], f"{mm.group(1)}-{mm.group(2)}-{mm.group(3)}" + (f" ({mm.group(4)})" if mm.group(4) else ""))
# US spelling throughout (ES&T Air house style)
US = {"colour": "color", "colours": "colors", "coloured": "colored", "centre": "center", "organised": "organized",
      "modelled": "modeled", "grey": "gray"}
n_us = 0
for d_ in (main, si):
    for p in all_paragraphs(d_):
        while True:
            m = re.search(r"\b(" + "|".join(US) + r")\b", runs_text(p), re.I)
            if not m: break
            w = m.group(0); nw = US[w.lower()]; nw = nw[0].upper() + nw[1:] if w[0].isupper() else nw
            if replace_in_paragraph(p, w, nw, m.start()) < 0: raise SystemExit(f"spelling edit failed: {w}")
            n_us += 1
print(f"US spelling: {n_us} word(s)")
L.subscript_chem(all_paragraphs(main)); L.subscript_chem(all_paragraphs(si))
L.exponent_fix(all_paragraphs(main)); L.exponent_fix(all_paragraphs(si))
main.save(a.main_out); si.save(a.si_out); tmp_m.unlink(); tmp_s.unlink()
print("wrote", a.main_out, a.si_out)
