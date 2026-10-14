"""93_coauthor_comment_edits.py -- edits for the October 2026 co-author comments (B. McDonald,
C. Lyu), applied to Priyanka's edited manuscript (aircraft_old / aircraft_manuscript/updated, with
the instrument edits of 13 Oct) and to the current SI (output of script 92), together with the
airborne-lidar mixing-height text of script 92, which her edited manuscript did not yet carry.

Changes
  * GRA2PES v1.1 methane is no longer used (it was a separate beta product derived from the NEI);
    v1.1 is kept as the CO anchor. Figure 4 loses its v1.1 bar; Table S8, Table S10 and Figure
    S17 lose their v1.1 entries.
  * CH4:CO emission scaled by the GRA2PES CO of each flight's day of the week (13 July 2024 was a
    Saturday); Table 1, Table S6 and every emission number follow.
  * NEI sensitivity: GRA2PES CO times the NEI/GRA2PES ratio over the same seven counties (1.88),
    in place of the county total; McDonald et al. (2018) cited for the MOVES-FIVE difference.
  * The cell-by-cell aircraft-inventory comparison leaves the main text (Figure 4B removed); the
    HRRR-STILT footprint comparison and the downwind-crossing test stay. SI S7.1 is kept,
    labelled descriptive.
  * Error bar (range of the four flights) on the airborne bar of Figure 4.
  * Shorter limitations; GRA2PES dataset link; v2.0beta changes relative to v1.1 in section 2.6.
  * ARL car: the ratios use CH4 and C2H6 both from the Aeris (clarified after the 13 Oct edit).
All inserted text is red. Tracked changes already in the file are left for the author.

usage: python3 scripts/93_coauthor_comment_edits.py <main.docx> <si.docx> <main_out> <si_out>
         --outputs <OUT_DIR> --figures <OUT_DIR>/figures
"""
import sys, csv, re, argparse, statistics as st, copy, zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import docx_edit_lib as L
import docx_ops as O

ap = argparse.ArgumentParser()
for k in ("main_in", "si_in", "main_out", "si_out"): ap.add_argument(k)
ap.add_argument("--outputs", required=True); ap.add_argument("--figures", required=True)
a = ap.parse_args(); OUT = Path(a.outputs); FIG = Path(a.figures)
rd = lambda f: list(csv.DictReader(open(OUT / f)))
f1 = lambda x: f"{x:.1f}"

# ---- numbers ---------------------------------------------------------------------------------
RM = {r["flight"]: r for r in rd("ratio_method_flux.csv")}
EM = [r for r in rd("emission_estimates.csv") if r["CH4CO_GRA2PES_t_hr"] not in ("NA", "")]
assert [r["flight"] for r in EM] == ["20240703_R0_L1", "20240703_R0_L2", "20240709_R0", "20240713_R0_L1"], EM
E = [float(r["CH4CO_GRA2PES_t_hr"]) for r in EM]
lo, hi, med = min(E), max(E), st.median(E)
NB = [float(r["CH4CO_NEIbox_t_hr"]) for r in EM]
anch = {r["flight"]: float(RM[r["flight"]]["E_CO_anchor_Gg_yr"]) for r in EM}
WK, SAT = anch["20240703_R0_L1"], anch["20240713_R0_L1"]
assert anch["20240709_R0"] == WK == 126.4 and SAT == 114.0, anch
SA = rd("source_attribution.csv")
Ef = st.median(float(r["E_fossil_t_hr"]) for r in SA); Eb = st.median(float(r["E_biogenic_t_hr"]) for r in SA)
ff = round(100 * Ef / (Ef + Eb))
REP = {r["sector"]: float(r["top_down_t_hr"]) for r in rd("source_attribution_representative.csv")}
INV = rd("inventory_comparison.csv"); inv_b = sum(float(r["t_hr"]) for r in INV if r["group"] == "biogenic")
EPA = 2.8; V2 = 16.5; V2W = 12.2
T1 = {r["Flight"]: r for r in rd("table1.csv")}
cell = lambda s: s.replace("-", " to ").strip()
N = dict(lo=f1(lo), hi=f1(hi), med=f1(med), nlo=f1(min(NB)), nhi=f1(max(NB)), wk=f1(WK), sat=f1(SAT),
         Ef=f1(Ef), Eb=f1(Eb), tot=f1(Ef + Eb), ff=ff, lf=f1(REP["5A1_Landfills_MSW"]), ww=f1(REP["5D_Wastewater_Treatment_Domestic"]),
         nd=f1(REP["1B2b_Natural_Gas_Distribution"]), bx=f1(Eb / inv_b),
         r_lo=f1(lo / EPA), r_hi=f1(hi / EPA), r_med=f1(med / EPA), v_lo=f1(V2 / hi), v_hi=f1(V2 / lo),
         wfrac=f"about {round(100 * Eb / V2W)}%", lo_i=round(lo), hi_i=round(hi))
same_sum = f1(Ef + Eb) == f1(med)
import numpy as np
IQR = np.percentile([float(r["slope"]) for r in rd("wind_source_exposure_legs.csv") if r["slope"] not in ("NA", "")], [25, 75])
BGr = {r["flight"]: r for r in rd("background_mode_sensitivity.csv")}
BGv = [float(BGr[r["flight"]]["E_co_track_t_hr"]) for r in EM]; BG = (min(BGv), max(BGv))
QCt = {r["flight"]: r for r in rd("qc_robustness.csv")}
print("numbers:", N, "| paired-median sum equals median total:", same_sum)

OPS = {"main": [], "si": []}
def R(doc, old, new, count=1): OPS[doc].append(("R", old, new, count))
def RW(doc, anchor, text, drop_h=()): OPS[doc].append(("RW", anchor, text, drop_h))
def DEL(doc, anchor, comments_to=None): OPS[doc].append(("DEL", anchor, comments_to))
M, S = "main", "si"

# ================================ main text ================================
# -- abstract
R(M, "constant-altitude yielded estimable urban ethane-methan slopes", "constant-altitude transects yielded estimable urban ethane-methane slopes")
R(M, "Scaling the CH4:CO enhancement ratio by a gridded CO inventory yielded 4.4 to 10.3 t CH4/h (median 6.5) on four flights passing a reliability screen.",
     f"Scaling the CH4:CO enhancement ratio by a gridded CO inventory matched to each flight's day of the week yielded {N['lo']} to {N['hi']} t CH4/h (median {N['med']}) on four flights passing a reliability screen.")
R(M, "exceed the two published box-consistent bottom-up inventories (1.7 and 2.8 t CH4/h) by about 2.3 to 3.8-fold at the CH4:CO median,",
     f"exceed the EPA gridded inventory for the same area (2.8 t CH4/h) by about {N['r_med']}-fold at the median,")

# -- 2.3 ARL car: the ratio pair is the Aeris CH4 and C2H6 (script 71)
R(M, " among a larger instrument suite.", " among a larger instrument suite; the ethane-to-methane ratios here use the CH4 and C2H6 measured together by the Aeris.")

# -- 2.4 mixing height (airborne lidar; script 92)
RW(M, "Mixing heights were obtained from a NOAA Chemical Sciences Laboratory (CSL) Doppler-lidar site",
   "Mixing heights come from the NOAA Chemical Sciences Laboratory (CSL) airborne Doppler lidar (MD02) on the Twin Otter, which retrieved the boundary-layer height below the aircraft about every 30 s along the flight track "
   "from the lidar's fit variance, range-corrected intensity gradient, wind-speed gradient and wind shear (final R1 release; 17 of the 22 flights).[[^11]] "
   "The ground-based CSL Doppler lidar at the Boulder DSRC (39.99°N, 105.26°W), which reports vertical-velocity variance, supplies the mixing height for the five flights without an airborne retrieval and is kept as a sensitivity case. "
   "{{c13}}The winds in the wind-direction, encounter and transect analyses are the aircraft's own measurements (wind roses for the flights and for June and July 2024 are in Figure S13).{{/c13}} "
   "The High-Resolution Rapid Refresh (HRRR) model drives the Stochastic Time-Inverted Lagrangian Transport (STILT) footprints (section S9).")
RW(M, "Each flight's mixing height is the median of the lidar's retrievals",
   "Each flight's mixing height is the median of the airborne retrievals inside the analysis box during that flight (the whole track where fewer than three lie in the box). "
   "It spans 0.63 to 2.66 km on the seven of the nine flights with estimable urban slopes that have an airborne retrieval; for the other two, 3 July L2 and 12 July L1, the ground lidar gives 3.35 and 2.75 km (Table S1). "
   "On those seven flights all 34 urban transects flown below 1 km above ground lay inside the boundary layer measured along the track, as did 97% of the samples entering the ethane fits, whereas 7 of the 10 transects flown above 1.5 km lay above it. "
   "The ground lidar is a single site about 40 km north-west of the urban core; under strong westerly flow over the mountains its vertical-velocity variance can register turbulence aloft that is not connected to the surface, and its mixing height exceeds the airborne value on 10 of the 13 flights covered by both, by up to 2.9 km (median difference 0.75 km). "
   "The mixing height is used to document the daytime boundary layer, in the sensitivity analysis that applies an in-boundary-layer filter (section S3), in the vertical-coverage test for curtains (section S1.2), and to scale the single-transect rates downwind of inventory hotspots (section S11).")

# -- 2.6 inventories
RW(M, "Several gridded inventories, each aggregated over the same analysis box",
   "Several gridded inventories, each aggregated over the same analysis box (for the 0.1-degree EPA grid, each cell weighted by the fraction of its area inside the box), are used in this analysis. "
   "1) The sector-resolved U.S. gridded EPA Greenhouse Gas Inventory (GHGI; 2020 Express Extension, 0.1 degree){{h0}} provides the bottom-up methane source split. "
   "2) GRA2PES v1.1 (4 km, July 2023){{h1}} {{c15}}{{c16}}{{c17}}provides the primary CO anchor{{/c15}}{{/c16}}{{/c17}}. "
   "3) {{c18}}GRA2PES v2.0beta{{/c18}} (4 km, July 2023; NOAA CSL, beta release of June 2026) provides sector-resolved methane. Its main advance over v1.1 is the addition and revision of the key methane sectors, with facility-level landfill and wastewater emissions, a post-meter natural-gas leakage term, and updated oil-and-gas and agricultural emissions, whereas the methods for combustion sources of CO and CO2 are largely unchanged. "
   "We use it as a second bottom-up methane estimate and to locate the waste sources tested in section 4.3, while treating it explicitly as a beta product {{x0}}. "
   "4) Vulcan v4.0 fossil-fuel CO2 (1 km, 2022){{h2}} provides the CO2 anchor. "
   "The EPA 2020 National Emissions Inventory (NEI) county-level CO totals are used to scale the GRA2PES CO in a sensitivity case (section 3.4).")

# -- 3.2 background (Figure 4B gone)
R(M, "The gridded signature map of Figure 4B and section S7.1, which bins", "The gridded signature map of section S7.1 (Figure S11), which bins")

# -- 3.3 / 3.4 QC filter (script 92)
R(M, "leaves the campaign median unchanged at 21%, although it shifts three flights by 8 to 17 percentage points (section S3).",
     "leaves the campaign median unchanged at 21% and changes only one flight by more than 1 percentage point (section S3).")
R(M, "We did not impose the filters in the primary analysis because our boundary-layer heights come from a single Doppler lidar about 40 km from central Denver and may not represent mixing conditions over the city, particularly on days with a shallow boundary layer. We therefore fit all urban-transect samples; with the filters applied, the campaign median is unchanged (21%), although three flights change by 8 to 17 percentage points (Table S3).",
     "We did not impose the filters in the primary analysis, so that the fits keep every urban transect, including those flown above the mixed layer. We therefore fit all urban-transect samples; with the filters applied, using the boundary-layer height measured along the track by the airborne lidar, the campaign median is unchanged (21%) and only 8 July L1 changes by more than 1 percentage point (17 to 0%; Table S3).")

# -- 3.4 CO anchor: day of week; NEI scaling
RW(M, "We anchored the enhancement ratio to two gridded inventories defined over the same analysis box.",
   "We anchored the enhancement ratio to two gridded inventories defined over the same analysis box. For CH4:CO2 we use Vulcan v4.0 fossil-fuel CO2,{{h0}} reprojected from its 1-km Lambert grid and summed over the 2,856 1-km cells intersecting the box, yielding 23,622 Gg CO2 yr−1 for 2022. "
   "For CH4:CO, our primary anchor is GRA2PES v1.1 total CO,{{h1}} summed over the same domain (171 4-km cells, July 2023). "
   f"Because traffic CO has a strong weekly cycle, each flight's slope is scaled by the GRA2PES CO of its own day type: {N['wk']} Gg CO yr−1 on July weekdays and {N['sat']} on Saturdays (24-hour means), against 121.5 for the July mean over all days. "
   "Three of the four flights passing the reliability screen were on weekdays and 13 July L1 on a Saturday. "
   f"{{{{c24}}}}On a common July-weekday basis, GRA2PES v2.0beta gives a box total of 123.4 Gg CO yr−1 against {N['wk']} for v1.1 (2% lower), as expected because v2.0beta revised the methane sectors while leaving the methods for combustion CO largely unchanged.{{{{/c24}}}}")
RW(M, "The resulting CO-anchored methane estimates are therefore insensitive to the inventory version",
   "The resulting CO-anchored methane estimates are therefore insensitive to the inventory version. "
   "{{c25}}The EPA 2020 NEI gives about 1.9 times more CO than GRA2PES over the seven metropolitan counties (292.7 against 156 Gg CO yr−1; section 5), consistent with the finding of McDonald et al. that the MOVES model used by the NEI gives about twice the on-road CO of the fuel-based FIVE inventory used by GRA2PES.[[^36]]{{/c25}} "
   "As a sensitivity case we therefore multiply each flight's GRA2PES CO by this ratio (1.88), which keeps the analysis-box domain and changes only the magnitude of the CO inventory.")

# -- Figure 1 caption (NEI footprint)
R(M, "The seven-county metropolitan footprint used for the EPA-NEI sensitivity anchor, about 11,700 km2, is roughly four times the area of the analysis box, shown for scale as a locator inset, so the county-scale inventory domain is far larger than the sampled region.",
     "The seven-county metropolitan footprint over which the NEI and GRA2PES CO totals are compared, about 11,700 km2, is roughly four times the area of the analysis box and is shown for scale as a locator inset.")

# -- Table 1 caption
R(M, "and the two enhancement-ratio emission estimates with their bootstrap intervals,",
     "and the two enhancement-ratio emission estimates with their bootstrap intervals (CH4:CO scaled by GRA2PES CO for the flight's day of the week),")

# -- 4.2
R(M, "Scaling the CH4:CO slope by gridded GRA2PES CO (121.5 Gg CO yr−1) yields 4.4 to 10.3 t CH4/h (median 6.5)",
     f"Scaling the CH4:CO slope by gridded GRA2PES CO for each flight's day type ({N['wk']} Gg CO yr−1 on weekdays, {N['sat']} on Saturdays) yields {N['lo']} to {N['hi']} t CH4/h (median {N['med']})")
R(M, "Taken together, the four CH4:CO estimates indicate urban methane emission rates of roughly 4 to 10 t CH4/h, with a median of 6.5 t/h.",
     f"Taken together, the four CH4:CO estimates indicate urban methane emission rates of roughly {N['lo_i']} to {N['hi_i']} t CH4/h, with a median of {N['med']} t/h.")
R(M, "Using the EPA NEI seven-county CO total instead (292.7 Gg CO yr−1) raises the CH4:CO-based estimates to 10.7 to 24.7 t CH4/h. We treat these as an upper bound sensitivity case because the county inventory covers a substantially larger area than the aircraft analysis domain and gives more CO than GRA2PES over the same counties (section 5).",
     f"Scaling GRA2PES CO by 1.9, the ratio of NEI to GRA2PES CO over the seven metropolitan counties (section 3.4), raises the CH4:CO-based estimates to {N['nlo']} to {N['nhi']} t CH4/h. We treat these as an upper-bound sensitivity case, because the difference reflects the two CO inventories rather than Denver's methane (section 5).")
R(M, "are already the most fossil-influenced flights at the adopted calibration and were sampled under northerly or north-easterly flow.",
     "are already the most fossil-influenced flights at the adopted calibration and were sampled under northerly or north-easterly flow, from the direction of the DJB, as expected if advected basin gas contributes to their fossil signal.")

# -- 4.3 allocation
R(M, "This gives approximately 1.9 t CH4/h from fossil sources and 4.6 t CH4/h from biogenic sources- a combined 6.5 t/h, of which about 29% is fossil (Figure 4A).",
     f"This gives approximately {N['Ef']} t CH4/h from fossil sources and {N['Eb']} t CH4/h from biogenic sources- a combined {N['tot']} t/h, of which about {N['ff']}% is fossil (Figure 4).")
R(M, "Because the two medians are calculated separately, their sum need not equal the median total emission estimate; with four flights it does here (6.5 t/h), because the same flights are central in all three rankings.",
     f"Because the two medians are calculated separately, their sum need not equal the median total emission estimate; with four flights it does here ({N['med']} t/h), because the same flights are central in all three rankings." if same_sum else
     f"Because the two medians are calculated separately, their sum need not equal the median total emission estimate ({N['med']} t/h).")
R(M, "This assigns approximately 1.9 t CH4/h to municipal solid-waste landfills, 1.6 t CH4/h to domestic wastewater, and 0.8 t CH4/h to natural-gas distribution.",
     f"This assigns approximately {N['lf']} t CH4/h to municipal solid-waste landfills, {N['ww']} t CH4/h to domestic wastewater, and {N['nd']} t CH4/h to natural-gas distribution.")
R(M, "The inferred biogenic total is approximately 4.5 times the inventory estimate,", f"The inferred biogenic total is approximately {N['bx']} times the inventory estimate,")

# -- 4.3 spatial comparison -> transport-aware comparisons only
RW(M, "The spatial comparison shows a similar contrast (Figure 4B).",
   "{{c32}}We do not compare the aircraft with the inventories cell by cell. An aircraft sample integrates emissions transported from upwind, not only those of the cell beneath it, so a local mismatch does not by itself establish an inventory error, and a meaningful spatial test requires convolving the inventory with transport footprints.{{/c32}} "
   "The aircraft-sampled ethane-to-methane character is mapped in Figure S11, and the cell-matched statistics are given for reference in section S7.1 (Table S8); the domain-wide comparison in section 5 is the primary inventory comparison.")
RW(M, "Accounting for atmospheric transport narrows the discrepancy.",
   "{{c33}}For the EPA inventory we made this footprint-based comparison.{{/c33}} Combining HRRR-STILT footprints with the EPA inventory predicts a median leg-level ethane-to-methane slope of 0.041 mol mol−1, compared with 0.034 observed, for the 25 urban level legs of the Vertical_Speed sensitivity definition (section S9). "
   "The modeled air reaching the aircraft includes waste methane transported from upwind cells, including a large source cell just north of the analysis box, so the transport-based prediction is much closer to the observed signature than the inventory's box-wide source mix (61% fossil), although it still indicates a more fossil-rich mixture. "
   "Extending the footprint-based comparison to GRA2PES v2.0beta, for example with FLEXPART-WRF footprints, is a natural next step.")
DEL(M, "The two GRA2PES versions differ sharply in how they represent biogenic methane.", comments_to="Adding revised waste emissions in v2.0beta lowers the box-wide fossil share")
DEL(M, "This concentration of biogenic emissions complicates a direct cell-by-cell comparison.", comments_to="Adding revised waste emissions in v2.0beta lowers the box-wide fossil share")
RW(M, "Adding revised waste emissions in v2.0beta lowers the box-wide fossil share",
   "{{c34}}{{c35}}{{c36}}In GRA2PES v2.0beta, waste methane is concentrated in a few landfill and wastewater cells, leaving most other cells almost entirely fossil, whereas the aircraft observe biogenic or mixed signatures across a broader area.{{/c34}}{{/c35}}{{/c36}} "
   "{{c37}}Rather than compare cells directly, we test these sources where the aircraft crossed downwind of them, using the aircraft's own winds and the mixing height measured along the track.{{/c37}}")
RW(M, "The three waste-source cells contain 12.0 of v2.0beta",
   "The three waste-source cells contain 12.0 of v2.0beta's 12.2 t CH4/h of box-wide waste emissions. "
   "We identified 15 aircraft crossings on seven flights within 12 km and 25° of the downwind direction from two of these cells (section S11, Figure S17, Table S10). "
   "The three crossings downwind of the DADS landfill (two flights) give a median biogenic methane rate of 0.06 t CH4/h- about 1% of the inventory's waste emissions in the upwind area, dominated by the landfill cell's 6.07 t CH4/h. "
   "Twelve crossings downwind of the waste cell immediately north of the Suncor/Metro Water Recovery complex (seven flights) give a median biogenic rate of 0.50 t CH4/h, a median 1.3 times the v2.0beta waste emissions in their upwind fetch, "
   "and total methane rates across these crossings range from 0.39 to 3.68 t CH4/h. "
   "These estimates scale directly with the mixing height, measured by the airborne lidar at each crossing, and assume uniform vertical mixing, so they provide approximate constraints rather than reliable facility emission rates.")
R(M, "It assigns approximately 1.6 t CH4/h to metropolitan domestic wastewater.", f"It assigns approximately {N['ww']} t CH4/h to metropolitan domestic wastewater.")
R(M, "More broadly, both the spatial comparison (Figure 4B) and the aggregate allocation indicate a larger biogenic contribution than the EPA inventory: the allocated top-down fossil share is approximately 29%,",
     f"More broadly, the aggregate allocation indicates a larger biogenic contribution than the EPA inventory: the allocated top-down fossil share is approximately {N['ff']}%,")

# -- Figure 4 caption (one panel; no v1.1; error bar)
RW(M, "Figure 4: Urban methane source mix, airborne-derived versus bottom-up",
   "**Figure 4:** Urban methane source mix, airborne-derived versus bottom-up inventories, shown as a consistency check rather than a calibrated budget. "
   "The airborne-derived bar is the sum of the median fossil and median biogenic emission rates of the four flights with a usable CH4:CO estimate (section 4.3) "
   f"{{{{c39}}}}(left bar){{{{/c39}}}}; its error bar spans the four flight-level estimates ({N['lo']} to {N['hi']} t CH4/h). "
   "The sector-resolved EPA gridded inventory and {{c40}}GRA2PES v2.0beta{{/c40}} are shown for comparison (right bars; v2.0beta classes are waste plus agriculture as biogenic and oil-and-gas plus post-meter as fossil). "
   "The airborne-derived estimate is larger and less fossil than the EPA inventory; GRA2PES v2.0beta is larger still, about three-quarters of it waste. "
   "Percentages are fossil/(fossil + biogenic); gray is combustion and other sectors.")

# -- 5 Discussion
R(M, "The sampled plume character across the metropolitan area is shown in Figure 4B,", "The sampled plume character across the metropolitan area is shown in Figure S11,")
R(M, "importantly, the maps depict measured ethane-to-methane ratios", "importantly, the map depicts measured ethane-to-methane ratios")
R(M, "Denver’s urban emission-rate estimate (~4 to 10 t/hr)", f"Denver’s urban emission-rate estimate (~{N['lo_i']} to {N['hi_i']} t/hr)")
R(M, "A second box-consistent bottom-up estimate, GRA2PES v1.1 for July 2023, totals 1.7 t CH4/h but contains no waste methane within the box; its 91% fossil share therefore reflects a gas-system inventory paired with an almost empty biogenic category. By contrast, GRA2PES v2.0beta totals 16.5 t CH4/h for the same month,",
     "GRA2PES v2.0beta, which adds facility-level landfill and wastewater emissions and post-meter leakage, totals 16.5 t CH4/h for July 2023,")
R(M, "corresponding to a fossil share of 24% of the fossil-plus-biogenic total (Figure 4A).", "corresponding to a fossil share of 24% of the fossil-plus-biogenic total (Figure 4).")
RW(M, "The two published inventories therefore place Denver near 2 t CH4/h",
   f"The EPA inventory therefore places Denver near 2.8 t CH4/h, whereas the aircraft enhancement-ratio estimates are approximately {N['lo_i']}-{N['hi_i']} t/h. "
   f"The airborne estimates exceed the EPA inventory by factors of roughly {N['r_lo']}-{N['r_hi']} across flights and by about {N['r_med']} at the CH4:CO median of {N['med']} t/h, comparable to top-down-to-bottom-up discrepancies reported for other U.S. cities.{{{{h0}}}} "
   f"GRA2PES v2.0beta instead lies above the CH4:CO airborne range by a factor of approximately {N['v_lo']}-{N['v_hi']}. Its waste sector alone (12.2 t/h) exceeds the entire CH4:CO range. "
   "Its source split, however, is much closer to the airborne ethane result: v2.0beta is 76% biogenic, compared with an airborne biogenic fraction of 74-81% across the regional reference endmembers. "
   f"The disagreement is therefore primarily one of magnitude. The aircraft-derived biogenic total of approximately {N['Eb']} t/h (section 4.3) is {N['wfrac']} of the v2.0beta waste total, and the crossings downwind of the DADS landfill in section 4.3 point the same way, although near the wastewater plant the median crossing exceeds the v2.0beta value there. "
   "The EPA inventory and GRA2PES v2.0beta thus bracket the aircraft estimate from opposite sides: the EPA inventory is lower and more fossil, whereas v2.0beta appears, on the evidence available here, to overstate waste methane.")
RW(M, "The anchor comparisons illustrate this point.",
   f"The anchor comparisons illustrate this point. Scaling CH4:CO by the box-consistent GRA2PES CO for each flight's day type yields {N['lo']}-{N['hi']} t CH4/h. "
   "No flight passes the CH4:CO2 reliability screen in this campaign, because summertime biogenic CO2 weakens the CH4:CO2 correlation (section 4.2), so the box-consistent Vulcan anchor cannot provide an independent tracer check here. "
   f"Scaling the GRA2PES CO by 1.9, the ratio of NEI to GRA2PES CO over the seven metropolitan counties, raises the estimates to {N['nlo']}-{N['nhi']} t CH4/h.")
RW(M, "That difference has two parts, neither of which is a change in Denver’s methane emissions.",
   "That difference is not a change in Denver's methane emissions but a difference in the magnitude of the two CO inventories over the same domain (292.7 against 156 Gg CO yr−1 over the seven counties), "
   "consistent with the roughly twofold higher on-road CO of the MOVES model used by the NEI relative to the fuel-based FIVE inventory used by GRA2PES.[[^36]] "
   "We therefore treat the GRA2PES-anchored values as primary and the NEI-scaled values as an upper-bound sensitivity case. "
   "More generally, spatial consistency and contemporaneity of the inventory anchor are the principal determinants of uncertainty in the absolute enhancement-ratio emission estimate.")

# -- limitations, shortened
RW(M, "The ethane endmember is the dominant systematic uncertainty in the absolute fossil-biogenic split",
   "The ethane endmember is the dominant systematic uncertainty in the absolute fossil-biogenic split. The campaign median is 19 to 26% fossil across the regional reference values (0.081-0.110 mol mol−1), "
   "which span the production-area ratio measured during this campaign, the 2021 basin mass balance{{h0}} and the composition of gas delivered to Denver (Table S4). "
   "It would exceed 50% only for an effective endmember below 0.042 mol mol−1, as ethane-poor as the gas delivered to East Coast cities; Denver's delivered gas is not, but these reference values do not exclude a leaner effective endmember for the urban fossil mixture (sections 3.3, S3 and S5).",
   drop_h=(1, 2))
RW(M, "Boundary-layer height is constrained at only one location.",
   "The mixing height is measured along the flight track by the airborne Doppler lidar on 17 of the 22 flights; for the other five, including 3 July L2 and 12 July L1 in Table 1, it comes from a single ground lidar about 40 km from the urban core. "
   "It enters only the sensitivity tests and the single-transect rates of section 4.3, not the enhancement-ratio results.")
RW(M, "The flight geometry was not designed to close an urban methane budget.",
   "The flight geometry was designed around the DJB and does not close an urban budget, so we report enhancement-ratio estimates rather than a mass-balance flux (section S1). "
   "Some advected basin methane may contribute to the fossil signal of the 13 July flights, which extended toward the basin; HRRR-STILT footprints convolved with the EPA inventory attribute roughly 22 to 36% of the predicted urban enhancement to basin oil and gas (section S9).")
RW(M, "The inventory anchors control the absolute emission magnitude.",
   "The absolute emission rates scale linearly with the CO inventory anchor, which is not contemporaneous with the flights (GRA2PES, July 2023; EPA GHGI and NEI, 2020), and the bootstrap intervals capture only plume-to-plume variability, not anchor uncertainty. "
   "{{c43}}GRA2PES v2.0beta is a beta release whose revised methane sectors are still being evaluated; our comparison is one such evaluation.{{/c43}}")
RW(M, "The enhancement-ratio method assumes sufficiently similar transport of methane and its anchor species.",
   "The CH4:CO scaling assumes that methane and CO share transport and dilution, an assumption that is imperfect where landfills and wastewater plants are offset from traffic CO (section S7).")
RW(M, "The campaign represents summer conditions only.",
   "{{c44}}All results are summertime values from nine flights on seven days, with 58 urban transects and about 6.7 h of transect flight inside the analysis box (section 4.1).{{/c44}} Urban thermogenic fractions are typically lower in summer than in winter,{{h0}} and urban methane emissions vary seasonally,{{h1}} "
   "so Denver's annual-mean fossil fraction, with its heating-driven gas demand, may be higher than reported here; cold-season flights are needed to resolve this.")
RW(M, "The supporting ground measurements are contextual rather than quantitatively integrated with the aircraft.",
   "The ground measurements are contextual: the CDPHE surveys are not coincident with the flights, and the ARL car sampled the city south of 39.80°N appreciably on only one day, so they characterize source types and check the airborne ratios rather than constrain the estimates.")

# -- accuracy fixes from the 14 Oct audit
R(M, "Of the 22 flights, 20 crossed the Denver metropolitan analysis box,", "Of the 22 flights, 21 crossed the Denver metropolitan analysis box,")
R(M, "Of the 22 in situ flights, 20 crossed the Denver-metro box on at least one constant-altitude transect,", "Of the 22 in situ flights, 21 crossed the Denver-metro box on at least one constant-altitude transect,")
R(M, "a very ethane-poor fossil source is unlikely to have dominated that leg; the lower floors on the other three flights do not exclude it (section S3).",
     "so these floors, like the lower ones on 3 July L1 and 9 July, do not exclude an ethane-poor fossil source (section S3).")
R(M, "Even the strongest of these leg-level constraints (0.043-0.044 mol mol−1 on both July 13 flights) lies below every endmember considered,",
     "Even the strongest of these transect-level floors (0.043-0.044 mol mol−1 on both 13 July flights) lies below every endmember considered,")
R(M, "Transect slopes cluster around 0.014 to 0.053 mol mol−1 (interquartile range)", f"Transect slopes cluster around {IQR[0]:.3f} to {IQR[1]:.3f} mol mol−1 (interquartile range)")
R(M, "supplies the mixing height for the five flights without an airborne retrieval and is kept as a sensitivity case.",
     "supplies the mixing height for four of the five flights without an airborne retrieval (the campaign's second Doppler lidar, PUMAS, for 28 June, when the DSRC lidar had no data) and is kept as a sensitivity case.")
R(S, "For them the mixing height comes from the ground Doppler lidar at the Boulder DSRC:", "For them the mixing height comes from the ground Doppler lidar at the Boulder DSRC (PUMAS on 28 June, when the DSRC lidar had no data):")
R(S, "below the lidar mixing height from the Boulder site,", "below the boundary-layer height measured by the airborne lidar (the ground lidar where there is no airborne retrieval),")
R(S, "and their emission estimates become 4.7 to 9.4 t CH4/h.", f"and their emission estimates become {BG[0]:.1f} to {BG[1]:.1f} t CH4/h.")

# -- conclusion, data availability
R(M, "place urban methane emissions at approximately 4-10 t CH4/h, above the two box-consistent published bottom-up inventories,",
     f"place urban methane emissions at approximately {N['lo_i']}-{N['hi_i']} t CH4/h, above the EPA gridded inventory for the same area,")
R(M, "The AMMBEC airborne in situ data (NOAA Twin Otter, 1 s ICARTT files) are available",
     "The AMMBEC airborne in situ data (NOAA Twin Otter, 1 s ICARTT files) and airborne Doppler-lidar boundary-layer heights are available")
R(M, "GRA2PES v1.1 and v2.0beta (July 2023) CO and sector methane (v2.0beta obtained from NOAA CSL under its beta data policy)",
     "GRA2PES v1.1 CO and v2.0beta sector methane (July 2023; https://csl.noaa.gov/groups/csl4/gra2pes/datasets/; v2.0beta obtained from NOAA CSL under its beta data policy)")

MCD = ("(36)\tMcDonald, B. C.; McKeen, S. A.; Cui, Y. Y.; Ahmadov, R.; Kim, S.-W.; Frost, G. J.; Pollack, I. B.; Peischl, J.; Ryerson, T. B.; Holloway, J. S.; Graus, M.; Warneke, C.; Gilman, J. B.; de Gouw, J. A.; Kaiser, J.; Keutsch, F. N.; Hanisco, T. F.; Wolfe, G. M.; Trainer, M. "
       "Modeling Ozone in the Eastern U.S. Using a Fuel-Based Mobile Source Emissions Inventory. Environ. Sci. Technol. 2018, 52 (13), 7360–7370. https://doi.org/10.1021/acs.est.8b00778")

# ================================ SI ================================
R(S, "S7.1 Cell-by-cell comparison of the aircraft with three inventories", "S7.1 Cell-matched comparison of the aircraft with two inventories (descriptive)")
RW(S, "Figure 4B compares the aircraft-sampled ethane character with the EPA inventory by eye.",
   "This section maps the aircraft-sampled ethane character (Figure S11A) and, for reference only, lists cell-matched statistics against two inventories (Table S8). "
   "A cell-matched comparison does not account for transport: an aircraft sample integrates emissions carried from upwind, so a local mismatch need not indicate an inventory error. "
   "The main text therefore compares the inventories with the aircraft domain-wide and, for the EPA inventory, through HRRR-STILT footprints (section S9). "
   "For each of the 26 aircraft cells (0.02 degree) with a local York ethane:methane slope, the fossil fraction at the adopted endmember is compared with the fossil fraction (fossil over fossil plus biogenic, the definition used for the EPA inventory in the main text) of the 0.1-degree EPA cell and of the 4-km GRA2PES v2.0beta cell containing it. "
   "GRA2PES per-cell sector methane was extracted from the July 2023 weekday files with the same cell and unit handling as the box sums; fossil is oil-and-gas plus post-meter, biogenic is waste plus agriculture. "
   "At these cells the median aircraft fossil fraction is 34% (31-42% across the regional reference endmembers), against 83% for the EPA inventory and 99% for GRA2PES v2.0beta, whose waste methane lies in a few cells; "
   "across all 49 GRA2PES cells crossed by the aircraft, which hold 45% of its box methane, the emission-weighted v2.0beta fossil share is 12%.")
R(S, "ρ = -0.14 (p = 0.78) for the EPA inventory, 0.37 (p = 0.18) for v1.1 and 0.44 (p = 0.10) for v2.0beta, none of them significant.",
     "ρ = -0.14 (p = 0.78) for the EPA inventory and 0.44 (p = 0.10) for v2.0beta, neither of them significant.")
R(S, "with three gridded inventories at the aircraft-sampled cells.", "with two gridded inventories at the aircraft-sampled cells.")
R(S, "summed for waste and for total methane; the same fetch is summed for v1.1.", "summed for waste and for total methane.")
R(S, "(B) the total aircraft rate against the v2.0beta total, with the v1.1 total in the same fetch as crosses.", "(B) the total aircraft rate against the v2.0beta total.")
RW(S, "Urban CH4:CO and CH4:CO2 slopes (mol/mol; CH4:CO2 in ppb/ppm), their correlations, and the York fossil fraction.",
   "**Table S6:** Urban CH4:CO and CH4:CO2 slopes (mol/mol; CH4:CO2 in ppb/ppm), their correlations, the York fossil fraction and the CH4:CO emission. OLS fits the emission slopes; the fossil fraction uses the within-transect York fit of section 3.3. "
   "The CO anchor is the GRA2PES v1.1 box CO for the flight's day of the week (July 2023, 24-hour mean); the emission, given for the four flights that pass the reliability screen (Table 1), is the CH4:CO slope × 16.04/28.01 × the CO anchor, converted to t/h over 8,766 h per year. "
   f"The NEI-scaled sensitivity multiplies it by 1.88 ({N['nlo']} to {N['nhi']} t/h).")
R(S, "is 0.63, below the 0.7 criterion.",
     "is 0.63, below the 0.7 criterion. "
     f"The CO anchor of each emission is the GRA2PES v1.1 box CO for the flight's day of the week (July 2023, 24-hour mean: {N['wk']} Gg CO yr−1 on weekdays and {N['sat']} on Saturdays, against 121.5 for the mean over all days), because traffic CO has a strong weekly cycle; "
     "with the monthly mean for every flight the four estimates would be 4.4 to 10.3 t/h.")

# ================================ apply ================================
def run_doc(src, dst, which, label, post=None):
    d = docx.Document(src); n = 0
    for op in OPS[which]:
        if op[0] == "R": O.do_replace(d, op[1], op[2], op[3], label)
        elif op[0] == "RW": O.do_rewrite(d, op[1], op[2], label, op[3])
        elif op[0] == "DEL": O.delete_para(d, op[1], label, op[2])
        n += 1
    if post: post(d)
    allp = L.all_paragraphs(d)
    m = L.subscript_chem(allp); e = L.exponent_fix(L.all_paragraphs(d)); c = L.cite_fix(L.all_paragraphs(d)); g = O.gra2pes_superscript(d)
    d.save(dst); print(f"{label}: {n} edit(s) -> {dst} ({m} subscripts, {e} exponents, {c} citations)")

def main_post(d):
    t = d.tables[0]; col = 3
    lab = {"2024-07-03 (L1)": "20240703_R0_L1", "2024-07-03 (L2)": "20240703_R0_L2", "2024-07-09": "20240709_R0", "2024-07-13 (L1)": "20240713_R0_L1"}
    k = 0
    for row in t.rows[1:]:
        f = lab.get(row.cells[0].text.strip())
        if f: k += O.set_cell(row.cells[col], cell(T1[f]["CH4CO_GRA2PES_t_hr"]))
    print(f"main text: Table 1 CO column, {k} cell(s) changed")
    # reference list: drop (35) Terracon (no longer cited), add McDonald et al. 2018, renumber by first citation
    refs = [p for p in d.paragraphs if re.match(r"\(\d+\)\s", p.text.strip())]
    terr = [p for p in refs if "Terracon" in p.text]; assert len(terr) == 1
    num = int(re.match(r"\((\d+)\)", terr[0].text.strip()).group(1)); assert num == 35 == len(refs)
    last = refs[-1]
    new = copy.deepcopy(last._p); last._p.addnext(new)
    np_ = Paragraph(new, last._parent)
    for ch in list(new):
        if ch.tag != qn("w:pPr"): new.remove(ch)
    r = np_.add_run(MCD.replace("(36)", "(35)", 1)); r.font.color.rgb = L.RED
    terr[0]._p.getparent().remove(terr[0]._p)
    for p in L.all_paragraphs(d):
        for x in p._p.iter(qn("w:t")):
            if x.text and "[[^36]]" in x.text: x.text = x.text.replace("[[^36]]", "[[^35]]")
    L.cite_fix(L.all_paragraphs(d))
    L.renumber_refs(d, "main text", start="Abstract")

run_doc(a.main_in, a.main_out, "main", "main text", post=main_post)
def nth_picture(path, n):          # media file of the n-th embedded picture in reading order
    z = zipfile.ZipFile(path); rels = z.read("word/_rels/document.xml.rels").decode(); x = z.read("word/document.xml").decode()
    rid = {m.group(1): m.group(2) for m in re.finditer(r'Id="(rId\d+)"[^>]*Target="media/([^"]+)"', rels)}
    rid.update({m.group(2): m.group(1) for m in re.finditer(r'Target="media/([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    order = [rid[m.group(1)] for m in re.finditer(r'r:embed="(rId\d+)"', x) if m.group(1) in rid]
    return order[n - 1]
f4 = nth_picture(a.main_out, 4)
O.swap_figures(a.main_out, {f4: FIG / "Fig4_source_combined.png"}, {f4}, "main text")   # Figure 4: one panel now

def si_post(d):
    t = d.tables[7]; rows = [r for r in t.rows if r.cells[0].text.startswith("GRA2PES v1.1")]; assert len(rows) == 1
    rows[0]._tr.getparent().remove(rows[0]._tr)
    t = d.tables[9]; assert t.rows[0].cells[-1].text.startswith("v1.1 total in fetch")
    for tr in t._tbl.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc")); tr.remove(tcs[-1])
    g = t._tbl.find(qn("w:tblGrid")); gc = g.findall(qn("w:gridCol")); g.remove(gc[-1])
    t3 = d.tables[2]; assert len(t3.columns) == 5 and t3.rows[0].cells[0].text == "Flight"
    for row in t3.rows[1:]:
        a_, b_ = row.cells[1].text.strip(), row.cells[2].text.strip()
        if a_.lstrip("-").isdigit() and b_.lstrip("-").isdigit(): O.set_cell(row.cells[3], str(abs(int(a_) - int(b_))))
    t = d.tables[5]; assert t.rows[0].cells[0].text == "Flight" and len(t.columns) == 7
    lab = lambda f: (lambda m: f"{m.group(1)}-{m.group(2)}-{m.group(3)}" + (f" ({m.group(4)})" if m.group(4) else ""))(re.match(r"(\d{4})(\d\d)(\d\d)_R0_?(L\d)?", f))
    byl = {lab(f): r for f, r in RM.items()}
    rows = []
    for row in t.rows[1:]:
        old = [c.text for c in row.cells]; r = byl[old[0]]
        e = f"{float(r['E_CH4_from_CO_t_hr']):.1f}" if r["usable_co"] == "TRUE" else ""
        rows.append(old + [f"{float(r['E_CO_anchor_Gg_yr']):.1f}", e])
    O.do_table(d, 5, rows, [None] * 7 + ["CO anchor (Gg/yr)", "CH4:CO emission (t/h)"], "SI", widths=[1.15, 0.5, 0.75, 0.45, 0.75, 0.45, 0.6, 0.7, 0.85])
run_doc(a.si_in, a.si_out, "si", "SI", post=si_post)
O.swap_figures(a.si_out, {"image8.png": FIG / "FigS13_spatial_match.png", "image4.png": FIG / "FigS14_v2_pointsource.png"}, set(), "SI")

# ---- checks ----------------------------------------------------------------------------------
dm = docx.Document(a.main_out); tm = "\n".join(O.vis_text(p) for p in L.all_paragraphs(dm))
for b in ("v1.1 for July 2023", "1.7 and 2.8", "Figure 4B", "Figure 4A", "10.7 to 24.7", "40 km from central Denver", "constrained at only one location",
          "0.85 t CH4/h", "Terracon", "two published", "ethane-methan slopes", "4.4 to 10.3", "median 6.5"):
    assert b not in tm, f"main: {b!r} remains"
for g in (N["lo"], N["hi"], N["nlo"], N["nhi"], "McDonald, B. C.", "csl4/gra2pes/datasets"):
    assert g in tm, f"main: {g!r} missing"
x = zipfile.ZipFile(a.main_out).read("word/document.xml").decode()
xc = zipfile.ZipFile(a.main_out).read("word/comments.xml").decode()
ids_doc = set(re.findall(r'<w:commentReference w:id="(\d+)"', x)); ids_c = set(re.findall(r'<w:comment [^>]*?w:id="(\d+)"', xc))
assert ids_doc == ids_c, (sorted(ids_c - ids_doc), sorted(ids_doc - ids_c))
ds = docx.Document(a.si_out); ts = "\n".join(O.vis_text(p) for p in L.all_paragraphs(ds))
for b in ("v1.1 total", "for v1.1", "Figure 4B", "three inventories", "three gridded"):
    assert b not in ts, f"SI: {b!r} remains"
assert len(ds.tables[5].columns) == 9 and len(ds.tables[7].rows) == 3 and len(ds.tables[9].columns) == 12
print(f"checks passed; {len(ids_doc)} comments anchored in the main text")
