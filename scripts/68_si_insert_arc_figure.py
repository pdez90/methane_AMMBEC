"""68_si_insert_arc_figure.py -- insert the ARC ground figure at the end of SI S2.3 and renumber
every SI figure, in both documents, by physical order of appearance.

Unlike scripts/60_si_reorder.py this does NOT reorder sections: section headings and table numbers
are left exactly as they are, and only figure numbers (captions and every reference to them in the
main text and SI) are rewritten. Inserting a figure inside S2 pushes all later figures up by one,
which is what this fixes.

usage: python3 scripts/68_si_insert_arc_figure.py <main.docx> <si.docx> <main_out> <si_out>
           --outputs <OUT_DIR> --figures <OUT_DIR>/figures
"""
import sys, re, csv, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx
from docx.oxml.ns import qn
import docx_edit_lib as L
from docx_edit_lib import I, runs_text, all_paragraphs, replace_in_paragraph

ap = argparse.ArgumentParser(); ap.add_argument("main_in"); ap.add_argument("si_in"); ap.add_argument("main_out"); ap.add_argument("si_out")
ap.add_argument("--outputs", required=True); ap.add_argument("--figures", required=True)
ap.add_argument("--png", default="FigS15_arc_ground.png")
a = ap.parse_args(); OUT = Path(a.outputs); L.FIGDIR = a.figures; L.OUTDIR = str(OUT)

# Counts for the caption, from the same CSV the figure was drawn from.
PL = list(csv.DictReader(open(OUT / "arc_ground_plumes.csv")))
URBAN = dict(lat_s=39.50, lat_n=39.95, lon_w=-105.20, lon_e=-104.55); BASIN_LAT = 40.05
xl = (URBAN["lon_w"] - 0.5, URBAN["lon_e"] + 0.5); yl = (URBAN["lat_s"] - 0.3, BASIN_LAT + 0.8)
shown = sum(1 for p in PL if xl[0] <= float(p["lon"]) <= xl[1] and yl[0] <= float(p["lat"]) <= yl[1])
print(f"caption: {shown} of {len(PL)} plumes inside the plotted area")

# The caption is inserted with a provisional number; the renumbering pass below fixes it from the
# figure's physical position, so the number written here only has to be unique.
CAPTION = (
  "Figure S99: Ground plumes crossed by the NOAA ARL Air Resources Car, 24 June to 15 July 2024. "
  "Left: plume locations, coloured by class (ethane-bearing, ethane-free, unclassified) and sized by "
  "peak methane enhancement, with the Denver-metro analysis box (green), the 40.05°N basin threshold "
  f"(dashed) and the 39.80°N line below which the box is metropolitan Denver (dotted); {shown} of the "
  f"{len(PL)} plumes fall in the plotted area, the remainder on transits outside the region. Right: "
  "ethane-to-methane ratio of the ethane-bearing plumes inside the box and in the basin, one point per "
  "plume, with the medians (thick lines), the adopted endmember of 0.102 mol mol-1 (dashed) and the "
  "production-area ratio of 0.0813 mol mol-1 (dotted)."
)
I("left to a dedicated analysis.", f"[[FIG:{a.png}]]\n\n{CAPTION}", "si")
L.apply(a.main_in, "/dev/null" if False else a.main_out, "main", "main text")   # no main-text edits; copies through
L.apply(a.si_in, a.si_out, "si", "SI")

# ---- renumber figures by physical order, in both documents --------------------------------
main = docx.Document(a.main_out); si = docx.Document(a.si_out)
CAP_FIG = re.compile(r"^\s*Figure\s+S(\d+)([a-z]?)\s*[:.]")
fig_map = {}
for p in si.paragraphs:
    m = CAP_FIG.match(p.text)
    if m:
        key = m.group(1) + m.group(2)
        if key not in fig_map: fig_map[key] = str(len(fig_map) + 1)
print("figure map:", ", ".join(f"S{o}->S{n}" for o, n in fig_map.items() if o != n) or "identity")

TOK = r"S\d+[a-z]?"
LIST = r"(?:%s(?:\s*(?:,|and|to|-|–)\s*%s)*)" % (TOK, TOK)
REF_RE = re.compile(r"(Figures?)(\s+)(%s)" % LIST)
CAPF_RE = re.compile(r"^(Figure\s+S)(\d+[a-z]?)(\s*[:.])")
def map_tok(tok):
    m = re.match(r"S(\d+)([a-z]?)", tok); return "S" + fig_map.get(m.group(1) + m.group(2), m.group(1) + m.group(2))
def map_list(lst):
    parts = re.split(r"(\s*(?:,|and|to|-|–)\s*)", lst)
    out = "".join(map_tok(t) if re.fullmatch(TOK, t) else t for t in parts)
    if not re.search(r"[-–]|\bto\b", lst):                       # a plain list stays ascending
        nums = re.findall(TOK, out); key = lambda t: (int(re.sub(r"[a-z]", "", t[1:])), t)
        if nums != sorted(nums, key=key):
            it = iter(sorted(nums, key=key)); out = re.sub(TOK, lambda m: next(it), out)
    return out
def rewrite(p):
    text = runs_text(p); edits = {}
    m = CAPF_RE.match(text)
    if m: edits[0] = (m.group(0), m.group(1) + fig_map.get(m.group(2), m.group(2)) + m.group(3))
    for m in REF_RE.finditer(text):
        if m.start() in edits: continue
        edits[m.start()] = (m.group(0), m.group(1) + m.group(2) + map_list(m.group(3)))
    n = 0; pos = 0
    for _, (old, new) in sorted(edits.items()):
        if old == new: pos = runs_text(p).find(old, pos) + len(old); continue
        pos2 = replace_in_paragraph(p, old, new, pos)
        if pos2 < 0: raise SystemExit("could not place %r in: %s" % (old, text[:80]))
        pos = pos2; n += 1
    return n
n_main = sum(rewrite(p) for p in all_paragraphs(main)); n_si = sum(rewrite(p) for p in all_paragraphs(si))
print(f"rewrote {n_main} figure reference(s) in the main text and {n_si} in the SI")

figs_now = ["S" + m.group(1) + m.group(2) for p in si.paragraphs for m in [CAP_FIG.match(p.text)] if m]
print("SI figures now:", figs_now)
missing = set()
for p in all_paragraphs(main) + all_paragraphs(si):
    for m in REF_RE.finditer(p.text):
        for t in re.findall(TOK, m.group(3)):
            if t not in figs_now: missing.add(t)
if missing: print("WARNING: references to figures that do not exist:", sorted(missing))

main.save(a.main_out); si.save(a.si_out)
with open(Path(a.si_out).parent / "si_figure_renumbering_key.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["kind", "old", "new"])
    for o, n in fig_map.items():
        if o != "99": w.writerow(["figure", "S" + o, "S" + n])
print("wrote", a.main_out, a.si_out, "and si_figure_renumbering_key.csv")
